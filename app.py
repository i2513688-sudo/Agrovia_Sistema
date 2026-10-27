import json
import hashlib
import html
import io
import math
import os
import re
import secrets
import smtplib
import ssl
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_CEILING, ROUND_HALF_UP
from email.message import EmailMessage
from email.utils import formataddr, parseaddr
from functools import wraps
from urllib.error import URLError
from urllib.parse import urlsplit
from urllib.request import Request as URLRequest, urlopen

from flask import (
    Flask,
    abort,
    render_template,
    request,
    g,
    redirect,
    url_for,
    session,
    flash,
    jsonify,
    send_file,
)

from jinja2.ext import Extension
from markupsafe import Markup, escape
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from mysql.connector import Error as MySQLDatabaseError

from db import query, execute, get_connection


# ==========================================================
# CONFIGURACIÓN
# ==========================================================

load_dotenv()

app = Flask(__name__)
entorno = os.getenv("APP_ENV", "development").lower()
secret_key = os.getenv("SECRET_KEY")
if not secret_key:
    if entorno == "production":
        raise RuntimeError("SECRET_KEY debe configurarse antes de iniciar en producción.")
    secret_key = secrets.token_hex(32)
    app.logger.warning("SECRET_KEY no configurada; se generó una clave temporal para desarrollo.")
elif entorno == "production" and len(secret_key) < 32:
    raise RuntimeError("SECRET_KEY debe tener al menos 32 caracteres en producción.")
app.secret_key = secret_key
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv(
        "SESSION_COOKIE_SECURE", "true" if entorno == "production" else "false"
    ).lower() in {"1", "true", "yes"},
)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024
PUNTOS_POR_OPERACION_CALIFICADA = 5
UMBRAL_LOTE_BAJO_STOCK = Decimal("0.20")
IMAGENES_PRODUCTOS = {
    "papa": "papa.png",
    "choclo": "choclo.png",
    "haba": "habas.png",
    "ajo": "ajo.png",
    "apio": "apio.png",
    "col/repollo": "col.png",
    "repollo": "col.png",
    "lechuga": "lechuga.png",
    "quinua": "quinua.png",
    "acelga": "Acelga.png",
}


def productos_con_imagen(productos):
    resultado = []
    for producto in productos:
        imagen = IMAGENES_PRODUCTOS.get(producto["nombre"].strip().casefold())
        if imagen:
            resultado.append({**producto, "imagen": imagen})
    return resultado


class CSRFFormExtension(Extension):
    def preprocess(self, source, *args, **kwargs):
        del args, kwargs

        def incluir_token(match):
            etiqueta = match.group(0)
            es_post = re.search(
                r"\bmethod\s*=\s*(?:['\"]\s*post\s*['\"]|post\b)",
                etiqueta,
                re.IGNORECASE
            )
            return f"{etiqueta}{{{{ csrf_input() }}}}" if es_post else etiqueta

        return re.sub(r"<form\b[^>]*>", incluir_token, source, flags=re.IGNORECASE)


app.jinja_env.add_extension(CSRFFormExtension)


def obtener_token_csrf():
    if "_csrf_token" not in session:
        session["_csrf_token"] = secrets.token_urlsafe(32)
    return session["_csrf_token"]


def csrf_input():
    return Markup(
        '<input type="hidden" name="csrf_token" value="{}">'.format(
            escape(obtener_token_csrf())
        )
    )


app.jinja_env.globals["csrf_input"] = csrf_input


@app.before_request
def proteger_solicitudes_de_escritura():
    if request.method not in {"POST", "PUT", "PATCH", "DELETE"}:
        return

    token_enviado = request.form.get("csrf_token") or request.headers.get("X-CSRF-Token")
    token_sesion = session.get("_csrf_token")
    if not token_enviado or not token_sesion or not secrets.compare_digest(
        token_enviado, token_sesion
    ):
        abort(400, description="Token CSRF ausente o inválido.")


def enviar_correo_recuperacion(destinatario, enlace):
    host = os.getenv("SMTP_HOST", "").strip()
    usuario = os.getenv("SMTP_USER", "").strip()
    remitente_configurado = os.getenv("SMTP_FROM", "").strip() or usuario
    _, remitente = parseaddr(remitente_configurado)
    if not host or "@" not in remitente or not destinatario:
        app.logger.error(
            "No se puede enviar recuperación: faltan SMTP_HOST, SMTP_FROM/SMTP_USER "
            "o el correo registrado del usuario."
        )
        return False

    try:
        puerto = int(os.getenv("SMTP_PORT", "587").strip())
    except ValueError:
        app.logger.error("SMTP_PORT debe ser un número válido.")
        return False
    if not 1 <= puerto <= 65535:
        app.logger.error("SMTP_PORT está fuera del rango permitido.")
        return False

    usar_tls = os.getenv("SMTP_USE_TLS", "true").strip().lower() in {
        "1", "true", "yes"
    }
    usar_ssl = os.getenv("SMTP_USE_SSL", "false").strip().lower() in {
        "1", "true", "yes"
    }
    contrasena = "".join(os.getenv("SMTP_PASSWORD", "").split())
    if usar_tls and usar_ssl:
        app.logger.error("Configura SMTP_USE_TLS o SMTP_USE_SSL, no ambos.")
        return False
    if bool(usuario) != bool(contrasena):
        app.logger.error("SMTP_USER y SMTP_PASSWORD deben configurarse juntos.")
        return False

    mensaje = EmailMessage()
    mensaje["Subject"] = "Recuperación de cuenta Agrovia"
    mensaje["From"] = formataddr(("Agrovia", remitente))
    mensaje["To"] = destinatario
    mensaje.set_content(
        "Hola,\n\n"
        "Recibimos una solicitud para cambiar la contraseña de tu cuenta de Agrovia. "
        "Abre el siguiente enlace dentro de los próximos 30 minutos:\n\n"
        f"{enlace}\n\n"
        "Si no solicitaste el cambio, ignora este mensaje.\n\n"
        "Atentamente,\n"
        "Equipo Agrovia"
    )

    try:
        if usar_ssl:
            conexion = smtplib.SMTP_SSL(
                host, puerto, timeout=10, context=ssl.create_default_context()
            )
        else:
            conexion = smtplib.SMTP(host, puerto, timeout=10)
        with conexion as servidor:
            if usar_tls:
                servidor.starttls(context=ssl.create_default_context())
            if usuario and contrasena:
                servidor.login(usuario, contrasena)
            rechazados = servidor.send_message(mensaje)
            if rechazados:
                codigos = sorted({
                    respuesta[0] for respuesta in rechazados.values()
                })
                app.logger.error(
                    "El servidor SMTP rechazó el destinatario del correo de "
                    "recuperación (código SMTP: %s).",
                    ", ".join(str(codigo) for codigo in codigos)
                )
                return False
        app.logger.info("El servidor SMTP aceptó el correo de recuperación.")
        return True
    except smtplib.SMTPAuthenticationError as error:
        app.logger.error(
            "El servidor SMTP rechazó la autenticación (código SMTP: %s). "
            "Verifica usuario y contraseña de aplicación.",
            error.smtp_code
        )
        return False
    except smtplib.SMTPRecipientsRefused as error:
        codigos = sorted({
            respuesta[0] for respuesta in error.recipients.values()
        })
        app.logger.error(
            "El servidor SMTP rechazó el destinatario del correo de "
            "recuperación (código SMTP: %s).",
            ", ".join(str(codigo) for codigo in codigos)
        )
        return False
    except smtplib.SMTPResponseException as error:
        app.logger.error(
            "El servidor SMTP rechazó el correo de recuperación "
            "(código SMTP: %s).",
            error.smtp_code
        )
        return False
    except (OSError, smtplib.SMTPException) as error:
        app.logger.error(
            "No se pudo enviar el correo de recuperación (tipo de error: %s).",
            type(error).__name__
        )
        return False


# ==========================================================
# FUNCIONES AUXILIARES
# ==========================================================

def login_required(view):

    @wraps(view)
    def wrapper(*args, **kwargs):

        if "user_id" not in session:
            return redirect(url_for("login"))

        identidad = query(
            """
            SELECT r.nombre AS nombre_rol, p.codigo AS codigo_permiso
            FROM usuarios u
            LEFT JOIN usuarios_roles ur
                ON ur.id_usuario = u.id_usuario
               AND ur.estado = 'ACTIVO'
            LEFT JOIN roles r
                ON r.id_rol = ur.id_rol
               AND r.estado = 'ACTIVO'
            LEFT JOIN roles_permisos rp
                ON rp.id_rol = r.id_rol
               AND rp.estado = 'ACTIVO'
            LEFT JOIN permisos p
                ON p.id_permiso = rp.id_permiso
               AND p.estado = 'ACTIVO'
            WHERE u.id_usuario = %s
              AND u.estado = 'ACTIVO'
            """,
            (session["user_id"],)
        )
        if not identidad:
            session.clear()
            flash("Tu sesión terminó porque la cuenta ya no está activa.", "warning")
            return redirect(url_for("login"))

        session["roles"] = sorted({
            fila["nombre_rol"] for fila in identidad if fila.get("nombre_rol")
        })
        session["permissions"] = sorted({
            fila["codigo_permiso"] for fila in identidad if fila.get("codigo_permiso")
        })
        g.authorization = {
            "roles": session["roles"],
            "permissions": session["permissions"],
        }

        return view(*args, **kwargs)

    return wrapper


def tiene_rol(nombre_rol):

    return nombre_rol in session.get("roles", [])


def tiene_permiso(codigo):
    usuario_id = session.get("user_id")
    if not usuario_id:
        return False
    return bool(query(
        """
        SELECT 1
        FROM usuarios u
        INNER JOIN usuarios_roles ur
            ON ur.id_usuario = u.id_usuario
        INNER JOIN roles r
            ON r.id_rol = ur.id_rol
        INNER JOIN roles_permisos rp
            ON rp.id_rol = r.id_rol
        INNER JOIN permisos p
            ON p.id_permiso = rp.id_permiso
        WHERE u.id_usuario = %s
          AND u.estado = 'ACTIVO'
          AND ur.estado = 'ACTIVO'
          AND r.estado = 'ACTIVO'
          AND rp.estado = 'ACTIVO'
          AND p.estado = 'ACTIVO'
          AND p.codigo = %s
        LIMIT 1
        """,
        (usuario_id, codigo),
        one=True
    ))


def es_admin():
    usuario_id = session.get("user_id")
    if not usuario_id:
        return False
    return bool(query(
        """
        SELECT 1
        FROM usuarios u
        INNER JOIN usuarios_roles ur
            ON ur.id_usuario = u.id_usuario
        INNER JOIN roles r
            ON r.id_rol = ur.id_rol
        WHERE u.id_usuario = %s
          AND u.estado = 'ACTIVO'
          AND ur.estado = 'ACTIVO'
          AND r.estado = 'ACTIVO'
          AND r.nombre = 'ADMINISTRADOR'
        LIMIT 1
        """,
        (usuario_id,),
        one=True
    ))


def registrar_auditoria(
    usuario_id=None,
    rol=None,
    accion="",
    modulo="",
    entidad_tipo="",
    entidad_id=None,
    valor_anterior=None,
    valor_nuevo=None,
    direccion_ip=None
):
    """Registra un evento de auditoría en la tabla correspondiente."""

    if usuario_id is None:
        usuario_id = session.get("user_id")

    if rol is None:
        roles = session.get("roles", [])
        rol = roles[0] if roles else None

    if direccion_ip is None:
        direccion_ip = request.remote_addr if request else None

    if valor_anterior is not None and not isinstance(valor_anterior, str):
        valor_anterior = json.dumps(valor_anterior, ensure_ascii=False)

    if valor_nuevo is not None and not isinstance(valor_nuevo, str):
        valor_nuevo = json.dumps(valor_nuevo, ensure_ascii=False)

    execute(
        """
        INSERT INTO auditoria
        (
            id_usuario,
            rol,
            accion,
            modulo,
            entidad_tipo,
            entidad_id,
            valor_anterior,
            valor_nuevo,
            direccion_ip,
            fecha_hora
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            NOW()
        )
        """,
        (
            usuario_id,
            rol,
            accion,
            modulo,
            entidad_tipo,
            entidad_id,
            valor_anterior,
            valor_nuevo,
            direccion_ip
        )
    )


def registrar_historial_estado(entidad_tipo, entidad_id, estado_anterior, estado_nuevo, observacion=None):
    execute(
        """
        INSERT INTO historial_estados
        (
            entidad_tipo,
            entidad_id,
            estado_anterior,
            estado_nuevo,
            id_usuario,
            observacion
        )
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (
            entidad_tipo,
            entidad_id,
            estado_anterior,
            estado_nuevo,
            session.get("user_id"),
            observacion
        )
    )


def calcular_ruta_osrm(latitud_origen, longitud_origen, latitud_destino, longitud_destino):
    base_url = os.getenv("OSRM_ROUTING_URL", "").strip().rstrip("/")
    if not base_url:
        return None
    servidor_ruta = urlsplit(base_url)
    if (
        not servidor_ruta.hostname
        or (
            servidor_ruta.scheme != "https"
            and not (
                servidor_ruta.scheme == "http"
                and servidor_ruta.hostname in {"localhost", "127.0.0.1", "::1"}
            )
        )
    ):
        app.logger.error(
            "OSRM_ROUTING_URL debe usar HTTPS, excepto para un servicio local."
        )
        return None

    coordenadas = (
        latitud_origen, longitud_origen, latitud_destino, longitud_destino
    )
    try:
        lat_origen, lon_origen, lat_destino, lon_destino = (
            float(valor) for valor in coordenadas
        )
    except (TypeError, ValueError):
        return None

    if (
        not all(math.isfinite(valor) for valor in (
            lat_origen, lon_origen, lat_destino, lon_destino
        ))
        or not -90 <= lat_origen <= 90
        or not -90 <= lat_destino <= 90
        or not -180 <= lon_origen <= 180
        or not -180 <= lon_destino <= 180
    ):
        return None

    url = (
        f"{base_url}/route/v1/driving/"
        f"{lon_origen},{lat_origen};{lon_destino},{lat_destino}"
        "?overview=false"
    )
    solicitud = URLRequest(url, headers={"User-Agent": "Agrovia/1.0"})
    try:
        with urlopen(solicitud, timeout=5) as respuesta:
            datos = json.loads(respuesta.read().decode("utf-8"))
    except (URLError, TimeoutError, ValueError, OSError):
        app.logger.warning("No se pudo consultar la ruta vial de OSRM.")
        return None

    if not isinstance(datos, dict):
        app.logger.warning("OSRM devolvió una respuesta de ruta no válida.")
        return None
    rutas = datos.get("routes") or []
    if datos.get("code") != "Ok" or not rutas:
        return None

    try:
        distancia = Decimal(str(rutas[0]["distance"])) / Decimal("1000")
        duracion_min = int(math.ceil(float(rutas[0]["duration"]) / 60))
    except (KeyError, TypeError, ValueError, InvalidOperation):
        app.logger.warning("OSRM devolvió una ruta con datos incompletos.")
        return None
    if not distancia.is_finite() or distancia <= 0 or duracion_min <= 0:
        return None
    return {
        "distancia_km": distancia.quantize(
            Decimal("0.01"), rounding=ROUND_HALF_UP
        ),
        "duracion_estimada_min": duracion_min,
    }


def crear_notificacion(id_usuario, titulo, mensaje, tipo="INFO", entidad_tipo=None, entidad_id=None):
    """Crea una notificación para un usuario."""

    execute(
        """
        INSERT INTO notificaciones
        (
            id_usuario,
            titulo,
            mensaje,
            tipo,
            entidad_tipo,
            entidad_id,
            leida,
            fecha
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            0,
            NOW()
        )
        """,
        (
            id_usuario,
            titulo,
            mensaje,
            tipo,
            entidad_tipo,
            entidad_id
        )
    )


def notificar_administradores(
    titulo, mensaje, tipo, entidad_tipo, entidad_id, permiso=None
):
    if permiso:
        administradores = query(
            """
            SELECT DISTINCT u.id_usuario
            FROM usuarios u
            INNER JOIN usuarios_roles ur ON ur.id_usuario = u.id_usuario
            INNER JOIN roles r ON r.id_rol = ur.id_rol
            INNER JOIN roles_permisos rp ON rp.id_rol = r.id_rol
            INNER JOIN permisos p ON p.id_permiso = rp.id_permiso
            WHERE u.estado = 'ACTIVO'
              AND ur.estado = 'ACTIVO'
              AND r.estado = 'ACTIVO'
              AND rp.estado = 'ACTIVO'
              AND p.estado = 'ACTIVO'
              AND p.codigo = %s
            """,
            (permiso,)
        )
    else:
        administradores = query(
            """
            SELECT DISTINCT u.id_usuario
            FROM usuarios u
            INNER JOIN usuarios_roles ur ON ur.id_usuario = u.id_usuario
            INNER JOIN roles r ON r.id_rol = ur.id_rol
            WHERE u.estado = 'ACTIVO'
              AND ur.estado = 'ACTIVO'
              AND r.estado = 'ACTIVO'
              AND r.nombre = 'ADMINISTRADOR'
            """
        )
    for administrador in administradores:
        crear_notificacion(
            administrador["id_usuario"],
            titulo,
            mensaje,
            tipo,
            entidad_tipo,
            entidad_id
        )


def notificar_participantes_operacion(
    id_operacion, titulo, mensaje, tipo="INFO"
):
    participantes = query(
        """
        SELECT DISTINCT u.id_usuario
        FROM operaciones o
        INNER JOIN productores pr ON pr.id_productor = o.id_productor
        INNER JOIN usuarios u ON u.id_usuario = pr.id_usuario
        WHERE o.id_operacion = %s
        UNION
        SELECT DISTINCT u.id_usuario
        FROM operaciones o
        INNER JOIN compradores co ON co.id_comprador = o.id_comprador
        INNER JOIN usuarios u ON u.id_usuario = co.id_usuario
        WHERE o.id_operacion = %s
        """,
        (id_operacion, id_operacion)
    )
    for participante in participantes:
        crear_notificacion(
            participante["id_usuario"],
            titulo,
            mensaje,
            tipo,
            "OPERACION",
            id_operacion
        )


def notificar_participantes_incidencia(
    id_incidencia, titulo, mensaje, tipo="INFO"
):
    participantes = query(
        """
        SELECT DISTINCT u.id_usuario
        FROM incidencias i
        INNER JOIN operaciones o ON o.id_operacion = i.id_operacion
        INNER JOIN productores pr ON pr.id_productor = o.id_productor
        INNER JOIN usuarios u ON u.id_usuario = pr.id_usuario
        WHERE i.id_incidencia = %s
        UNION
        SELECT DISTINCT u.id_usuario
        FROM incidencias i
        INNER JOIN operaciones o ON o.id_operacion = i.id_operacion
        INNER JOIN compradores co ON co.id_comprador = o.id_comprador
        INNER JOIN usuarios u ON u.id_usuario = co.id_usuario
        WHERE i.id_incidencia = %s
        """,
        (id_incidencia, id_incidencia)
    )
    for participante in participantes:
        crear_notificacion(
            participante["id_usuario"],
            titulo,
            mensaje,
            tipo,
            "INCIDENCIA",
            id_incidencia
        )


def acreditar_puntos_operacion_calificada(id_operacion):
    """Acredita puntos una sola vez cuando la operación y ambas calificaciones están completas."""
    conn = get_connection()
    cur = conn.cursor(dictionary=True)
    usuarios_acreditados = []
    try:
        conn.start_transaction()
        cur.execute(
            """
            SELECT o.codigo_operacion, o.estado,
                   up.id_usuario AS id_usuario_productor,
                   uc.id_usuario AS id_usuario_comprador
            FROM operaciones o
            INNER JOIN productores pr ON pr.id_productor = o.id_productor
            INNER JOIN usuarios up ON up.id_usuario = pr.id_usuario
            INNER JOIN compradores co ON co.id_comprador = o.id_comprador
            INNER JOIN usuarios uc ON uc.id_usuario = co.id_usuario
            WHERE o.id_operacion = %s
            FOR UPDATE
            """,
            (id_operacion,)
        )
        operacion = cur.fetchone()
        if not operacion or operacion["estado"] != "COMPLETADA":
            conn.commit()
            return usuarios_acreditados

        participantes = {
            operacion["id_usuario_productor"],
            operacion["id_usuario_comprador"]
        }
        cur.execute(
            """
            SELECT DISTINCT id_usuario_califica
            FROM calificaciones
            WHERE id_operacion = %s
              AND estado = 'ACTIVO'
            """,
            (id_operacion,)
        )
        usuarios_que_calificaron = {
            fila["id_usuario_califica"] for fila in cur.fetchall()
        }
        if not participantes.issubset(usuarios_que_calificaron):
            conn.commit()
            return usuarios_acreditados

        motivo = f"Bono por operación calificada {operacion['codigo_operacion']}"
        cur.execute(
            """
            SELECT id_usuario
            FROM movimientos_puntos
            WHERE id_operacion = %s
              AND tipo = 'CREDITO'
              AND motivo = %s
            FOR UPDATE
            """,
            (id_operacion, motivo)
        )
        usuarios_ya_acreditados = {
            fila["id_usuario"] for fila in cur.fetchall()
        }
        for id_usuario in participantes - usuarios_ya_acreditados:
            cur.execute(
                """
                INSERT INTO movimientos_puntos
                    (id_usuario, tipo, puntos, motivo, id_operacion)
                VALUES (%s, 'CREDITO', %s, %s, %s)
                """,
                (id_usuario, PUNTOS_POR_OPERACION_CALIFICADA, motivo, id_operacion)
            )
            cur.execute(
                """
                INSERT INTO notificaciones
                    (id_usuario, titulo, mensaje, tipo, entidad_tipo, entidad_id, leida, fecha)
                VALUES (%s, %s, %s, 'INFO', 'OPERACION', %s, 0, NOW())
                """,
                (
                    id_usuario,
                    "Puntos por operación completada",
                    f"Recibiste {PUNTOS_POR_OPERACION_CALIFICADA} puntos por completar y calificar la operación {operacion['codigo_operacion']}.",
                    id_operacion
                )
            )
            usuarios_acreditados.append(id_usuario)

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()

    return usuarios_acreditados


def guardar_archivo_entidad(archivo, entidad_tipo, entidad_id, id_usuario):

    if not archivo or not archivo.filename:
        return None

    nombre_original = secure_filename(archivo.filename)
    extension = os.path.splitext(nombre_original)[1].lower()
    extensiones_permitidas = {".jpg", ".jpeg", ".png", ".webp"}

    if extension not in extensiones_permitidas:
        raise ValueError("La imagen debe ser JPG, JPEG, PNG o WEBP.")

    upload_dir = os.path.join(
        app.static_folder,
        "uploads",
        "solicitudes"
    )

    os.makedirs(upload_dir, exist_ok=True)

    nombre_archivo = f"{uuid.uuid4().hex}{extension}"
    ruta_guardada = os.path.join(upload_dir, nombre_archivo)

    archivo.save(ruta_guardada)

    ruta_clave = os.path.join(
        "uploads",
        "solicitudes",
        nombre_archivo
    ).replace("\\", "/")

    codigo = generar_codigo("ARC", "archivos", "codigo_archivo")

    execute(
        """
        INSERT INTO archivos
        (
            codigo_archivo,
            tipo_archivo,
            entidad_tipo,
            entidad_id,
            nombre_original,
            nombre_almacenado,
            ruta_clave,
            proveedor_almacenamiento,
            mime_type,
            id_usuario_carga,
            visibilidad,
            estado
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            'LOCAL',
            %s,
            %s,
            'INTERNO',
            'ACTIVO'
        )
        """,
        (
            codigo,
            "IMAGEN",
            entidad_tipo,
            entidad_id,
            nombre_original,
            nombre_archivo,
            ruta_clave,
            archivo.mimetype,
            id_usuario
        )
    )

    return ruta_clave


def generar_codigo(prefijo, tabla, campo):
    """
    Genera un código sencillo y evita duplicados.
    """

    ultimo = query(
        f"""
        SELECT {campo}
        FROM {tabla}
        WHERE {campo} LIKE %s
        ORDER BY {campo} DESC
        LIMIT 1
        """,
        (prefijo + "%",),
        one=True
    )

    numero = 1

    if ultimo and ultimo[campo]:

        try:
            numero = int(ultimo[campo].split("-")[-1]) + 1
        except:
            numero = 1

    return f"{prefijo}-{numero:05d}"


def obtener_comprador_actual():
    """Devuelve el comprador activo del usuario autenticado."""

    if "COMPRADOR" not in session.get("roles", []):
        return None

    return query(
        """
        SELECT *
        FROM compradores
        WHERE id_usuario = %s
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        (session["user_id"],),
        one=True
    )


def documento_facturacion_valido(tipo_documento, numero_documento):
    numero_documento = numero_documento.strip()
    if tipo_documento == "DNI":
        return bool(re.fullmatch(r"\d{8}", numero_documento))
    if tipo_documento == "RUC":
        if not re.fullmatch(r"\d{11}", numero_documento):
            return False
        pesos = (5, 4, 3, 2, 7, 6, 5, 4, 3, 2)
        residuo = sum(int(digito) * peso for digito, peso in zip(numero_documento[:10], pesos)) % 11
        verificador = 11 - residuo
        if verificador == 10:
            verificador = 0
        elif verificador == 11:
            verificador = 1
        return numero_documento[-1] == str(verificador)
    return True


def calcular_valorizacion_solicitud(id_solicitud):
    """
    Calcula una valorización real a partir del precio de mercado,
    parámetros del cultivo y la producción estimada del lote.
    """

    solicitud = query(
        """
        SELECT s.*, p.nombre AS producto, pred.nombre_predio
        FROM solicitudes_lote s
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        INNER JOIN predios pred
            ON pred.id_predio = s.id_predio
        WHERE s.id_solicitud = %s
        LIMIT 1
        """,
        (id_solicitud,),
        one=True
    )

    if not solicitud:
        return None

    id_variedad = solicitud.get("id_variedad")
    referencia_precio = query(
        """
        SELECT precio_actual AS precio, fecha_precio_actual AS fecha_precio,
               promedio_30_dias AS promedio, id_variedad
        FROM vw_precio_referencia_actual
        WHERE id_producto = %s
          AND (id_variedad = %s OR id_variedad IS NULL)
        ORDER BY (id_variedad = %s) DESC
        LIMIT 1
        """,
        (solicitud["id_producto"], id_variedad, id_variedad),
        one=True
    )

    if referencia_precio and referencia_precio.get("precio"):
        precio_actual = referencia_precio
        promedio_precio = referencia_precio
    else:
        precio_actual = query(
            """
            SELECT id_variedad, precio, fecha_precio
            FROM precios_mercado
            WHERE id_producto = %s
              AND estado = 'ACTIVO'
              AND (id_variedad = %s OR id_variedad IS NULL)
            ORDER BY (id_variedad = %s) DESC, fecha_precio DESC, id_precio DESC
            LIMIT 1
            """,
            (solicitud["id_producto"], id_variedad, id_variedad),
            one=True
        )
        promedio_precio = None
        if precio_actual:
            promedio_precio = query(
                """
                SELECT AVG(precio) AS promedio
                FROM precios_mercado
                WHERE id_producto = %s
                  AND estado = 'ACTIVO'
                  AND (id_variedad = %s OR (id_variedad IS NULL AND %s IS NULL))
                  AND fecha_precio >= (
                      SELECT MAX(fecha_precio)
                      FROM precios_mercado
                      WHERE id_producto = %s
                        AND estado = 'ACTIVO'
                        AND (id_variedad = %s OR (id_variedad IS NULL AND %s IS NULL))
                  ) - INTERVAL 30 DAY
                """,
                (
                    solicitud["id_producto"],
                    precio_actual["id_variedad"],
                    precio_actual["id_variedad"],
                    solicitud["id_producto"],
                    precio_actual["id_variedad"],
                    precio_actual["id_variedad"]
                ),
                one=True
            )

    param_cultivo = query(
        """
        SELECT *
        FROM parametros_cultivo
        WHERE id_producto = %s
          AND estado = 'ACTIVO'
        ORDER BY id_parametro_cultivo DESC
        LIMIT 1
        """,
        (solicitud["id_producto"],),
        one=True
    )

    precio_actual_val = float(precio_actual["precio"]) if precio_actual and precio_actual.get("precio") else 0.0
    promedio_val = float(promedio_precio["promedio"]) if promedio_precio and promedio_precio.get("promedio") else 0.0
    if not precio_actual_val and not promedio_val:
        return None

    factor_zona = 1.0
    factor_temporada = 1.0
    factor_condicion = 1.0

    parametro_precio_actual = query(
        """
        SELECT valor
        FROM parametros_comerciales
        WHERE codigo = 'PESO_PRECIO_ACTUAL'
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        one=True
    )

    parametro_promedio = query(
        """
        SELECT valor
        FROM parametros_comerciales
        WHERE codigo = 'PESO_PROMEDIO'
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        one=True
    )

    peso_actual = float(parametro_precio_actual["valor"]) if parametro_precio_actual and parametro_precio_actual.get("valor") else 60.0
    peso_promedio = float(parametro_promedio["valor"]) if parametro_promedio and parametro_promedio.get("valor") else 40.0

    precio_base = 0.0
    if precio_actual_val or promedio_val:
        precio_base = (
            (precio_actual_val * peso_actual) +
            (promedio_val * peso_promedio)
        ) / (peso_actual + peso_promedio)

    precio_ajustado = precio_base * factor_zona * factor_temporada * factor_condicion

    rendimiento_min = float((param_cultivo or {}).get("rendimiento_min_kg_ha", 0) or 0)
    rendimiento_max = float((param_cultivo or {}).get("rendimiento_max_kg_ha", 0) or 0)

    if rendimiento_min or rendimiento_max:
        rendimiento_estimado = (rendimiento_min + rendimiento_max) / 2.0
    else:
        rendimiento_estimado = float(solicitud["area_cultivada"]) * 5000.0

    produccion_estimada = float(solicitud["area_cultivada"]) * rendimiento_estimado / 1000.0
    if produccion_estimada <= 0:
        produccion_estimada = float(solicitud["area_cultivada"]) * 5000.0 / 1000.0

    valor_lote = round(produccion_estimada * precio_ajustado, 2)

    registro = query(
        """
        SELECT *
        FROM valorizaciones
        WHERE id_solicitud = %s
        LIMIT 1
        """,
        (id_solicitud,),
        one=True
    )

    if registro:
        if registro["modo_calculo"] == "MONTO_FIJO":
            return registro["id_valorizacion"]

        execute(
            """
            UPDATE valorizaciones
            SET modo_calculo = 'AUTOMATICO',
                precio_actual = %s,
                precio_promedio = %s,
                peso_actual = %s,
                peso_historico = %s,
                factor_zona = %s,
                factor_temporada = %s,
                factor_condicion = %s,
                precio_base = %s,
                precio_ajustado = %s,
                produccion_estimada = %s,
                valor_lote = %s,
                estado = 'CALCULADA',
                fecha_calculo = NOW()
            WHERE id_valorizacion = %s
            """,
            (
                precio_actual_val,
                promedio_val,
                peso_actual,
                peso_promedio,
                factor_zona,
                factor_temporada,
                factor_condicion,
                precio_base,
                precio_ajustado,
                produccion_estimada,
                valor_lote,
                registro["id_valorizacion"]
            )
        )
        return registro["id_valorizacion"]

    codigo_valorizacion = generar_codigo(
        "VAL",
        "valorizaciones",
        "codigo_valorizacion"
    )

    id_valorizacion = execute(
        """
        INSERT INTO valorizaciones
        (
            codigo_valorizacion,
            id_solicitud,
            modo_calculo,
            precio_actual,
            precio_promedio,
            peso_actual,
            peso_historico,
            factor_zona,
            factor_temporada,
            factor_condicion,
            precio_base,
            precio_ajustado,
            produccion_estimada,
            valor_lote,
            estado,
            fecha_calculo,
            fecha_presentacion,
            fecha_respuesta,
            respuesta_productor
        )
        VALUES
        (
            %s,
            %s,
            'AUTOMATICO',
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            'CALCULADA',
            NOW(),
            NOW(),
            NOW(),
            'PENDIENTE'
        )
        """,
        (
            codigo_valorizacion,
            id_solicitud,
            precio_actual_val,
            promedio_val,
            peso_actual,
            peso_promedio,
            factor_zona,
            factor_temporada,
            factor_condicion,
            precio_base,
            precio_ajustado,
            produccion_estimada,
            valor_lote
        )
    )

    return id_valorizacion


def calcular_cotizacion_lote(id_lote, id_comprador):
    """Genera una cotización de operación con valores base y márgenes."""

    lote = query(
        """
        SELECT
            l.*,
            s.id_productor,
            s.id_producto,
            s.area_cultivada,
            p.nombre AS producto,
            CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, '')) AS productor
        FROM lotes l
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        INNER JOIN productores pr
            ON pr.id_productor = s.id_productor
        INNER JOIN usuarios u
            ON u.id_usuario = pr.id_usuario
        WHERE l.id_lote = %s
        LIMIT 1
        """,
        (id_lote,),
        one=True
    )

    if not lote:
        return None

    valorizacion = query(
        """
        SELECT *
        FROM valorizaciones
        WHERE id_solicitud = %s
        ORDER BY fecha_calculo DESC
        LIMIT 1
        """,
        (lote["id_solicitud"],),
        one=True
    )

    valor_produccion = float((valorizacion or {}).get("valor_lote") or lote.get("precio_publicado") or 0)

    comision = query(
        """
        SELECT valor
        FROM parametros_comerciales
        WHERE codigo = 'COMISION_VENTA_AGROVIA'
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        one=True
    )

    margen_cosecha = query(
        """
        SELECT valor
        FROM parametros_comerciales
        WHERE codigo = 'MARGEN_COSECHA_AGROVIA'
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        one=True
    )

    margen_transporte = query(
        """
        SELECT valor
        FROM parametros_comerciales
        WHERE codigo = 'MARGEN_TRANSPORTE_AGROVIA'
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        one=True
    )

    comision_pct = float((comision or {}).get("valor") or 5.0)
    margen_cosecha_pct = float((margen_cosecha or {}).get("valor") or 5.0)
    margen_transporte_pct = float((margen_transporte or {}).get("valor") or 5.0)

    costo_base_cosecha = round(valor_produccion * 0.08, 2)
    margen_cosecha_monto = round(costo_base_cosecha * (margen_cosecha_pct / 100), 2)
    precio_cosecha_cliente = round(valor_produccion + costo_base_cosecha + margen_cosecha_monto, 2)

    costo_base_flete = 0.0
    margen_transporte_monto = round(costo_base_flete * (margen_transporte_pct / 100), 2)
    precio_flete_cliente = round(costo_base_flete + margen_transporte_monto, 2)

    comision_venta_monto = round(valor_produccion * (comision_pct / 100), 2)
    ingreso_total_agrovia = round(comision_venta_monto + margen_cosecha_monto + margen_transporte_monto, 2)
    total_operacion = round(precio_cosecha_cliente + precio_flete_cliente + comision_venta_monto, 2)

    return {
        "id_lote": id_lote,
        "id_comprador": id_comprador,
        "valor_produccion": round(valor_produccion, 2),
        "costo_base_cosecha": costo_base_cosecha,
        "margen_cosecha_pct": margen_cosecha_pct,
        "margen_cosecha_monto": margen_cosecha_monto,
        "precio_cosecha_cliente": precio_cosecha_cliente,
        "costo_base_flete": costo_base_flete,
        "margen_transporte_pct": margen_transporte_pct,
        "margen_transporte_monto": margen_transporte_monto,
        "precio_flete_cliente": precio_flete_cliente,
        "comision_venta_pct": comision_pct,
        "comision_venta_monto": comision_venta_monto,
        "ingreso_total_agrovia": ingreso_total_agrovia,
        "total_operacion": total_operacion,
        "producto": lote["producto"],
        "productor": lote["productor"],
        "id_productor": lote["id_productor"],
    }


def calcular_costos_logisticos(
    distancia_km,
    peso_kg,
    tipo_vehiculo,
    precio_combustible,
    factor_ruta,
    reserva_capacidad_pct,
    carga_descarga,
    peajes,
    seguro
):
    """Calcula el costo del flete usando los valores vigentes de logística."""
    distancia = Decimal(str(distancia_km))
    peso = Decimal(str(peso_kg))
    rendimiento = Decimal(str(tipo_vehiculo["rendimiento_km_galon"]))
    capacidad = Decimal(str(tipo_vehiculo["capacidad_util_kg"]))
    factor = Decimal(str(factor_ruta))
    reserva = Decimal(str(reserva_capacidad_pct))
    precio_combustible = Decimal(str(precio_combustible))
    costo_operativo_km = Decimal(str(tipo_vehiculo["costo_operativo_km"] or 0))
    costo_fijo_viaje = Decimal(str(tipo_vehiculo["costo_fijo_viaje"] or 0))
    costo_carga_descarga = Decimal(str(carga_descarga or 0))
    costo_peajes = Decimal(str(peajes or 0))
    costo_seguro = Decimal(str(seguro or 0))

    if (
        not distancia.is_finite() or distancia <= 0
        or not peso.is_finite() or peso <= 0
        or not rendimiento.is_finite() or rendimiento <= 0
        or not capacidad.is_finite() or capacidad <= 0
        or not factor.is_finite() or factor <= 0
        or not reserva.is_finite() or reserva < 0 or reserva >= 100
        or not precio_combustible.is_finite() or precio_combustible <= 0
        or not costo_operativo_km.is_finite() or costo_operativo_km < 0
        or not costo_fijo_viaje.is_finite() or costo_fijo_viaje < 0
        or not costo_carga_descarga.is_finite() or costo_carga_descarga < 0
        or not costo_peajes.is_finite() or costo_peajes < 0
        or not costo_seguro.is_finite() or costo_seguro < 0
    ):
        raise ValueError("La distancia, carga o configuración logística no es válida.")

    capacidad_planificada = capacidad * (Decimal("1") - reserva / Decimal("100"))
    cantidad_vehiculos = int(
        (peso / capacidad_planificada).to_integral_value(rounding=ROUND_CEILING)
    )
    if cantidad_vehiculos < 1 or cantidad_vehiculos > 65535:
        raise ValueError("La carga excede la capacidad logística admitida.")

    distancia_total = distancia * factor * cantidad_vehiculos
    costo_combustible = distancia_total / rendimiento * precio_combustible
    costo_operativo = distancia_total * costo_operativo_km
    otros_costos = costo_fijo_viaje * cantidad_vehiculos
    costo_carga_descarga *= cantidad_vehiculos
    costo_base = (
        costo_combustible + costo_operativo + costo_carga_descarga
        + costo_peajes + costo_seguro + otros_costos
    )

    return {
        "cantidad_vehiculos": cantidad_vehiculos,
        "costo_combustible": costo_combustible.quantize(Decimal("0.01")),
        "costo_operativo": costo_operativo.quantize(Decimal("0.01")),
        "otros_costos": otros_costos.quantize(Decimal("0.01")),
        "carga_descarga": costo_carga_descarga.quantize(Decimal("0.01")),
        "peajes": costo_peajes.quantize(Decimal("0.01")),
        "seguro": costo_seguro.quantize(Decimal("0.01")),
        "costo_base_flete": costo_base.quantize(Decimal("0.01")),
        "precio_combustible": precio_combustible,
        "rendimiento": rendimiento
    }


@app.context_processor
def common():
    usuario_id = session.get("user_id")
    autorizacion = getattr(g, "authorization", None)
    if autorizacion is not None:
        roles = autorizacion["roles"]
        permisos = set(autorizacion["permissions"])
    else:
        roles = []
        permisos = set()
    if usuario_id and autorizacion is None:
        roles_activos = query(
            """
            SELECT r.nombre
            FROM usuarios_roles ur
            INNER JOIN roles r ON r.id_rol = ur.id_rol
            INNER JOIN usuarios u ON u.id_usuario = ur.id_usuario
            WHERE ur.id_usuario = %s
              AND ur.estado = 'ACTIVO'
              AND r.estado = 'ACTIVO'
              AND u.estado = 'ACTIVO'
            """,
            (usuario_id,)
        )
        permisos_activos = query(
            """
            SELECT DISTINCT p.codigo
            FROM usuarios_roles ur
            INNER JOIN roles r ON r.id_rol = ur.id_rol
            INNER JOIN roles_permisos rp ON rp.id_rol = r.id_rol
            INNER JOIN permisos p ON p.id_permiso = rp.id_permiso
            INNER JOIN usuarios u ON u.id_usuario = ur.id_usuario
            WHERE ur.id_usuario = %s
              AND u.estado = 'ACTIVO'
              AND ur.estado = 'ACTIVO'
              AND r.estado = 'ACTIVO'
              AND rp.estado = 'ACTIVO'
              AND p.estado = 'ACTIVO'
            """,
            (usuario_id,)
        )
        roles = [rol["nombre"] for rol in roles_activos]
        permisos = {permiso["codigo"] for permiso in permisos_activos}
    return {
        "session_user": session.get("user_name"),
        "session_roles": roles,
        "session_permissions": permisos
    }


# ==========================================================
# INICIO
# ==========================================================

@app.route("/")
def index():

    products = productos_con_imagen(query(
        """
        SELECT
            p.id_producto,
            p.nombre,
            p.descripcion,
            c.nombre AS categoria
        FROM productos p
        INNER JOIN categorias_producto c
            ON c.id_categoria = p.id_categoria
        WHERE p.estado = 'ACTIVO'
        ORDER BY p.nombre
        """
    ))

    lots = query(
        """
        SELECT
            l.id_lote,
            l.codigo_lote,
            p.nombre AS producto,
            l.produccion_min,
            l.produccion_max,
            l.precio_publicado,
            l.fecha_cosecha_inicio,
            l.fecha_cosecha_fin,
            l.estado
        FROM lotes l
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        WHERE l.estado IN ('DISPONIBLE', 'RESERVADO')
        ORDER BY l.fecha_publicacion DESC
        LIMIT 12
        """
    )

    return render_template(
        "index.html",
        products=products,
        lots=lots
    )


# ==========================================================
# LOGIN
# ==========================================================

@app.route("/recuperar-cuenta", methods=["GET", "POST"])
def recuperar_cuenta():
    if request.method == "POST":
        correo = request.form.get("correo", "").strip()
        host_smtp = os.getenv("SMTP_HOST", "").strip()
        remitente = os.getenv("SMTP_FROM", "").strip() or os.getenv(
            "SMTP_USER", ""
        ).strip()
        base_url = os.getenv("APP_BASE_URL", "").strip().rstrip("/")
        partes_base_url = urlsplit(base_url)
        base_url_valida = (
            partes_base_url.scheme in {"http", "https"}
            and bool(partes_base_url.netloc)
            and not partes_base_url.username
            and not partes_base_url.password
            and not any(caracter.isspace() for caracter in base_url)
        )
        if base_url_valida and partes_base_url.hostname in {
            "localhost", "127.0.0.1", "::1"
        }:
            app.logger.warning(
                "APP_BASE_URL apunta a localhost; el enlace solo será accesible "
                "desde el mismo equipo."
            )
        try:
            usuario = query(
                "SELECT id_usuario, correo FROM usuarios WHERE correo = %s AND estado = 'ACTIVO' LIMIT 1",
                (correo,),
                one=True
            ) if correo else None

            if not usuario:
                app.logger.info(
                    "Solicitud de recuperación sin una cuenta activa coincidente."
                )

            if usuario and host_smtp and remitente and base_url_valida:
                recientes = query(
                    """
                    SELECT COUNT(*) AS total
                    FROM tokens_seguridad
                    WHERE id_usuario = %s
                      AND tipo = 'RECUPERAR_PASSWORD'
                      AND fecha_creacion >= DATE_SUB(NOW(), INTERVAL 15 MINUTE)
                    """,
                    (usuario["id_usuario"],),
                    one=True
                )
                if int((recientes or {}).get("total") or 0) >= 3:
                    app.logger.warning(
                        "Solicitud de recuperación limitada por superar "
                        "el máximo de enlaces."
                    )
                else:
                    token = secrets.token_urlsafe(32)
                    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
                    id_token = execute(
                        """
                        INSERT INTO tokens_seguridad
                        (id_usuario, tipo, codigo_hash, fecha_expiracion, utilizado)
                        VALUES (%s, 'RECUPERAR_PASSWORD', %s, DATE_ADD(NOW(), INTERVAL 30 MINUTE), 0)
                        """,
                        (usuario["id_usuario"], token_hash)
                    )
                    enlace = f"{base_url}{url_for('restablecer_password', token=token)}"
                    if not enviar_correo_recuperacion(usuario["correo"], enlace):
                        execute(
                            "UPDATE tokens_seguridad SET utilizado = 1 WHERE id_token = %s",
                            (id_token,)
                        )
            elif usuario:
                app.logger.error(
                    "Recuperación no enviada: configuración incompleta "
                    "(SMTP_HOST=%s, remitente=%s, APP_BASE_URL válida=%s).",
                    bool(host_smtp),
                    bool(remitente),
                    base_url_valida
                )
        except MySQLDatabaseError:
            app.logger.exception(
                "No se pudo procesar la recuperación de contraseña por un error de base de datos."
            )
            flash(
                "No se pudo procesar la solicitud por un problema con la base de datos. "
                "Inténtalo más tarde o contacta con administración.",
                "danger"
            )
            return redirect(url_for("login"))

        return redirect(url_for("login"))

    return render_template("recuperar_password.html")


@app.route("/restablecer-cuenta/<token>", methods=["GET", "POST"])
def restablecer_password(token):
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    token_valido = query(
        """
        SELECT id_token
        FROM tokens_seguridad
        WHERE codigo_hash = %s
          AND tipo = 'RECUPERAR_PASSWORD'
          AND utilizado = 0
          AND fecha_expiracion > NOW()
        LIMIT 1
        """,
        (token_hash,),
        one=True
    )
    if not token_valido:
        flash("El enlace de recuperación no es válido o ya venció.", "warning")
        return redirect(url_for("recuperar_cuenta"))

    if request.method == "POST":
        password = request.form.get("password", "")
        confirmacion = request.form.get("confirmacion", "")
        if len(password) < 8 or password != confirmacion:
            flash("La contraseña debe tener al menos 8 caracteres y coincidir con la confirmación.", "warning")
            return render_template("restablecer_password.html", token=token)

        conn = get_connection()
        cur = conn.cursor(dictionary=True)
        try:
            conn.start_transaction()
            cur.execute(
                """
                SELECT id_token, id_usuario
                FROM tokens_seguridad
                WHERE codigo_hash = %s
                  AND tipo = 'RECUPERAR_PASSWORD'
                  AND utilizado = 0
                  AND fecha_expiracion > NOW()
                LIMIT 1
                FOR UPDATE
                """,
                (token_hash,)
            )
            registro = cur.fetchone()
            if not registro:
                conn.rollback()
                flash("El enlace de recuperación no es válido o ya venció.", "warning")
                return redirect(url_for("recuperar_cuenta"))

            cur.execute(
                "UPDATE usuarios SET password_hash = %s WHERE id_usuario = %s",
                (generate_password_hash(password), registro["id_usuario"])
            )
            cur.execute(
                """
                UPDATE tokens_seguridad
                SET utilizado = 1
                WHERE id_usuario = %s
                  AND tipo = 'RECUPERAR_PASSWORD'
                  AND utilizado = 0
                """,
                (registro["id_usuario"],)
            )
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()
        flash("Contraseña actualizada. Ya puedes iniciar sesión.", "success")
        return redirect(url_for("login"))

    return render_template("restablecer_password.html", token=token)


@app.route("/login", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        login_value = request.form["login"].strip()
        password = request.form["password"]
        session.clear()

        user = query(
            """
            SELECT
                u.*,
                GROUP_CONCAT(r.nombre) AS roles_text
            FROM usuarios u

            LEFT JOIN usuarios_roles ur
                ON ur.id_usuario = u.id_usuario
                AND ur.estado = 'ACTIVO'

            LEFT JOIN roles r
                ON r.id_rol = ur.id_rol
                AND r.estado = 'ACTIVO'

            WHERE u.correo = %s
               OR u.numero_documento = %s

            GROUP BY u.id_usuario
            """,
            (login_value, login_value),
            one=True
        )

        if not user:

            flash(
                "Correo/documento o contraseña incorrectos.",
                "danger"
            )

            return render_template("login.html")

        if not check_password_hash(
            user["password_hash"],
            password
        ):

            flash(
                "Correo/documento o contraseña incorrectos.",
                "danger"
            )

            return render_template("login.html")

        if user["estado"] != "ACTIVO":

            flash(
                "El usuario no está activo.",
                "warning"
            )

            return render_template("login.html")

        roles_usuario = [
            rol.strip()
            for rol in (user["roles_text"] or "").split(",")
            if rol and rol.strip()
        ]
        permisos_usuario = query(
            """
            SELECT DISTINCT p.codigo
            FROM usuarios_roles ur
            INNER JOIN roles r ON r.id_rol = ur.id_rol
            INNER JOIN roles_permisos rp ON rp.id_rol = r.id_rol
            INNER JOIN permisos p ON p.id_permiso = rp.id_permiso
            WHERE ur.id_usuario = %s
              AND ur.estado = 'ACTIVO'
              AND r.estado = 'ACTIVO'
              AND rp.estado = 'ACTIVO'
              AND p.estado = 'ACTIVO'
            """,
            (user["id_usuario"],)
        )
        session["user_id"] = user["id_usuario"]
        session["user_name"] = (
            f'{user["nombres"]} '
            f'{user.get("apellido_paterno") or ""}'
        ).strip()
        session["roles"] = roles_usuario
        session["permissions"] = [
            permiso["codigo"] for permiso in permisos_usuario
        ]

        execute(
            """
            UPDATE usuarios
            SET ultimo_acceso = NOW()
            WHERE id_usuario = %s
            """,
            (user["id_usuario"],)
        )

        if es_admin():
            return redirect(url_for("dashboard"))

        return redirect(url_for("index"))

    return render_template("login.html")


# ==========================================================
# LOGOUT
# ==========================================================

@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("index"))


# ==========================================================
# REGISTRO DE USUARIO
# ==========================================================

@app.route("/registro", methods=["GET", "POST"])
def registro():

    if request.method == "POST":

        nombres = request.form["nombres"].strip()
        apellido_paterno = request.form["apellido_paterno"].strip()
        apellido_materno = request.form["apellido_materno"].strip()

        documento = request.form["numero_documento"].strip()
        telefono = request.form["telefono"].strip()
        correo = request.form["correo"].strip()

        password = request.form["password"]

        tipo_persona = request.form["tipo_persona"]
        codigo_referido = request.form.get("codigo_referido", "").strip().upper()
        referidor = None

        if tipo_persona not in {"PRODUCTOR", "COMPRADOR"}:
            flash("Selecciona un tipo de cuenta válido.", "warning")
            return render_template("registro.html")

        if len(password) < 8:
            flash("La contraseña debe tener al menos 8 caracteres.", "warning")
            return render_template("registro.html")

        if codigo_referido:
            prefijo_referido = "AGROVIA-"
            if not codigo_referido.startswith(prefijo_referido):
                flash("El código de referido no es válido.", "warning")
                return render_template("registro.html")
            referidor = query(
                """
                SELECT id_usuario
                FROM usuarios
                WHERE codigo_usuario = %s
                  AND estado = 'ACTIVO'
                LIMIT 1
                """,
                (codigo_referido[len(prefijo_referido):],),
                one=True
            )
            if not referidor:
                flash("No encontramos una cuenta activa con ese código.", "warning")
                return render_template("registro.html")

        rol = query(
            """
            SELECT id_rol
            FROM roles
            WHERE nombre = %s
              AND estado = 'ACTIVO'
            LIMIT 1
            """,
            (tipo_persona,),
            one=True
        )
        if not rol:
            flash("El tipo de cuenta seleccionado no está disponible.", "warning")
            return render_template("registro.html")

        # Verificar usuario existente

        existente = query(
            """
            SELECT id_usuario
            FROM usuarios
            WHERE numero_documento = %s
               OR correo = %s
            """,
            (documento, correo),
            one=True
        )

        if existente:

            flash(
                "El documento o correo ya está registrado.",
                "danger"
            )

            return render_template("registro.html")

        codigo = generar_codigo(
            "USR",
            "usuarios",
            "codigo_usuario"
        )

        password_hash = generate_password_hash(password)

        id_usuario = execute(
            """
            INSERT INTO usuarios
            (
                codigo_usuario,
                tipo_documento,
                numero_documento,
                nombres,
                apellido_paterno,
                apellido_materno,
                telefono,
                correo,
                password_hash,
                origen_registro,
                estado
            )
            VALUES
            (
                %s,
                'DNI',
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                'WEB',
                'ACTIVO'
            )
            """,
            (
                codigo,
                documento,
                nombres,
                apellido_paterno,
                apellido_materno,
                telefono,
                correo,
                password_hash
            )
        )

        # ---------------------------------------------
        # PRODUCTOR
        # ---------------------------------------------

        if tipo_persona == "PRODUCTOR":

            codigo_productor = generar_codigo(
                "PRO",
                "productores",
                "codigo_productor"
            )

            execute(
                """
                INSERT INTO productores
                (
                    id_usuario,
                    codigo_productor,
                    estado_verificacion,
                    estado
                )
                VALUES
                (
                    %s,
                    %s,
                    'PENDIENTE',
                    'ACTIVO'
                )
                """,
                (
                    id_usuario,
                    codigo_productor
                )
            )

        # ---------------------------------------------
        # COMPRADOR
        # ---------------------------------------------

        else:

            codigo_comprador = generar_codigo(
                "COM",
                "compradores",
                "codigo_comprador"
            )

            execute(
                """
                INSERT INTO compradores
                (
                    id_usuario,
                    codigo_comprador,
                    tipo_comprador,
                    estado_verificacion,
                    estado
                )
                VALUES
                (
                    %s,
                    %s,
                    'PERSONA_CON_NEGOCIO',
                    'PENDIENTE',
                    'ACTIVO'
                )
                """,
                (
                    id_usuario,
                    codigo_comprador
                )
            )

        if rol:

            execute(
                """
                INSERT INTO usuarios_roles
                (
                    id_usuario,
                    id_rol,
                    estado
                )
                VALUES
                (
                    %s,
                    %s,
                    'ACTIVO'
                )
                """,
                (
                    id_usuario,
                    rol["id_rol"]
                )
            )

        if referidor:
            execute(
                """
                INSERT INTO referidos
                (id_usuario_referidor, id_usuario_referido, codigo_utilizado, estado)
                VALUES (%s, %s, %s, 'REGISTRADO')
                """,
                (referidor["id_usuario"], id_usuario, codigo_referido)
            )

        flash(
            "Registro realizado correctamente. Ahora puedes iniciar sesión.",
            "success"
        )

        return redirect(url_for("login"))

    return render_template("registro.html")


# ==========================================================
# DASHBOARD
# ==========================================================

@app.route("/dashboard")
@login_required
def dashboard():

    if not tiene_permiso("REPORTES_VER"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("index"))

    hoy = date.today()
    fecha_fin_default = hoy
    fecha_inicio_default = date(hoy.year, hoy.month, 1)
    for _ in range(5):
        if fecha_inicio_default.month == 1:
            fecha_inicio_default = date(fecha_inicio_default.year - 1, 12, 1)
        else:
            fecha_inicio_default = date(
                fecha_inicio_default.year, fecha_inicio_default.month - 1, 1
            )
    try:
        fecha_desde = date.fromisoformat(
            request.args.get("fecha_desde", "") or fecha_inicio_default.isoformat()
        )
        fecha_hasta = date.fromisoformat(
            request.args.get("fecha_hasta", "") or fecha_fin_default.isoformat()
        )
    except ValueError:
        flash("Selecciona un rango de fechas válido.", "warning")
        return redirect(url_for("dashboard"))
    if fecha_desde > fecha_hasta:
        flash("La fecha inicial no puede ser posterior a la fecha final.", "warning")
        return redirect(url_for("dashboard"))

    meses = []
    cursor_mes = date(fecha_desde.year, fecha_desde.month, 1)
    ultimo_mes = date(fecha_hasta.year, fecha_hasta.month, 1)
    while cursor_mes <= ultimo_mes:
        meses.append(cursor_mes.strftime("%Y-%m"))
        if cursor_mes.month == 12:
            cursor_mes = date(cursor_mes.year + 1, 1, 1)
        else:
            cursor_mes = date(cursor_mes.year, cursor_mes.month + 1, 1)
    nombres_meses = (
        "enero", "febrero", "marzo", "abril", "mayo", "junio",
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"
    )
    etiquetas_meses = [
        f"{nombres_meses[int(mes[5:7]) - 1].capitalize()} {mes[:4]}"
        for mes in meses
    ]
    parametros_fecha = (fecha_desde, fecha_hasta)

    ventas_periodo = query(
        """
        SELECT COALESCE(SUM(total_operacion), 0) AS total
        FROM operaciones
        WHERE estado != 'CANCELADA'
          AND fecha_operacion >= %s
          AND fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
        """,
        parametros_fecha,
        one=True
    )
    stats = {
        "ventas": float((ventas_periodo or {}).get("total") or 0),
        "operaciones": query(
            """
            SELECT COUNT(*) AS total
            FROM operaciones
            WHERE estado != 'CANCELADA'
              AND fecha_operacion >= %s
              AND fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
            """,
            parametros_fecha,
            one=True
        )["total"],
        "solicitudes": query(
            """
            SELECT COUNT(*) AS total
            FROM solicitudes_lote
            WHERE estado IN (
                'ENVIADA', 'EN_REVISION', 'OBSERVADA', 'VISITA_PROGRAMADA',
                'EN_VERIFICACION', 'VERIFICADA', 'EN_VALORIZACION',
                'PENDIENTE_ACEPTACION'
            )
            """,
            one=True
        )["total"],
        "entregas": query(
            """
            SELECT COUNT(*) AS total
            FROM transportes
            WHERE estado IN (
                'ASIGNADO', 'RECOJO_PROGRAMADO', 'CARGA_REALIZADA',
                'EN_RUTA', 'LLEGADA_DESTINO'
            )
            """,
            one=True
        )["total"],
    }

    recent = query(
        """
        SELECT
            o.codigo_operacion,
            o.total_operacion,
            o.estado,
            o.fecha_operacion,
            p.nombre AS producto,

            CONCAT(
                u.nombres,
                ' ',
                COALESCE(u.apellido_paterno,'')
            ) AS productor

        FROM operaciones o

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios u
            ON u.id_usuario = pr.id_usuario

        ORDER BY o.fecha_operacion DESC

        LIMIT 8
        """
    )

    incidencias_pendientes = query(
        """
        SELECT i.codigo_incidencia, i.categoria, i.prioridad, i.estado,
               i.fecha_apertura, o.codigo_operacion
        FROM incidencias i
        LEFT JOIN operaciones o ON o.id_operacion = i.id_operacion
        WHERE i.estado IN ('ABIERTO', 'EN_ATENCION')
        ORDER BY
            CASE i.prioridad
                WHEN 'CRITICA' THEN 1
                WHEN 'ALTA' THEN 2
                WHEN 'MEDIA' THEN 3
                ELSE 4
            END,
            i.fecha_apertura DESC
        LIMIT 8
        """
    )

    ventas_por_mes = query(
        """
        SELECT DATE_FORMAT(fecha_operacion, '%Y-%m') AS mes,
               SUM(total_operacion) AS total
        FROM operaciones
        WHERE estado != 'CANCELADA'
          AND fecha_operacion >= %s
          AND fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
        GROUP BY DATE_FORMAT(fecha_operacion, '%Y-%m')
        ORDER BY mes
        """,
        parametros_fecha
    )
    ventas_por_mes = {fila["mes"]: float(fila["total"] or 0) for fila in ventas_por_mes}
    ventas_mensuales = [
        {
            "mes": mes,
            "etiqueta": etiqueta,
            "total": ventas_por_mes.get(mes, 0),
        }
        for mes, etiqueta in zip(meses, etiquetas_meses)
    ]
    max_ventas = max((fila["total"] for fila in ventas_mensuales), default=0)
    for fila in ventas_mensuales:
        fila["porcentaje"] = (
            fila["total"] / max_ventas * 100 if max_ventas else 0
        )

    conteo_operaciones_por_mes = query(
        """
        SELECT DATE_FORMAT(fecha_operacion, '%Y-%m') AS mes,
               COUNT(*) AS total
        FROM operaciones
        WHERE estado != 'CANCELADA'
          AND fecha_operacion >= %s
          AND fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
        GROUP BY DATE_FORMAT(fecha_operacion, '%Y-%m')
        ORDER BY mes
        """,
        parametros_fecha
    )
    conteo_operaciones_por_mes = {
        fila["mes"]: int(fila["total"] or 0)
        for fila in conteo_operaciones_por_mes
    }
    operaciones_mensuales = [
        {
            "mes": mes,
            "etiqueta": etiqueta,
            "total": conteo_operaciones_por_mes.get(mes, 0),
        }
        for mes, etiqueta in zip(meses, etiquetas_meses)
    ]
    max_operaciones_mes = max(
        (fila["total"] for fila in operaciones_mensuales), default=0
    )
    for fila in operaciones_mensuales:
        fila["porcentaje"] = (
            fila["total"] / max_operaciones_mes * 100
            if max_operaciones_mes else 0
        )

    operaciones_por_estado = query(
        """
        SELECT estado, COUNT(*) AS total
        FROM operaciones
        WHERE fecha_operacion >= %s
          AND fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
        GROUP BY estado
        ORDER BY total DESC, estado
        """,
        parametros_fecha
    )
    max_operaciones_estado = max(
        (int(fila["total"] or 0) for fila in operaciones_por_estado), default=0
    )
    total_operaciones_estado = sum(
        int(fila["total"] or 0) for fila in operaciones_por_estado
    )
    colores_estados = (
        "#318447", "#4285a4", "#d58b21", "#8656a3", "#d15c4b",
        "#4c9b8a", "#b34f83", "#74834c",
    )
    circunferencia_estado = 2 * math.pi * 72
    desfase_estado = 0
    for indice, fila in enumerate(operaciones_por_estado):
        total_estado = int(fila["total"] or 0)
        fila["porcentaje"] = (
            total_estado / max_operaciones_estado * 100
            if max_operaciones_estado else 0
        )
        fila["porcentaje_total"] = (
            total_estado / total_operaciones_estado * 100
            if total_operaciones_estado else 0
        )
        fila["etiqueta_corta"] = (
            fila["estado"] if len(fila["estado"]) <= 22
            else f"{fila['estado'][:19]}..."
        )
        fila["barra_ancho"] = 420 * fila["porcentaje"] / 100
        fila["color"] = colores_estados[indice % len(colores_estados)]
        fila["tramo_circular"] = (
            circunferencia_estado * total_estado / total_operaciones_estado
            if total_operaciones_estado else 0
        )
        fila["desfase_circular"] = desfase_estado
        desfase_estado += fila["tramo_circular"]

    pagos_por_estado = query(
        """
        SELECT p.estado, COUNT(*) AS total, SUM(p.monto) AS monto
        FROM pagos p
        INNER JOIN operaciones o ON o.id_operacion = p.id_operacion
        WHERE p.estado IN ('PENDIENTE', 'RECIBIDO', 'RETENIDO')
          AND o.fecha_operacion >= %s
          AND o.fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
        GROUP BY p.estado
        ORDER BY total DESC, p.estado
        """,
        parametros_fecha
    )
    max_pagos_estado = max(
        (int(fila["total"] or 0) for fila in pagos_por_estado), default=0
    )
    for fila in pagos_por_estado:
        fila["porcentaje"] = (
            int(fila["total"] or 0) / max_pagos_estado * 100
            if max_pagos_estado else 0
        )
        fila["monto"] = float(fila["monto"] or 0)

    lotes_con_stock = query(
        """
        SELECT l.id_lote, l.codigo_lote, p.nombre AS producto,
               l.produccion_max,
               COALESCE(operaciones.comprometido_kg, 0)
                 + COALESCE(reservas.reservado_kg, 0) AS comprometido_kg
        FROM lotes l
        INNER JOIN solicitudes_lote s ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p ON p.id_producto = s.id_producto
        LEFT JOIN (
            SELECT o.id_lote, SUM(cl.peso_kg) AS comprometido_kg
            FROM operaciones o
            INNER JOIN cotizaciones_logisticas cl
                ON cl.id_cotizacion = (
                    SELECT MAX(cl_actual.id_cotizacion)
                    FROM cotizaciones_logisticas cl_actual
                    WHERE cl_actual.id_cotizacion_operacion = o.id_cotizacion_operacion
                )
            WHERE o.estado != 'CANCELADA'
              AND cl.peso_kg IS NOT NULL
            GROUP BY o.id_lote
        ) operaciones ON operaciones.id_lote = l.id_lote
        LEFT JOIN (
            SELECT r.id_lote, SUM(cl.peso_kg) AS reservado_kg
            FROM reservas_lote r
            INNER JOIN cotizaciones_logisticas cl
                ON cl.id_cotizacion = (
                    SELECT MAX(cl_actual.id_cotizacion)
                    FROM cotizaciones_logisticas cl_actual
                    WHERE cl_actual.id_cotizacion_operacion = r.id_cotizacion_operacion
                )
            WHERE r.estado = 'ACTIVA'
              AND r.fecha_vencimiento > NOW()
              AND cl.peso_kg IS NOT NULL
              AND NOT EXISTS (
                  SELECT 1
                  FROM operaciones o
                  WHERE o.id_cotizacion_operacion = r.id_cotizacion_operacion
                    AND o.estado != 'CANCELADA'
              )
            GROUP BY r.id_lote
        ) reservas ON reservas.id_lote = l.id_lote
        WHERE l.estado IN ('DISPONIBLE', 'RESERVADO')
          AND l.produccion_max > 0
        ORDER BY l.fecha_publicacion DESC, l.id_lote DESC
        """
    )
    alertas_inventario = []
    for lote in lotes_con_stock:
        produccion_max = Decimal(str(lote["produccion_max"]))
        comprometido = Decimal(str(lote["comprometido_kg"] or 0))
        restante = max(produccion_max - comprometido, Decimal("0"))
        if restante > produccion_max * UMBRAL_LOTE_BAJO_STOCK:
            continue
        alertas_inventario.append({
            "id_lote": lote["id_lote"],
            "codigo_lote": lote["codigo_lote"],
            "producto": lote["producto"],
            "restante_kg": float(restante),
            "produccion_max_kg": float(produccion_max),
            "porcentaje_restante": float(restante / produccion_max * 100),
            "agotado": restante == 0,
        })

    productos_top = query(
        """
        SELECT s.id_producto, p.nombre, SUM(o.total_operacion) AS total
        FROM operaciones o
        INNER JOIN lotes l ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p ON p.id_producto = s.id_producto
        WHERE o.estado != 'CANCELADA'
          AND o.fecha_operacion >= %s
          AND o.fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
        GROUP BY s.id_producto, p.nombre
        ORDER BY total DESC, p.nombre
        LIMIT 5
        """,
        parametros_fecha
    )
    evolucion_productos = []
    if productos_top:
        ids_productos = [producto["id_producto"] for producto in productos_top]
        placeholders = ", ".join(["%s"] * len(ids_productos))
        evolucion = query(
            f"""
            SELECT s.id_producto, DATE_FORMAT(o.fecha_operacion, '%Y-%m') AS mes,
                   SUM(o.total_operacion) AS total
            FROM operaciones o
            INNER JOIN lotes l ON l.id_lote = o.id_lote
            INNER JOIN solicitudes_lote s ON s.id_solicitud = l.id_solicitud
            WHERE o.estado != 'CANCELADA'
              AND o.fecha_operacion >= %s
              AND o.fecha_operacion < DATE_ADD(%s, INTERVAL 1 DAY)
              AND s.id_producto IN ({placeholders})
            GROUP BY s.id_producto, DATE_FORMAT(o.fecha_operacion, '%Y-%m')
            ORDER BY mes
            """,
            (*parametros_fecha, *ids_productos)
        )
        valores_producto_mes = {
            (fila["id_producto"], fila["mes"]): float(fila["total"] or 0)
            for fila in evolucion
        }
        max_producto = max(
            (valor for valor in valores_producto_mes.values()), default=0
        )
        colores = ("#318447", "#e69f24", "#4285a4", "#ab5c78", "#7659a6")
        ancho_svg = 800
        alto_svg = 220
        margen_svg = 20
        for indice, producto in enumerate(productos_top):
            valores = [
                valores_producto_mes.get((producto["id_producto"], mes), 0)
                for mes in meses
            ]
            coordenadas = []
            puntos = []
            for indice_mes, valor in enumerate(valores):
                x = (
                    margen_svg + indice_mes * (ancho_svg - margen_svg * 2)
                    / (len(meses) - 1)
                    if len(meses) > 1 else ancho_svg / 2
                )
                y = alto_svg - margen_svg - (
                    valor / max_producto * (alto_svg - margen_svg * 2)
                    if max_producto else 0
                )
                coordenadas.append(f"{x:.1f},{y:.1f}")
                puntos.append({
                    "x": f"{x:.1f}",
                    "y": f"{y:.1f}",
                    "valor": valor,
                    "mes": etiquetas_meses[indice_mes],
                })
            evolucion_productos.append({
                "nombre": producto["nombre"],
                "color": colores[indice % len(colores)],
                "valores": valores,
                "puntos": " ".join(coordenadas),
                "marcadores": puntos,
            })

    return render_template(
        "dashboard.html",
        stats=stats,
        recent=recent,
        incidencias_pendientes=incidencias_pendientes,
        fecha_desde=fecha_desde.isoformat(),
        fecha_hasta=fecha_hasta.isoformat(),
        ventas_mensuales=ventas_mensuales,
        operaciones_por_estado=operaciones_por_estado,
        total_operaciones_estado=total_operaciones_estado,
        circunferencia_estado=circunferencia_estado,
        pagos_por_estado=pagos_por_estado,
        operaciones_mensuales=operaciones_mensuales,
        evolucion_productos=evolucion_productos,
        etiquetas_meses=etiquetas_meses,
        alertas_inventario=alertas_inventario,
    )


# ==========================================================
# PRODUCTOS
# ==========================================================

@app.route("/productos")
def productos():

    products = productos_con_imagen(query(
        """
        SELECT
            p.*,
            c.nombre AS categoria
        FROM productos p

        INNER JOIN categorias_producto c
            ON c.id_categoria = p.id_categoria

        WHERE p.estado = 'ACTIVO'

        ORDER BY p.nombre
        """
    ))
    categories = query(
        """
        SELECT nombre
        FROM categorias_producto
        WHERE estado = 'ACTIVO'
        ORDER BY nombre
        """
    )

    return render_template(
        "productos.html",
        products=products,
        categories=categories
    )


# ==========================================================
# LOTES
# ==========================================================

@app.route("/lotes")
def lotes():

    lots = query(
        """
        SELECT
            l.*,
            p.nombre AS producto,
            c.nombre AS categoria,

            CONCAT(
                u.nombres,
                ' ',
                COALESCE(u.apellido_paterno,'')
            ) AS productor

        FROM lotes l

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN categorias_producto c
            ON c.id_categoria = p.id_categoria

        INNER JOIN productores pr
            ON pr.id_productor = s.id_productor

        INNER JOIN usuarios u
            ON u.id_usuario = pr.id_usuario

        WHERE l.estado = 'DISPONIBLE'

        ORDER BY
            l.fecha_publicacion DESC,
            l.id_lote DESC
        """
    )

    return render_template(
        "lotes.html",
        lots=lots
    )


@app.route("/lotes/<int:id_lote>/reservar", methods=["POST"])
@login_required
def reservar_lote(id_lote):

    comprador = obtener_comprador_actual()
    if not comprador:
        flash(
            "Solo los compradores pueden reservar lotes.",
            "warning"
        )

        return redirect(url_for("lotes"))

    conn = get_connection()
    cur = conn.cursor(dictionary=True)
    try:
        conn.start_transaction()
        cur.execute(
            """
            SELECT estado
            FROM lotes
            WHERE id_lote = %s
            FOR UPDATE
            """,
            (id_lote,)
        )
        lote = cur.fetchone()
        if not lote:
            conn.rollback()
            flash("El lote no existe.", "danger")
            return redirect(url_for("lotes"))

        if lote["estado"] != "DISPONIBLE":
            conn.rollback()
            flash("Este lote ya no está disponible.", "warning")
            return redirect(url_for("lotes"))

        cur.execute(
            """
            SELECT id_cotizacion_operacion, fecha_vencimiento
            FROM cotizaciones_operacion
            WHERE id_lote = %s
              AND id_comprador = %s
              AND estado = 'VIGENTE'
              AND fecha_vencimiento > NOW()
            ORDER BY fecha_calculo DESC
            LIMIT 1
            """,
            (id_lote, comprador["id_comprador"])
        )
        cotizacion = cur.fetchone()
        if not cotizacion:
            conn.rollback()
            flash("Genera primero una cotización vigente para este lote.", "warning")
            return redirect(url_for("lotes"))

        cur.execute(
            """
            INSERT INTO reservas_lote
            (
                id_lote,
                id_comprador,
                id_cotizacion_operacion,
                fecha_vencimiento,
                estado
            )
            VALUES (%s, %s, %s, %s, 'ACTIVA')
            """,
            (
                id_lote,
                comprador["id_comprador"],
                cotizacion["id_cotizacion_operacion"],
                cotizacion["fecha_vencimiento"]
            )
        )
        cur.execute(
            """
            UPDATE lotes
            SET estado = 'RESERVADO'
            WHERE id_lote = %s
              AND estado = 'DISPONIBLE'
            """,
            (id_lote,)
        )
        if cur.rowcount != 1:
            conn.rollback()
            flash("Este lote ya no está disponible.", "warning")
            return redirect(url_for("lotes"))

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()

    flash(
        "Lote reservado correctamente. El administrador puede continuar con la operación.",
        "success"
    )

    return redirect(url_for("lotes"))


# ==========================================================
# REGISTRAR PREDIO
# ==========================================================

@app.route("/predios", methods=["GET", "POST"])
@login_required
def predios():

    productor = query(
        """
        SELECT *
        FROM productores
        WHERE id_usuario = %s
        AND estado = 'ACTIVO'
        """,
        (session["user_id"],),
        one=True
    )

    if not productor:

        flash(
            "Esta opción solo está disponible para productores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    if request.method == "POST":

        nombre = request.form.get("nombre_predio", "").strip()
        referencia = request.form.get("referencia", "").strip()
        unidad = request.form.get("unidad_area", "")
        id_ubigeo = request.form.get("id_ubigeo", type=int)
        try:
            area = Decimal(request.form.get("area_total", ""))
        except InvalidOperation:
            area = Decimal("NaN")

        ubigeo = query(
            """
            SELECT id_ubigeo
            FROM ubigeos
            WHERE id_ubigeo = %s
              AND estado = 'ACTIVO'
            LIMIT 1
            """,
            (id_ubigeo,),
            one=True
        ) if id_ubigeo else None

        if (
            not nombre or len(nombre) > 120
            or len(referencia) > 255
            or not area.is_finite() or area <= 0
            or unidad not in {"HECTAREA", "METRO_CUADRADO"}
            or not ubigeo
        ):
            flash("Completa los datos del predio e indica un ubigeo válido.", "warning")
            return redirect(url_for("predios"))

        codigo = generar_codigo(
            "PRE",
            "predios",
            "codigo_predio"
        )

        execute(
            """
            INSERT INTO predios
            (
                codigo_predio,
                id_productor,
                nombre_predio,
                referencia,
                area_total,
                unidad_area,
                id_ubigeo,
                estado_verificacion,
                estado
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                'PENDIENTE',
                'ACTIVO'
            )
            """,
            (
                codigo,
                productor["id_productor"],
                nombre,
                referencia,
                area,
                unidad,
                id_ubigeo
            )
        )

        flash(
            "Predio registrado correctamente.",
            "success"
        )

        return redirect(url_for("predios"))

    rows = query(
        """
        SELECT p.*, u.departamento, u.provincia, u.distrito
        FROM predios p
        LEFT JOIN ubigeos u ON u.id_ubigeo = p.id_ubigeo
        WHERE p.id_productor = %s
          AND p.estado = 'ACTIVO'
        ORDER BY p.fecha_registro DESC
        """,
        (productor["id_productor"],)
    )

    return render_template(
        "predios.html",
        predios=rows,
        departamentos=query(
            """
            SELECT DISTINCT departamento
            FROM ubigeos
            WHERE estado = 'ACTIVO'
            ORDER BY departamento
            """
        )
    )


@app.route("/api/ubigeos/provincias")
def api_ubigeo_provincias():
    departamento = request.args.get("departamento", "").strip()
    if not departamento:
        return jsonify([])
    return jsonify(query(
        """
        SELECT DISTINCT provincia
        FROM ubigeos
        WHERE departamento = %s
          AND estado = 'ACTIVO'
        ORDER BY provincia
        """,
        (departamento,)
    ))


@app.route("/api/ubigeos/distritos")
def api_ubigeo_distritos():
    departamento = request.args.get("departamento", "").strip()
    provincia = request.args.get("provincia", "").strip()
    if not departamento or not provincia:
        return jsonify([])
    return jsonify(query(
        """
        SELECT id_ubigeo, distrito
        FROM ubigeos
        WHERE departamento = %s
          AND provincia = %s
          AND estado = 'ACTIVO'
        ORDER BY distrito
        """,
        (departamento, provincia)
    ))


# ==========================================================
# SOLICITAR LOTE
# ==========================================================

@app.route("/solicitar-lote", methods=["GET", "POST"])
@login_required
def solicitar_lote():

    productor = query(
        """
        SELECT *
        FROM productores
        WHERE id_usuario = %s
        AND estado = 'ACTIVO'
        """,
        (session["user_id"],),
        one=True
    )

    if not productor:

        flash(
            "Debes estar registrado como productor.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    productos = query(
        """
        SELECT *
        FROM productos
        WHERE estado = 'ACTIVO'
        ORDER BY nombre
        """
    )

    predios = query(
        """
        SELECT *
        FROM predios
        WHERE id_productor = %s
        AND estado = 'ACTIVO'
        ORDER BY nombre_predio
        """,
        (productor["id_productor"],)
    )

    if request.method == "POST":

        producto = request.form["id_producto"]
        predio = request.form["id_predio"]

        area = request.form["area_cultivada"]

        fecha_siembra = request.form["fecha_siembra"] or None

        fecha_inicio = (
            request.form["fecha_cosecha_estimada_inicio"]
            or None
        )

        fecha_fin = (
            request.form["fecha_cosecha_estimada_fin"]
            or None
        )

        descripcion = request.form["descripcion"]

        imagen = request.files.get("imagen")

        codigo = generar_codigo(
            "SOL",
            "solicitudes_lote",
            "codigo_solicitud"
        )

        id_solicitud = execute(
            """
            INSERT INTO solicitudes_lote
            (
                codigo_solicitud,
                id_productor,
                id_predio,
                id_producto,
                area_cultivada,
                fecha_siembra,
                fecha_cosecha_estimada_inicio,
                fecha_cosecha_estimada_fin,
                descripcion,
                estado
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                'ENVIADA'
            )
            """,
            (
                codigo,
                productor["id_productor"],
                predio,
                producto,
                area,
                fecha_siembra,
                fecha_inicio,
                fecha_fin,
                descripcion
            )
        )

        notificar_administradores(
            "Nueva solicitud de lote",
            f"El productor envió la solicitud {codigo} y está pendiente de revisión.",
            "INFO",
            "SOLICITUD_LOTE",
            id_solicitud,
            "SOLICITUDES_REVISAR"
        )
        registrar_historial_estado(
            "SOLICITUD_LOTE",
            id_solicitud,
            None,
            "ENVIADA",
            "Solicitud enviada por el productor."
        )

        try:
            if imagen and imagen.filename:
                guardar_archivo_entidad(
                    imagen,
                    "SOLICITUD_LOTE",
                    id_solicitud,
                    session["user_id"]
                )
        except ValueError as exc:
            flash(str(exc), "warning")
            return redirect(url_for("solicitar_lote"))

        flash(
            "Solicitud de lote enviada correctamente al administrador para su revisión.",
            "success"
        )

        return redirect(url_for("dashboard"))

    return render_template(
        "solicitar_lote.html",
        productos=productos,
        predios=predios
    )


@app.route("/mis-solicitudes")
@login_required
def mis_solicitudes_lote():
    productor = query(
        """
        SELECT id_productor
        FROM productores
        WHERE id_usuario = %s
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        (session["user_id"],),
        one=True
    )
    if not productor:
        flash("El seguimiento de solicitudes está disponible para productores registrados.", "warning")
        return redirect(url_for("dashboard"))

    solicitudes = query(
        """
        SELECT s.id_solicitud, s.codigo_solicitud, s.estado,
               s.fecha_solicitud, s.descripcion, p.nombre AS producto,
               pred.nombre_predio
        FROM solicitudes_lote s
        INNER JOIN productos p ON p.id_producto = s.id_producto
        INNER JOIN predios pred ON pred.id_predio = s.id_predio
        WHERE s.id_productor = %s
        ORDER BY s.fecha_solicitud DESC, s.id_solicitud DESC
        """,
        (productor["id_productor"],)
    )
    if solicitudes:
        ids_solicitud = [item["id_solicitud"] for item in solicitudes]
        placeholders = ", ".join(["%s"] * len(ids_solicitud))
        historial = query(
            f"""
            SELECT entidad_id, estado_anterior, estado_nuevo, fecha, observacion
            FROM historial_estados
            WHERE entidad_tipo = 'SOLICITUD_LOTE'
              AND entidad_id IN ({placeholders})
            ORDER BY fecha ASC, id_historial ASC
            """,
            tuple(ids_solicitud)
        )
        historial_por_solicitud = {}
        for evento in historial:
            historial_por_solicitud.setdefault(evento["entidad_id"], []).append(evento)
        for item in solicitudes:
            item["historial"] = historial_por_solicitud.get(item["id_solicitud"], [])

    return render_template(
        "mis_solicitudes_lote.html",
        solicitudes=solicitudes
    )


# ==========================================================
# ADMIN: SOLICITUDES DE LOTE
# ==========================================================

@app.route("/admin/solicitudes-lote")
@login_required
def admin_solicitudes_lote():

    if not tiene_permiso("SOLICITUDES_REVISAR"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    solicitudes = query(
        """
        SELECT
            s.id_solicitud,
            s.codigo_solicitud,
            s.area_cultivada,
            s.fecha_siembra,
            s.fecha_cosecha_estimada_inicio,
            s.fecha_cosecha_estimada_fin,
            s.descripcion,
            s.estado,
            s.fecha_solicitud,

            (
                SELECT a.ruta_clave
                FROM archivos a
                WHERE a.entidad_tipo = 'SOLICITUD_LOTE'
                  AND a.entidad_id = s.id_solicitud
                  AND a.estado = 'ACTIVO'
                ORDER BY a.fecha_carga DESC
                LIMIT 1
            ) AS imagen,

            CONCAT(
                u.nombres,
                ' ',
                COALESCE(u.apellido_paterno,'')
            ) AS productor,

            p.nombre AS producto,
            pred.nombre_predio

        FROM solicitudes_lote s

        INNER JOIN productores pr
            ON pr.id_productor = s.id_productor

        INNER JOIN usuarios u
            ON u.id_usuario = pr.id_usuario

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN predios pred
            ON pred.id_predio = s.id_predio

        ORDER BY s.fecha_solicitud DESC
        """
    )

    return render_template(
        "admin_solicitudes_lote.html",
        solicitudes=solicitudes
    )


@app.route("/admin/solicitudes-lote/<int:id_solicitud>/aprobar", methods=["POST"])
@login_required
def aprobar_solicitud_lote(id_solicitud):

    if not tiene_permiso("SOLICITUDES_REVISAR"):

        flash(
            "No tienes permisos para aprobar solicitudes.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    solicitud = query(
        """
        SELECT *
        FROM solicitudes_lote
        WHERE id_solicitud = %s
        """,
        (id_solicitud,),
        one=True
    )

    if not solicitud:

        flash(
            "La solicitud no existe.",
            "danger"
        )

        return redirect(url_for("admin_solicitudes_lote"))

    valorizacion_existente = query(
        """
        SELECT id_valorizacion
        FROM valorizaciones
        WHERE id_solicitud = %s
        LIMIT 1
        """,
        (id_solicitud,),
        one=True
    )

    if valorizacion_existente:
        id_valorizacion = valorizacion_existente["id_valorizacion"]
    else:
        codigo_valorizacion = generar_codigo(
            "VAL",
            "valorizaciones",
            "codigo_valorizacion"
        )

        id_valorizacion = execute(
            """
            INSERT INTO valorizaciones
            (
                codigo_valorizacion,
                id_solicitud,
                modo_calculo,
                precio_actual,
                precio_promedio,
                peso_actual,
                peso_historico,
                factor_zona,
                factor_temporada,
                factor_condicion,
                precio_base,
                precio_ajustado,
                produccion_estimada,
                valor_lote,
                id_usuario_respuesta,
                respuesta_productor,
                estado,
                fecha_presentacion,
                fecha_respuesta
            )
            VALUES
            (
                %s,
                %s,
                'AUTOMATICO',
                0,
                0,
                0,
                0,
                1,
                1,
                1,
                0,
                0,
                %s,
                0,
                %s,
                'ACEPTADA',
                'ACEPTADA',
                NOW(),
                NOW()
            )
            """,
            (
                codigo_valorizacion,
                id_solicitud,
                solicitud["area_cultivada"],
                session["user_id"]
            )
        )

    produccion_verificada = query(
        """
        SELECT produccion_min_estimada, produccion_max_estimada
        FROM verificaciones
        WHERE id_solicitud = %s
          AND resultado = 'APROBADO'
          AND estado = 'COMPLETADA'
          AND produccion_min_estimada > 0
          AND produccion_max_estimada >= produccion_min_estimada
        ORDER BY fecha_realizada DESC, id_verificacion DESC
        LIMIT 1
        """,
        (id_solicitud,),
        one=True
    )
    produccion_min_kg = (
        produccion_verificada["produccion_min_estimada"]
        if produccion_verificada else None
    )
    produccion_max_kg = (
        produccion_verificada["produccion_max_estimada"]
        if produccion_verificada else None
    )

    codigo_lote = generar_codigo(
        "LOT",
        "lotes",
        "codigo_lote"
    )

    execute(
        """
        INSERT INTO lotes
        (
            codigo_lote,
            id_solicitud,
            id_valorizacion,
            produccion_min,
            produccion_max,
            precio_publicado,
            fecha_cosecha_inicio,
            fecha_cosecha_fin,
            estado_cultivo_comercial,
            fecha_publicacion,
            estado
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            0,
            %s,
            %s,
            'PROXIMO_COSECHA',
            NOW(),
            'DISPONIBLE'
        )
        """,
        (
            codigo_lote,
            id_solicitud,
            id_valorizacion,
            produccion_min_kg,
            produccion_max_kg,
            solicitud["fecha_cosecha_estimada_inicio"],
            solicitud["fecha_cosecha_estimada_fin"]
        )
    )

    execute(
        """
        UPDATE solicitudes_lote
        SET estado = 'APROBADA',
            fecha_actualizacion = NOW()
        WHERE id_solicitud = %s
        """,
        (id_solicitud,)
    )

    verificacion = query(
        """
        SELECT id_verificacion
        FROM verificaciones
        WHERE id_solicitud = %s
        LIMIT 1
        """,
        (id_solicitud,),
        one=True
    )

    if not verificacion:
        execute(
            """
            INSERT INTO verificaciones
            (
                codigo_verificacion,
                id_solicitud,
                id_verificador,
                fecha_programada,
                estado
            )
            VALUES
            (
                %s,
                %s,
                %s,
                NOW(),
                'PROGRAMADA'
            )
            """,
            (
                generar_codigo("VER", "verificaciones", "codigo_verificacion"),
                id_solicitud,
                session["user_id"]
            )
        )

    calcular_valorizacion_solicitud(id_solicitud)

    registrar_historial_estado(
        "SOLICITUD_LOTE",
        id_solicitud,
        solicitud["estado"],
        "APROBADA",
        "Solicitud aprobada y lote publicado por administración."
    )
    productor_solicitud = query(
        """
        SELECT u.id_usuario
        FROM productores p
        INNER JOIN usuarios u ON u.id_usuario = p.id_usuario
        WHERE p.id_productor = %s
        LIMIT 1
        """,
        (solicitud["id_productor"],),
        one=True
    )
    if productor_solicitud:
        crear_notificacion(
            productor_solicitud["id_usuario"],
            "Solicitud de lote aprobada",
            f"La solicitud {solicitud['codigo_solicitud']} fue aprobada y el lote se publicó.",
            "SUCCESS",
            "SOLICITUD_LOTE",
            id_solicitud
        )

    flash(
        "Solicitud aprobada, lote publicado y flujo de verificación/valorización iniciado.",
        "success"
    )

    return redirect(url_for("admin_solicitudes_lote"))


@app.route("/admin/verificaciones")
@login_required
def admin_verificaciones():

    if not (
        tiene_permiso("VERIFICACIONES_ASIGNAR")
        or tiene_permiso("VERIFICACIONES_EJECUTAR")
    ):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    verificaciones = query(
        """
        SELECT
            v.id_verificacion,
            v.codigo_verificacion,
            v.fecha_programada,
            v.fecha_realizada,
            v.area_observada,
            v.resultado,
            v.estado,
            s.codigo_solicitud,
            s.estado AS estado_solicitud,
            CONCAT(
                u.nombres,
                ' ',
                COALESCE(u.apellido_paterno, '')
            ) AS productor,
            p.nombre AS producto,
            pred.nombre_predio

        FROM verificaciones v

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = v.id_solicitud

        INNER JOIN productores pr
            ON pr.id_productor = s.id_productor

        INNER JOIN usuarios u
            ON u.id_usuario = pr.id_usuario

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN predios pred
            ON pred.id_predio = s.id_predio

        ORDER BY v.fecha_programada DESC
        """
    )

    return render_template(
        "admin_verificaciones.html",
        verificaciones=verificaciones
    )


@app.route("/admin/verificaciones/<int:id_verificacion>/evaluar", methods=["GET", "POST"])
@login_required
def evaluar_verificacion(id_verificacion):

    if not tiene_permiso("VERIFICACIONES_EJECUTAR"):

        flash(
            "No tienes permisos para evaluar verificaciones.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    verificacion = query(
        """
        SELECT
            v.*,
            s.codigo_solicitud,
            s.estado AS estado_solicitud,
            s.area_cultivada,
            u.id_usuario AS id_productor_usuario,
            CONCAT(
                u.nombres,
                ' ',
                COALESCE(u.apellido_paterno, '')
            ) AS productor,
            p.nombre AS producto,
            pred.nombre_predio

        FROM verificaciones v

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = v.id_solicitud

        INNER JOIN productores pr
            ON pr.id_productor = s.id_productor

        INNER JOIN usuarios u
            ON u.id_usuario = pr.id_usuario

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN predios pred
            ON pred.id_predio = s.id_predio

        WHERE v.id_verificacion = %s
        LIMIT 1
        """,
        (id_verificacion,),
        one=True
    )

    if not verificacion:
        flash("La verificación no existe.", "danger")
        return redirect(url_for("admin_verificaciones"))

    if request.method == "POST":
        resultado = request.form.get("resultado", "OBSERVADO")
        estado_cultivo = request.form.get("estado_cultivo", "BUENO")
        etapa_productiva = request.form.get("etapa_productiva", "CULTIVO_ACTIVO")
        area_observada = request.form.get("area_observada")
        produccion_min = request.form.get("produccion_min_estimada")
        produccion_max = request.form.get("produccion_max_estimada")
        observaciones = request.form.get("observaciones", "").strip()

        execute(
            """
            UPDATE verificaciones
            SET fecha_realizada = NOW(),
                area_observada = %s,
                estado_cultivo = %s,
                etapa_productiva = %s,
                produccion_min_estimada = %s,
                produccion_max_estimada = %s,
                observaciones = %s,
                resultado = %s,
                estado = 'COMPLETADA'
            WHERE id_verificacion = %s
            """,
            (
                area_observada or None,
                estado_cultivo,
                etapa_productiva,
                produccion_min or None,
                produccion_max or None,
                observaciones or None,
                resultado,
                id_verificacion
            )
        )

        if resultado == "APROBADO":
            execute(
                """
                UPDATE solicitudes_lote
                SET estado = 'VERIFICADA',
                    fecha_actualizacion = NOW()
                WHERE id_solicitud = %s
                """,
                (verificacion["id_solicitud"],)
            )
            execute(
                """
                UPDATE predios
                SET estado_verificacion = 'VERIFICADO'
                WHERE id_predio = %s
                """,
                (query("SELECT id_predio FROM solicitudes_lote WHERE id_solicitud = %s LIMIT 1", (verificacion["id_solicitud"],), one=True)["id_predio"],)
            )
            calcular_valorizacion_solicitud(verificacion["id_solicitud"])
            flash("La verificación fue aprobada y la valorización fue recalculada.", "success")
        elif resultado == "OBSERVADO":
            execute(
                """
                UPDATE solicitudes_lote
                SET estado = 'OBSERVADA',
                    fecha_actualizacion = NOW()
                WHERE id_solicitud = %s
                """,
                (verificacion["id_solicitud"],)
            )
            execute(
                """
                UPDATE predios
                SET estado_verificacion = 'OBSERVADO'
                WHERE id_predio = %s
                """,
                (query("SELECT id_predio FROM solicitudes_lote WHERE id_solicitud = %s LIMIT 1", (verificacion["id_solicitud"],), one=True)["id_predio"],)
            )
            flash("La verificación quedó observada y requiere revisión.", "warning")
        elif resultado == "RECHAZADO":
            execute(
                """
                UPDATE solicitudes_lote
                SET estado = 'RECHAZADA',
                    fecha_actualizacion = NOW()
                WHERE id_solicitud = %s
                """,
                (verificacion["id_solicitud"],)
            )
            execute(
                """
                UPDATE predios
                SET estado_verificacion = 'RECHAZADO'
                WHERE id_predio = %s
                """,
                (query("SELECT id_predio FROM solicitudes_lote WHERE id_solicitud = %s LIMIT 1", (verificacion["id_solicitud"],), one=True)["id_predio"],)
            )
            flash("La verificación fue rechazada.", "danger")
        else:
            execute(
                """
                UPDATE solicitudes_lote
                SET estado = 'EN_VERIFICACION',
                    fecha_actualizacion = NOW()
                WHERE id_solicitud = %s
                """,
                (verificacion["id_solicitud"],)
            )
            flash("La verificación quedó registrada con resultado pendiente.", "warning")

        estados_solicitud = {
            "APROBADO": "VERIFICADA",
            "OBSERVADO": "OBSERVADA",
            "RECHAZADO": "RECHAZADA",
        }
        nuevo_estado = estados_solicitud.get(resultado, "EN_VERIFICACION")
        registrar_historial_estado(
            "SOLICITUD_LOTE",
            verificacion["id_solicitud"],
            verificacion["estado_solicitud"],
            nuevo_estado,
            observaciones or f"Resultado de verificación: {resultado}."
        )
        crear_notificacion(
            verificacion["id_productor_usuario"],
            "Actualización de verificación de solicitud",
            f"La solicitud {verificacion['codigo_solicitud']} quedó {nuevo_estado}."
            + (f" Observaciones: {observaciones}" if observaciones else ""),
            "SUCCESS" if resultado == "APROBADO" else (
                "DANGER" if resultado == "RECHAZADO" else "WARNING"
            ),
            "SOLICITUD_LOTE",
            verificacion["id_solicitud"]
        )

        return redirect(url_for("admin_verificaciones"))

    return render_template(
        "admin_verificacion_evaluar.html",
        verificacion=verificacion
    )


@app.route("/admin/valorizaciones")
@login_required
def admin_valorizaciones():

    if not tiene_permiso("PRECIOS_VER"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    valorizaciones = query(
        """
        SELECT
            v.id_valorizacion,
            v.codigo_valorizacion,
            v.precio_base,
            v.precio_ajustado,
            v.produccion_estimada,
            v.valor_lote,
            v.modo_calculo,
            v.estado,
            v.fecha_calculo,
            s.codigo_solicitud,
            CONCAT(
                u.nombres,
                ' ',
                COALESCE(u.apellido_paterno, '')
            ) AS productor,
            p.nombre AS producto,
            pred.nombre_predio

        FROM valorizaciones v

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = v.id_solicitud

        INNER JOIN productores pr
            ON pr.id_productor = s.id_productor

        INNER JOIN usuarios u
            ON u.id_usuario = pr.id_usuario

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN predios pred
            ON pred.id_predio = s.id_predio

        ORDER BY v.fecha_calculo DESC
        """
    )

    return render_template(
        "admin_valorizaciones.html",
        valorizaciones=valorizaciones
    )


@app.route("/admin/precios-mercado", methods=["GET", "POST"])
@login_required
def admin_precios_mercado():
    if not tiene_permiso("PRECIOS_VER"):
        flash("No tienes permiso para consultar precios de mercado.", "warning")
        return redirect(url_for("dashboard"))

    if request.method == "POST" and not tiene_permiso("PRECIOS_MODIFICAR"):
        flash("No tienes permiso para modificar precios de mercado.", "danger")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        action = request.form.get("action")
        if action == "cambiar_estado":
            id_precio = request.form.get("id_precio", type=int)
            if id_precio:
                execute(
                    """
                    UPDATE precios_mercado
                    SET estado = IF(estado = 'ACTIVO', 'INACTIVO', 'ACTIVO')
                    WHERE id_precio = %s
                    """,
                    (id_precio,)
                )
                flash("Estado del precio actualizado.", "success")
            return redirect(url_for("admin_precios_mercado"))

        try:
            id_producto = int(request.form.get("id_producto", ""))
            raw_id_variedad = request.form.get("id_variedad", "").strip()
            id_variedad = int(raw_id_variedad) if raw_id_variedad else None
            precio = Decimal(request.form.get("precio", ""))
            fecha_precio = date.fromisoformat(request.form.get("fecha_precio", ""))
        except (ValueError, InvalidOperation):
            flash("Revisa producto, precio, variedad y fecha.", "warning")
            return redirect(url_for("admin_precios_mercado"))

        unidad = request.form.get("unidad", "")
        if not precio.is_finite() or precio <= 0:
            flash("El precio debe ser mayor que cero.", "warning")
            return redirect(url_for("admin_precios_mercado"))
        if unidad not in {"KG", "TONELADA", "SACO", "UNIDAD"}:
            flash("Selecciona una unidad válida.", "warning")
            return redirect(url_for("admin_precios_mercado"))

        producto = query(
            "SELECT id_producto FROM productos WHERE id_producto = %s AND estado = 'ACTIVO'",
            (id_producto,),
            one=True
        )
        if not producto:
            flash("El producto no existe o está inactivo.", "warning")
            return redirect(url_for("admin_precios_mercado"))

        if id_variedad:
            variedad = query(
                """
                SELECT id_variedad
                FROM variedades
                WHERE id_variedad = %s
                  AND id_producto = %s
                  AND estado = 'ACTIVO'
                """,
                (id_variedad, id_producto),
                one=True
            )
            if not variedad:
                flash("La variedad seleccionada no pertenece al producto.", "warning")
                return redirect(url_for("admin_precios_mercado"))

        execute(
            """
            INSERT INTO precios_mercado
            (
                id_producto, id_variedad, precio, unidad, mercado,
                zona, fuente, fecha_precio, id_usuario_registro, estado
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, 'ACTIVO')
            """,
            (
                id_producto,
                id_variedad,
                precio,
                unidad,
                request.form.get("mercado", "").strip() or None,
                request.form.get("zona", "").strip() or None,
                request.form.get("fuente", "").strip() or None,
                fecha_precio,
                session["user_id"]
            )
        )
        flash("Precio de mercado registrado.", "success")
        return redirect(url_for("admin_precios_mercado"))

    precios = query(
        """
        SELECT pm.*, p.nombre AS producto, v.nombre AS variedad
        FROM precios_mercado pm
        INNER JOIN productos p ON p.id_producto = pm.id_producto
        LEFT JOIN variedades v ON v.id_variedad = pm.id_variedad
        ORDER BY pm.fecha_precio DESC, pm.id_precio DESC
        """
    )
    productos = query(
        "SELECT id_producto, nombre FROM productos WHERE estado = 'ACTIVO' ORDER BY nombre"
    )
    variedades = query(
        """
        SELECT id_variedad, id_producto, nombre
        FROM variedades
        WHERE estado = 'ACTIVO'
        ORDER BY nombre
        """
    )
    return render_template(
        "admin_precios_mercado.html",
        precios=precios,
        productos=productos,
        variedades=variedades
    )


@app.route("/admin/variedades", methods=["GET", "POST"])
@login_required
def admin_variedades():
    if not tiene_permiso("PARAMETROS_COMERCIALES_MODIFICAR"):
        flash("Esta sección es solo para administradores.", "warning")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        action = request.form.get("action")
        if action == "cambiar_estado":
            id_variedad = request.form.get("id_variedad", type=int)
            if id_variedad:
                execute(
                    """
                    UPDATE variedades
                    SET estado = IF(estado = 'ACTIVO', 'INACTIVO', 'ACTIVO')
                    WHERE id_variedad = %s
                    """,
                    (id_variedad,)
                )
                flash("Estado de la variedad actualizado.", "success")
            return redirect(url_for("admin_variedades"))

        try:
            id_producto = int(request.form.get("id_producto", ""))
        except ValueError:
            flash("Selecciona un producto válido.", "warning")
            return redirect(url_for("admin_variedades"))

        nombre = request.form.get("nombre", "").strip()
        descripcion = request.form.get("descripcion", "").strip() or None
        if not nombre or len(nombre) > 120 or (descripcion and len(descripcion) > 255):
            flash("Revisa el nombre y la descripción de la variedad.", "warning")
            return redirect(url_for("admin_variedades"))

        producto = query(
            "SELECT id_producto FROM productos WHERE id_producto = %s AND estado = 'ACTIVO'",
            (id_producto,),
            one=True
        )
        if not producto:
            flash("El producto no existe o está inactivo.", "warning")
            return redirect(url_for("admin_variedades"))

        existente = query(
            """
            SELECT id_variedad
            FROM variedades
            WHERE id_producto = %s
              AND nombre = %s
            LIMIT 1
            """,
            (id_producto, nombre),
            one=True
        )
        if existente:
            flash("Esa variedad ya existe para el producto seleccionado.", "warning")
            return redirect(url_for("admin_variedades"))

        execute(
            """
            INSERT INTO variedades (id_producto, nombre, descripcion, estado)
            VALUES (%s, %s, %s, 'ACTIVO')
            """,
            (id_producto, nombre, descripcion)
        )
        flash("Variedad registrada.", "success")
        return redirect(url_for("admin_variedades"))

    variedades = query(
        """
        SELECT v.id_variedad, v.nombre, v.descripcion, v.estado,
               p.id_producto, p.nombre AS producto
        FROM variedades v
        INNER JOIN productos p ON p.id_producto = v.id_producto
        ORDER BY p.nombre, v.nombre
        """
    )
    productos = query(
        "SELECT id_producto, nombre FROM productos WHERE estado = 'ACTIVO' ORDER BY nombre"
    )
    return render_template(
        "admin_variedades.html",
        variedades=variedades,
        productos=productos
    )


@app.route("/admin/parametros-agricolas", methods=["GET", "POST"])
@login_required
def admin_parametros_agricolas():
    if not tiene_permiso("PARAMETROS_COMERCIALES_MODIFICAR"):
        flash("No tienes permiso para administrar parámetros agrícolas.", "warning")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        action = request.form.get("action")
        if action == "actualizar_comercial":
            id_parametro = request.form.get("id_parametro", type=int)
            parametro = query(
                "SELECT tipo_dato FROM parametros_comerciales WHERE id_parametro = %s",
                (id_parametro,),
                one=True
            )
            try:
                valor = Decimal(request.form.get("valor", ""))
            except InvalidOperation:
                valor = Decimal("NaN")
            if not parametro or not valor.is_finite() or valor < 0:
                flash("El valor del parámetro comercial no es válido.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            if parametro["tipo_dato"] == "PORCENTAJE" and valor > 100:
                flash("Los porcentajes deben estar entre 0 y 100.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            if parametro["tipo_dato"] == "FACTOR" and valor <= 0:
                flash("El factor debe ser mayor que cero.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            execute(
                """
                UPDATE parametros_comerciales
                SET valor = %s, id_usuario_actualiza = %s, fecha_actualizacion = NOW()
                WHERE id_parametro = %s
                """,
                (valor, session["user_id"], id_parametro)
            )
            flash("Parámetro comercial actualizado.", "success")
            return redirect(url_for("admin_parametros_agricolas"))

        if action == "actualizar_cultivo":
            id_parametro = request.form.get("id_parametro_cultivo", type=int)
            try:
                dias_min = int(request.form.get("dias_cosecha_min", ""))
                dias_max = int(request.form.get("dias_cosecha_max", ""))
                raw_rend_min = request.form.get("rendimiento_min_kg_ha", "").strip()
                raw_rend_max = request.form.get("rendimiento_max_kg_ha", "").strip()
                rend_min = Decimal(raw_rend_min) if raw_rend_min else None
                rend_max = Decimal(raw_rend_max) if raw_rend_max else None
            except (ValueError, InvalidOperation):
                flash("Revisa los días y rendimientos del cultivo.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            if dias_min < 1 or dias_max < dias_min:
                flash("El rango de días de cosecha no es válido.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            if any(valor is not None and (not valor.is_finite() or valor < 0) for valor in (rend_min, rend_max)):
                flash("Los rendimientos deben ser valores válidos no negativos.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            if rend_min is not None and rend_max is not None and rend_max < rend_min:
                flash("El rendimiento máximo no puede ser menor que el mínimo.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            execute(
                """
                UPDATE parametros_cultivo
                SET dias_cosecha_min = %s, dias_cosecha_max = %s,
                    rendimiento_min_kg_ha = %s, rendimiento_max_kg_ha = %s,
                    id_usuario_actualiza = %s, fecha_actualizacion = NOW()
                WHERE id_parametro_cultivo = %s
                """,
                (dias_min, dias_max, rend_min, rend_max, session["user_id"], id_parametro)
            )
            flash("Parámetros de cultivo actualizados.", "success")
            return redirect(url_for("admin_parametros_agricolas"))

        if action == "actualizar_cosecha":
            id_parametro = request.form.get("id_parametro_cosecha", type=int)
            nombres_decimal = (
                "jornales_referenciales_ha", "jornal_diario", "costo_carga_tonelada",
                "costo_acondicionamiento_tonelada", "otros_costos_base"
            )
            try:
                valores = {
                    nombre: Decimal(request.form.get(nombre, ""))
                    if request.form.get(nombre, "").strip() else None
                    for nombre in nombres_decimal
                }
                duracion = int(request.form.get("duracion_objetivo_dias", ""))
            except (ValueError, InvalidOperation):
                flash("Revisa los valores de los parámetros de cosecha.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            obligatorios = ("jornal_diario", "costo_carga_tonelada", "costo_acondicionamiento_tonelada", "otros_costos_base")
            if duracion < 1 or duracion > 255 or any(valores[nombre] is None for nombre in obligatorios):
                flash("La duración y los costos obligatorios deben ser válidos.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            if any(valor is not None and (not valor.is_finite() or valor < 0) for valor in valores.values()):
                flash("Los parámetros de cosecha deben ser valores no negativos.", "warning")
                return redirect(url_for("admin_parametros_agricolas"))
            execute(
                """
                UPDATE parametros_cosecha
                SET jornales_referenciales_ha = %s, jornal_diario = %s,
                    costo_carga_tonelada = %s, costo_acondicionamiento_tonelada = %s,
                    otros_costos_base = %s, duracion_objetivo_dias = %s,
                    id_usuario_actualiza = %s, fecha_actualizacion = NOW()
                WHERE id_parametro_cosecha = %s
                """,
                (
                    valores["jornales_referenciales_ha"], valores["jornal_diario"],
                    valores["costo_carga_tonelada"], valores["costo_acondicionamiento_tonelada"],
                    valores["otros_costos_base"], duracion, session["user_id"], id_parametro
                )
            )
            flash("Parámetros de cosecha actualizados.", "success")
            return redirect(url_for("admin_parametros_agricolas"))

        if action == "cambiar_estado":
            tablas = {
                "cultivo": ("parametros_cultivo", "id_parametro_cultivo"),
                "cosecha": ("parametros_cosecha", "id_parametro_cosecha")
            }
            tabla = tablas.get(request.form.get("tipo"))
            id_registro = request.form.get("id_registro", type=int)
            if tabla and id_registro:
                execute(
                    f"UPDATE {tabla[0]} SET estado = IF(estado = 'ACTIVO', 'INACTIVO', 'ACTIVO'), id_usuario_actualiza = %s WHERE {tabla[1]} = %s",
                    (session["user_id"], id_registro)
                )
                flash("Estado del parámetro actualizado.", "success")
            return redirect(url_for("admin_parametros_agricolas"))

        flash("La acción solicitada no está disponible.", "warning")
        return redirect(url_for("admin_parametros_agricolas"))

    comerciales = query(
        "SELECT * FROM parametros_comerciales ORDER BY nombre"
    )
    cultivos = query(
        """
        SELECT pc.*, p.nombre AS producto, v.nombre AS variedad
        FROM parametros_cultivo pc
        INNER JOIN productos p ON p.id_producto = pc.id_producto
        LEFT JOIN variedades v ON v.id_variedad = pc.id_variedad
        ORDER BY p.nombre, v.nombre
        """
    )
    cosechas = query(
        """
        SELECT pc.*, p.nombre AS producto, v.nombre AS variedad
        FROM parametros_cosecha pc
        INNER JOIN productos p ON p.id_producto = pc.id_producto
        LEFT JOIN variedades v ON v.id_variedad = pc.id_variedad
        ORDER BY p.nombre, v.nombre
        """
    )
    return render_template(
        "admin_parametros_agricolas.html",
        comerciales=comerciales,
        cultivos=cultivos,
        cosechas=cosechas
    )


@app.route("/admin/configuracion-logistica", methods=["GET", "POST"])
@login_required
def admin_configuracion_logistica():
    if not tiene_permiso("TRANSPORTE_ASIGNAR"):
        flash("No tienes permiso para administrar la configuración logística.", "warning")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        action = request.form.get("action")
        if action == "actualizar_parametro":
            id_parametro = request.form.get("id_parametro", type=int)
            try:
                valor = Decimal(request.form.get("valor", ""))
            except InvalidOperation:
                valor = Decimal("NaN")
            if not id_parametro or not valor.is_finite() or valor < 0:
                flash("El valor del parámetro logístico no es válido.", "warning")
                return redirect(url_for("admin_configuracion_logistica"))
            execute(
                """
                UPDATE parametros_logisticos
                SET valor = %s, id_usuario_actualiza = %s, fecha_actualizacion = NOW()
                WHERE id_parametro = %s
                """,
                (valor, session["user_id"], id_parametro)
            )
            flash("Parámetro logístico actualizado.", "success")
            return redirect(url_for("admin_configuracion_logistica"))

        if action == "actualizar_tipo_vehiculo":
            id_tipo = request.form.get("id_tipo_vehiculo", type=int)
            nombres_decimal = (
                "capacidad_util_kg", "rendimiento_km_galon",
                "costo_operativo_km", "costo_fijo_viaje"
            )
            try:
                valores = {
                    nombre: Decimal(request.form.get(nombre, ""))
                    for nombre in nombres_decimal
                }
            except InvalidOperation:
                flash("Los valores del tipo de vehículo no son válidos.", "warning")
                return redirect(url_for("admin_configuracion_logistica"))
            calidad = request.form.get("calidad_dato", "")
            calidades = {"LEGAL_REFERENCIAL", "REFERENCIAL", "ESTIMADO_MVP"}
            if (
                not id_tipo
                or not all(valor.is_finite() and valor >= 0 for valor in valores.values())
                or valores["capacidad_util_kg"] <= 0
                or valores["rendimiento_km_galon"] <= 0
                or calidad not in calidades
            ):
                flash("Capacidad, rendimiento, costos o calidad no válidos.", "warning")
                return redirect(url_for("admin_configuracion_logistica"))
            execute(
                """
                UPDATE tipos_vehiculo_logistica
                SET capacidad_util_kg = %s, rendimiento_km_galon = %s,
                    costo_operativo_km = %s, costo_fijo_viaje = %s,
                    calidad_dato = %s, actualizado_en = NOW()
                WHERE id_tipo_vehiculo = %s
                """,
                (
                    valores["capacidad_util_kg"], valores["rendimiento_km_galon"],
                    valores["costo_operativo_km"], valores["costo_fijo_viaje"],
                    calidad, id_tipo
                )
            )
            flash("Tipo de vehículo actualizado.", "success")
            return redirect(url_for("admin_configuracion_logistica"))

        if action == "registrar_combustible":
            try:
                precio_galon = Decimal(request.form.get("precio_galon", ""))
                fecha_precio = date.fromisoformat(request.form.get("fecha_precio", ""))
            except (ValueError, InvalidOperation):
                flash("Revisa el precio y la fecha del combustible.", "warning")
                return redirect(url_for("admin_configuracion_logistica"))
            combustible = request.form.get("combustible", "").strip()
            fuente = request.form.get("fuente", "").strip()
            calidad = request.form.get("calidad_dato", "")
            if (
                not combustible or len(combustible) > 80
                or not fuente or len(fuente) > 255
                or not precio_galon.is_finite() or precio_galon <= 0
                or calidad not in {"OFICIAL", "REFERENCIAL", "ESTIMADO_MVP"}
            ):
                flash("Completa combustible, precio, fuente y calidad válidos.", "warning")
                return redirect(url_for("admin_configuracion_logistica"))
            execute(
                """
                INSERT INTO precios_combustible_logistica
                (
                    combustible, departamento, provincia, distrito,
                    precio_galon, fuente, fecha_precio, calidad_dato, estado
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 'ACTIVO')
                """,
                (
                    combustible,
                    request.form.get("departamento", "").strip() or None,
                    request.form.get("provincia", "").strip() or None,
                    request.form.get("distrito", "").strip() or None,
                    precio_galon,
                    fuente,
                    fecha_precio,
                    calidad
                )
            )
            flash("Precio de combustible registrado.", "success")
            return redirect(url_for("admin_configuracion_logistica"))

        if action == "cambiar_estado":
            tablas = {
                "tipo": ("tipos_vehiculo_logistica", "id_tipo_vehiculo"),
                "combustible": ("precios_combustible_logistica", "id_precio_combustible")
            }
            tabla = tablas.get(request.form.get("tipo"))
            id_registro = request.form.get("id_registro", type=int)
            if tabla and id_registro:
                execute(
                    f"UPDATE {tabla[0]} SET estado = IF(estado = 'ACTIVO', 'INACTIVO', 'ACTIVO') WHERE {tabla[1]} = %s",
                    (id_registro,)
                )
                flash("Estado logístico actualizado.", "success")
            return redirect(url_for("admin_configuracion_logistica"))

        flash("La acción solicitada no está disponible.", "warning")
        return redirect(url_for("admin_configuracion_logistica"))

    parametros = query("SELECT * FROM parametros_logisticos ORDER BY nombre")
    tipos = query("SELECT * FROM tipos_vehiculo_logistica ORDER BY capacidad_util_kg")
    combustibles = query(
        "SELECT * FROM precios_combustible_logistica ORDER BY fecha_precio DESC, id_precio_combustible DESC"
    )
    tipos_referencia = query(
        "SELECT * FROM vw_vehiculos_logistica ORDER BY capacidad_util_kg"
    )
    combustibles_actuales = query(
        """
        SELECT *
        FROM vw_combustible_logistica_actual
        ORDER BY departamento, provincia, distrito, combustible
        """
    )
    return render_template(
        "admin_configuracion_logistica.html",
        parametros=parametros,
        tipos=tipos,
        combustibles=combustibles,
        tipos_referencia=tipos_referencia,
        combustibles_actuales=combustibles_actuales
    )


@app.route("/admin/beneficios", methods=["GET", "POST"])
@login_required
def admin_beneficios():
    if not es_admin():
        flash("Esta sección es solo para administradores.", "warning")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        action = request.form.get("action")
        if action == "actualizar_stock":
            id_beneficio = request.form.get("id_beneficio", type=int)
            raw_stock = request.form.get("stock", "").strip()
            try:
                stock = int(raw_stock) if raw_stock else None
            except ValueError:
                flash("El stock debe ser un entero no negativo.", "warning")
                return redirect(url_for("admin_beneficios"))
            if stock is not None and stock < 0:
                flash("El stock no puede ser negativo.", "warning")
                return redirect(url_for("admin_beneficios"))
            execute(
                """
                UPDATE beneficios
                SET stock = %s,
                    estado = CASE WHEN %s = 0 THEN 'AGOTADO' ELSE 'ACTIVO' END
                WHERE id_beneficio = %s
                """,
                (stock, stock, id_beneficio)
            )
            flash("Stock del beneficio actualizado.", "success")
            return redirect(url_for("admin_beneficios"))

        if action == "cambiar_estado":
            id_beneficio = request.form.get("id_beneficio", type=int)
            execute(
                """
                UPDATE beneficios
                SET estado = IF(estado = 'ACTIVO', 'INACTIVO', 'ACTIVO')
                WHERE id_beneficio = %s
                  AND estado IN ('ACTIVO', 'INACTIVO')
                """,
                (id_beneficio,)
            )
            flash("Estado del beneficio actualizado.", "success")
            return redirect(url_for("admin_beneficios"))

        nombre = request.form.get("nombre", "").strip()
        descripcion = request.form.get("descripcion", "").strip() or None
        tipo_usuario = request.form.get("tipo_usuario", "")
        raw_stock = request.form.get("stock", "").strip()
        raw_inicio = request.form.get("vigencia_inicio", "").strip()
        raw_fin = request.form.get("vigencia_fin", "").strip()
        try:
            puntos = int(request.form.get("puntos_requeridos", ""))
            stock = int(raw_stock) if raw_stock else None
            vigencia_inicio = date.fromisoformat(raw_inicio) if raw_inicio else None
            vigencia_fin = date.fromisoformat(raw_fin) if raw_fin else None
        except ValueError:
            flash("Revisa puntos, stock y fechas de vigencia.", "warning")
            return redirect(url_for("admin_beneficios"))

        if (
            not nombre or len(nombre) > 150 or puntos <= 0
            or (stock is not None and stock < 0)
            or tipo_usuario not in {"PRODUCTOR", "COMPRADOR", "AMBOS"}
            or (vigencia_inicio and vigencia_fin and vigencia_fin < vigencia_inicio)
        ):
            flash("Los datos del beneficio no son válidos.", "warning")
            return redirect(url_for("admin_beneficios"))

        execute(
            """
            INSERT INTO beneficios
            (
                codigo_beneficio, nombre, descripcion, puntos_requeridos,
                tipo_usuario, vigencia_inicio, vigencia_fin, stock, estado
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                generar_codigo("BEN", "beneficios", "codigo_beneficio"),
                nombre,
                descripcion,
                puntos,
                tipo_usuario,
                vigencia_inicio,
                vigencia_fin,
                stock,
                "AGOTADO" if stock == 0 else "ACTIVO"
            )
        )
        flash("Beneficio registrado.", "success")
        return redirect(url_for("admin_beneficios"))

    beneficios = query("SELECT * FROM beneficios ORDER BY nombre")
    return render_template("admin_beneficios.html", beneficios=beneficios)


@app.route("/admin/campanas-agricolas", methods=["GET", "POST"])
@login_required
def admin_campanas_agricolas():
    if not tiene_permiso("SOLICITUDES_REVISAR"):
        flash("No tienes permiso para administrar campañas agrícolas.", "danger")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        try:
            id_campana = request.form.get("id_campana", type=int)
            fecha_inicio = date.fromisoformat(request.form.get("fecha_inicio", ""))
            fecha_fin = date.fromisoformat(request.form.get("fecha_fin", ""))
        except ValueError:
            flash("Revisa las fechas de la campaña.", "warning")
            return redirect(url_for("admin_campanas_agricolas"))

        nombre = request.form.get("nombre", "").strip()
        estado = request.form.get("estado", "PLANIFICADA")
        if (
            not nombre or len(nombre) > 100 or fecha_fin < fecha_inicio
            or estado not in {"PLANIFICADA", "ACTIVA", "CERRADA", "INACTIVA"}
        ):
            flash("Los datos de la campaña no son válidos.", "warning")
            return redirect(url_for("admin_campanas_agricolas"))

        duplicada = query(
            "SELECT id_campana FROM campanas_agricolas WHERE nombre = %s AND id_campana <> %s LIMIT 1",
            (nombre, id_campana or 0),
            one=True
        )
        if duplicada:
            flash("Ya existe una campaña con ese nombre.", "warning")
            return redirect(url_for("admin_campanas_agricolas"))

        if id_campana:
            execute(
                """
                UPDATE campanas_agricolas
                SET nombre = %s, fecha_inicio = %s, fecha_fin = %s, estado = %s
                WHERE id_campana = %s
                """,
                (nombre, fecha_inicio, fecha_fin, estado, id_campana)
            )
            flash("Campaña actualizada.", "success")
        else:
            execute(
                """
                INSERT INTO campanas_agricolas (nombre, fecha_inicio, fecha_fin, estado)
                VALUES (%s, %s, %s, %s)
                """,
                (nombre, fecha_inicio, fecha_fin, estado)
            )
            flash("Campaña creada.", "success")
        return redirect(url_for("admin_campanas_agricolas"))

    campanas = query("SELECT * FROM campanas_agricolas ORDER BY fecha_inicio DESC, nombre")
    return render_template("admin_campanas_agricolas.html", campanas=campanas)


@app.route("/admin/publicaciones", methods=["GET", "POST"])
@login_required
def admin_publicaciones():
    if not tiene_permiso("PUBLICACIONES_GESTIONAR"):
        flash("No tienes permiso para gestionar publicaciones.", "danger")
        return redirect(url_for("dashboard"))

    tipos_publicacion = {
        "NOTICIA", "CAPACITACION", "EVENTO",
        "INFORMACION_AGRICOLA", "ACTIVIDAD_AGROVIA"
    }
    estados_publicacion = {"BORRADOR", "PUBLICADA", "ARCHIVADA"}
    if request.method == "POST":
        id_publicacion = request.form.get("id_publicacion", type=int)
        tipo = request.form.get("tipo", "")
        titulo = request.form.get("titulo", "").strip()
        contenido = request.form.get("contenido", "").strip()
        estado = request.form.get("estado", "BORRADOR")
        if (
            tipo not in tipos_publicacion or estado not in estados_publicacion
            or not titulo or len(titulo) > 180 or not contenido
        ):
            flash("Revisa el tipo, título, contenido y estado de la publicación.", "warning")
            return redirect(url_for("admin_publicaciones"))

        if id_publicacion:
            execute(
                """
                UPDATE publicaciones
                SET tipo = %s, titulo = %s, contenido = %s, estado = %s,
                    fecha_publicacion = CASE
                        WHEN %s = 'PUBLICADA' THEN COALESCE(fecha_publicacion, NOW())
                        ELSE fecha_publicacion
                    END
                WHERE id_publicacion = %s
                """,
                (tipo, titulo, contenido, estado, estado, id_publicacion)
            )
            flash("Publicación actualizada.", "success")
        else:
            execute(
                """
                INSERT INTO publicaciones
                (codigo_publicacion, tipo, titulo, contenido, id_autor, fecha_publicacion, estado)
                VALUES (%s, %s, %s, %s, %s, IF(%s = 'PUBLICADA', NOW(), NULL), %s)
                """,
                (
                    generar_codigo("PUB", "publicaciones", "codigo_publicacion"),
                    tipo,
                    titulo,
                    contenido,
                    session["user_id"],
                    estado,
                    estado
                )
            )
            flash("Publicación creada.", "success")
        return redirect(url_for("admin_publicaciones"))

    publicaciones = query(
        """
        SELECT p.*, CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, '')) AS autor
        FROM publicaciones p
        INNER JOIN usuarios u ON u.id_usuario = p.id_autor
        ORDER BY COALESCE(p.fecha_publicacion, '1000-01-01') DESC, p.id_publicacion DESC
        """
    )
    return render_template(
        "admin_publicaciones.html",
        publicaciones=publicaciones,
        tipos_publicacion=sorted(tipos_publicacion),
        estados_publicacion=sorted(estados_publicacion)
    )


@app.route("/admin/referidos", methods=["GET", "POST"])
@login_required
def admin_referidos():
    if not es_admin():
        flash("Esta sección es solo para administradores.", "warning")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        action = request.form.get("action")
        id_referido = request.form.get("id_referido", type=int)
        if action == "validar":
            execute(
                """
                UPDATE referidos
                SET estado = 'VALIDADO'
                WHERE id_referido = %s
                  AND estado = 'REGISTRADO'
                """,
                (id_referido,)
            )
            flash("Referido validado.", "success")
            return redirect(url_for("admin_referidos"))

        if action == "bonificar":
            try:
                puntos = int(request.form.get("puntos", ""))
            except ValueError:
                puntos = 0
            if not id_referido or puntos <= 0:
                flash("Indica una cantidad de puntos válida.", "warning")
                return redirect(url_for("admin_referidos"))

            conn = get_connection()
            cur = conn.cursor(dictionary=True)
            try:
                conn.start_transaction()
                cur.execute(
                    "SELECT * FROM referidos WHERE id_referido = %s FOR UPDATE",
                    (id_referido,)
                )
                referido = cur.fetchone()
                if not referido or referido["estado"] != "VALIDADO":
                    conn.rollback()
                    flash("Solo se pueden bonificar referidos validados.", "warning")
                    return redirect(url_for("admin_referidos"))

                cur.execute(
                    "UPDATE referidos SET estado = 'BONIFICADO' WHERE id_referido = %s AND estado = 'VALIDADO'",
                    (id_referido,)
                )
                if cur.rowcount != 1:
                    conn.rollback()
                    flash("El referido ya fue procesado.", "warning")
                    return redirect(url_for("admin_referidos"))
                cur.execute(
                    """
                    INSERT INTO movimientos_puntos (id_usuario, tipo, puntos, motivo)
                    VALUES (%s, 'CREDITO', %s, %s)
                    """,
                    (
                        referido["id_usuario_referidor"],
                        puntos,
                        f"Bono de referido: {id_referido}"
                    )
                )
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                cur.close()
                conn.close()
            flash("Puntos de referido acreditados.", "success")
            return redirect(url_for("admin_referidos"))

        flash("La acción solicitada no está disponible.", "warning")
        return redirect(url_for("admin_referidos"))

    referidos = query(
        """
        SELECT r.id_referido, r.codigo_utilizado, r.estado, r.fecha,
               r.id_usuario_referidor,
               CONCAT(ur.nombres, ' ', COALESCE(ur.apellido_paterno, '')) AS referidor,
               CONCAT(ue.nombres, ' ', COALESCE(ue.apellido_paterno, '')) AS referido
        FROM referidos r
        INNER JOIN usuarios ur ON ur.id_usuario = r.id_usuario_referidor
        INNER JOIN usuarios ue ON ue.id_usuario = r.id_usuario_referido
        ORDER BY r.fecha DESC
        """
    )
    return render_template("admin_referidos.html", referidos=referidos)


@app.route("/admin/permisos", methods=["GET", "POST"])
@login_required
def admin_permisos():
    if not tiene_permiso("ROLES_GESTIONAR"):
        flash("No tienes permiso para gestionar roles y permisos.", "danger")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        try:
            id_rol = int(request.form.get("id_rol", ""))
            ids_permiso = {int(value) for value in request.form.getlist("permisos")}
        except ValueError:
            flash("La selección de rol o permisos no es válida.", "warning")
            return redirect(url_for("admin_permisos"))

        rol = query(
            "SELECT id_rol FROM roles WHERE id_rol = %s AND estado = 'ACTIVO'",
            (id_rol,),
            one=True
        )
        permisos_activos = query(
            "SELECT id_permiso FROM permisos WHERE estado = 'ACTIVO'"
        )
        ids_validos = {permiso["id_permiso"] for permiso in permisos_activos}
        if not rol or not ids_permiso.issubset(ids_validos):
            flash("El rol o alguno de los permisos seleccionados no está activo.", "warning")
            return redirect(url_for("admin_permisos"))

        conn = get_connection()
        cur = conn.cursor()
        try:
            conn.start_transaction()
            cur.execute(
                "UPDATE roles_permisos SET estado = 'INACTIVO' WHERE id_rol = %s AND estado = 'ACTIVO'",
                (id_rol,)
            )
            for id_permiso in ids_permiso:
                cur.execute(
                    """
                    INSERT INTO roles_permisos (id_rol, id_permiso, estado)
                    VALUES (%s, %s, 'ACTIVO')
                    ON DUPLICATE KEY UPDATE estado = 'ACTIVO'
                    """,
                    (id_rol, id_permiso)
                )
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()
        flash("Permisos del rol actualizados.", "success")
        return redirect(url_for("admin_permisos"))

    roles = query("SELECT id_rol, nombre, descripcion FROM roles WHERE estado = 'ACTIVO' ORDER BY nombre")
    permisos = query("SELECT id_permiso, codigo, nombre, descripcion, modulo FROM permisos WHERE estado = 'ACTIVO' ORDER BY modulo, nombre")
    asignaciones = query("SELECT id_rol, id_permiso FROM roles_permisos WHERE estado = 'ACTIVO'")
    permisos_por_rol = {}
    for asignacion in asignaciones:
        permisos_por_rol.setdefault(asignacion["id_rol"], set()).add(asignacion["id_permiso"])
    return render_template(
        "admin_permisos.html",
        roles=roles,
        permisos=permisos,
        permisos_por_rol=permisos_por_rol
    )


@app.route("/beneficios", methods=["GET", "POST"])
@login_required
def beneficios_usuario():
    roles = session.get("roles", [])
    tipos_usuario = []
    if "PRODUCTOR" in roles:
        tipos_usuario.append("PRODUCTOR")
    if "COMPRADOR" in roles:
        tipos_usuario.append("COMPRADOR")
    if not tipos_usuario:
        flash("Los puntos y beneficios están disponibles para productores y compradores.", "warning")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        try:
            id_beneficio = int(request.form.get("id_beneficio", ""))
        except ValueError:
            flash("Selecciona un beneficio válido.", "warning")
            return redirect(url_for("beneficios_usuario"))

        conn = get_connection()
        cur = conn.cursor(dictionary=True)
        try:
            conn.start_transaction()
            cur.execute(
                "SELECT id_usuario FROM usuarios WHERE id_usuario = %s FOR UPDATE",
                (session["user_id"],)
            )
            if not cur.fetchone():
                conn.rollback()
                flash("No se encontró tu cuenta.", "danger")
                return redirect(url_for("beneficios_usuario"))

            role_placeholders = ", ".join(["%s"] * len(tipos_usuario))
            cur.execute(
                f"""
                SELECT *
                FROM beneficios
                WHERE id_beneficio = %s
                  AND estado = 'ACTIVO'
                  AND (vigencia_inicio IS NULL OR vigencia_inicio <= CURDATE())
                  AND (vigencia_fin IS NULL OR vigencia_fin >= CURDATE())
                  AND (tipo_usuario = 'AMBOS' OR tipo_usuario IN ({role_placeholders}))
                FOR UPDATE
                """,
                (id_beneficio, *tipos_usuario)
            )
            beneficio = cur.fetchone()
            if not beneficio or beneficio["stock"] == 0:
                conn.rollback()
                flash("El beneficio ya no está disponible para tu cuenta.", "warning")
                return redirect(url_for("beneficios_usuario"))

            cur.execute(
                """
                SELECT COALESCE(SUM(CASE WHEN tipo = 'CREDITO' THEN puntos ELSE -puntos END), 0) AS saldo
                FROM movimientos_puntos
                WHERE id_usuario = %s
                """,
                (session["user_id"],)
            )
            saldo = int(cur.fetchone()["saldo"] or 0)
            if saldo < beneficio["puntos_requeridos"]:
                conn.rollback()
                flash("No tienes puntos suficientes para este canje.", "warning")
                return redirect(url_for("beneficios_usuario"))

            cur.execute(
                """
                INSERT INTO canjes_beneficios
                (id_beneficio, id_usuario, puntos_utilizados, estado)
                VALUES (%s, %s, %s, 'SOLICITADO')
                """,
                (id_beneficio, session["user_id"], beneficio["puntos_requeridos"])
            )
            cur.execute(
                """
                INSERT INTO movimientos_puntos (id_usuario, tipo, puntos, motivo)
                VALUES (%s, 'DEBITO', %s, %s)
                """,
                (
                    session["user_id"],
                    beneficio["puntos_requeridos"],
                    f"Canje de beneficio: {beneficio['nombre']}"
                )
            )
            if beneficio["stock"] is not None:
                cur.execute(
                    """
                    UPDATE beneficios
                    SET stock = stock - 1,
                        estado = IF(stock = 1, 'AGOTADO', estado)
                    WHERE id_beneficio = %s AND stock > 0
                    """,
                    (id_beneficio,)
                )
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()
        flash("Canje solicitado. Los puntos se descontaron de tu saldo.", "success")
        return redirect(url_for("beneficios_usuario"))

    placeholders = ", ".join(["%s"] * len(tipos_usuario))
    beneficios = query(
        f"""
        SELECT *
        FROM beneficios
        WHERE estado = 'ACTIVO'
          AND (vigencia_inicio IS NULL OR vigencia_inicio <= CURDATE())
          AND (vigencia_fin IS NULL OR vigencia_fin >= CURDATE())
          AND (tipo_usuario = 'AMBOS' OR tipo_usuario IN ({placeholders}))
        ORDER BY puntos_requeridos, nombre
        """,
        tuple(tipos_usuario)
    )
    movimientos = query(
        """
        SELECT tipo, puntos, motivo, fecha
        FROM movimientos_puntos
        WHERE id_usuario = %s
        ORDER BY fecha DESC, id_movimiento DESC
        LIMIT 50
        """,
        (session["user_id"],)
    )
    saldo_actual = query(
        """
        SELECT COALESCE(SUM(CASE WHEN tipo = 'CREDITO' THEN puntos ELSE -puntos END), 0) AS saldo
        FROM movimientos_puntos
        WHERE id_usuario = %s
        """,
        (session["user_id"],),
        one=True
    )
    canjes = query(
        """
        SELECT cb.puntos_utilizados, cb.fecha, cb.estado, b.nombre AS beneficio
        FROM canjes_beneficios cb
        INNER JOIN beneficios b ON b.id_beneficio = cb.id_beneficio
        WHERE cb.id_usuario = %s
        ORDER BY cb.fecha DESC, cb.id_canje DESC
        LIMIT 50
        """,
        (session["user_id"],)
    )
    return render_template(
        "beneficios.html",
        saldo=int((saldo_actual or {}).get("saldo") or 0),
        movimientos=movimientos,
        beneficios=beneficios,
        canjes=canjes
    )


@app.route("/mis-referidos")
@login_required
def mis_referidos():
    usuario = query(
        "SELECT codigo_usuario FROM usuarios WHERE id_usuario = %s",
        (session["user_id"],),
        one=True
    )
    referidos = query(
        """
        SELECT r.estado, r.fecha,
               CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, '')) AS nombre
        FROM referidos r
        INNER JOIN usuarios u ON u.id_usuario = r.id_usuario_referido
        WHERE r.id_usuario_referidor = %s
        ORDER BY r.fecha DESC
        """,
        (session["user_id"],)
    )
    bonos = query(
        """
        SELECT COALESCE(SUM(puntos), 0) AS puntos
        FROM movimientos_puntos
        WHERE id_usuario = %s
          AND tipo = 'CREDITO'
          AND motivo LIKE 'Bono de referido:%'
        """,
        (session["user_id"],),
        one=True
    )
    return render_template(
        "referidos.html",
        codigo_referido=f"AGROVIA-{usuario['codigo_usuario']}",
        referidos=referidos,
        puntos_referidos=int((bonos or {}).get("puntos") or 0)
    )


@app.route("/admin/cotizaciones")
@login_required
def admin_cotizaciones():

    if not tiene_permiso("TRANSPORTE_ASIGNAR"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    cotizaciones = query(
        """
        SELECT
            c.id_cotizacion_operacion,
            c.codigo_cotizacion,
            c.valor_produccion,
            c.costo_base_flete,
            c.precio_flete_cliente,
            c.total_operacion,
            c.estado,
            c.fecha_calculo,
            c.fecha_vencimiento,
            (
                SELECT cl.distancia_km
                FROM cotizaciones_logisticas cl
                WHERE cl.id_cotizacion_operacion = c.id_cotizacion_operacion
                ORDER BY cl.id_cotizacion DESC
                LIMIT 1
            ) AS distancia_km,
            (
                SELECT cl.peso_kg
                FROM cotizaciones_logisticas cl
                WHERE cl.id_cotizacion_operacion = c.id_cotizacion_operacion
                ORDER BY cl.id_cotizacion DESC
                LIMIT 1
            ) AS peso_logistico_kg,
            (
                SELECT cl.tipo_vehiculo
                FROM cotizaciones_logisticas cl
                WHERE cl.id_cotizacion_operacion = c.id_cotizacion_operacion
                ORDER BY cl.id_cotizacion DESC
                LIMIT 1
            ) AS tipo_vehiculo_logistico,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            p.nombre AS producto,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            l.codigo_lote,
            (
                SELECT v.produccion_min_estimada
                FROM verificaciones v
                WHERE v.id_solicitud = l.id_solicitud
                  AND v.resultado = 'APROBADO'
                  AND v.estado = 'COMPLETADA'
                  AND v.produccion_min_estimada > 0
                  AND v.produccion_max_estimada >= v.produccion_min_estimada
                ORDER BY v.fecha_realizada DESC, v.id_verificacion DESC
                LIMIT 1
            ) AS produccion_min_kg,
            (
                SELECT v.produccion_max_estimada
                FROM verificaciones v
                WHERE v.id_solicitud = l.id_solicitud
                  AND v.resultado = 'APROBADO'
                  AND v.estado = 'COMPLETADA'
                  AND v.produccion_min_estimada > 0
                  AND v.produccion_max_estimada >= v.produccion_min_estimada
                ORDER BY v.fecha_realizada DESC, v.id_verificacion DESC
                LIMIT 1
            ) AS produccion_max_kg,
            de.nombre_referencia AS referencia_entrega,
            de.direccion AS direccion_entrega,
            ud.departamento AS departamento_entrega,
            ud.provincia AS provincia_entrega,
            ud.distrito AS distrito_entrega

        FROM cotizaciones_operacion c

        INNER JOIN lotes l
            ON l.id_lote = c.id_lote

        LEFT JOIN direcciones_entrega de
            ON de.id_direccion = c.id_direccion_entrega

        LEFT JOIN ubigeos ud
            ON ud.id_ubigeo = de.id_ubigeo

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = s.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = c.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        ORDER BY c.fecha_calculo DESC
        """
    )

    tipos_vehiculo = query(
        "SELECT * FROM vw_vehiculos_logistica ORDER BY capacidad_util_kg"
    )
    return render_template(
        "admin_cotizaciones.html",
        cotizaciones=cotizaciones,
        tipos_vehiculo=tipos_vehiculo
    )


@app.route("/admin/cotizaciones/<int:id_cotizacion>/aceptar", methods=["POST"])
@login_required
def aceptar_cotizacion(id_cotizacion):

    if not tiene_permiso("TRANSPORTE_ASIGNAR"):

        flash(
            "No tienes permisos para aceptar cotizaciones.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    cotizacion = query(
        """
        SELECT
            c.*,
            l.id_lote,
            l.id_solicitud,
            s.id_productor,
            s.id_producto,
            p.nombre AS producto,
            pred.id_ubigeo AS id_ubigeo_origen,
            pred.latitud AS latitud_origen,
            pred.longitud AS longitud_origen,
            de.latitud AS latitud_destino,
            de.longitud AS longitud_destino,
            CONCAT(
                u.nombres,
                ' ',
                COALESCE(u.apellido_paterno, '')
            ) AS comprador
        FROM cotizaciones_operacion c
        INNER JOIN lotes l
            ON l.id_lote = c.id_lote
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        LEFT JOIN predios pred
            ON pred.id_predio = s.id_predio
        LEFT JOIN direcciones_entrega de
            ON de.id_direccion = c.id_direccion_entrega
        INNER JOIN compradores co
            ON co.id_comprador = c.id_comprador
        INNER JOIN usuarios u
            ON u.id_usuario = co.id_usuario
        WHERE c.id_cotizacion_operacion = %s
        LIMIT 1
        """,
        (id_cotizacion,),
        one=True
    )

    if not cotizacion:
        flash("La cotización no existe.", "danger")
        return redirect(url_for("admin_cotizaciones"))

    if cotizacion["estado"] != "VIGENTE":
        flash("Solo se pueden aceptar cotizaciones vigentes.", "warning")
        return redirect(url_for("admin_cotizaciones"))

    direccion_entrega = query(
        """
        SELECT id_direccion
        FROM direcciones_entrega
        WHERE id_direccion = %s
          AND id_comprador = %s
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        (cotizacion["id_direccion_entrega"], cotizacion["id_comprador"]),
        one=True
    )
    if not direccion_entrega:
        flash("El comprador debe actualizar una dirección de entrega válida.", "warning")
        return redirect(url_for("cuenta"))

    if query(
        """
        SELECT id_operacion
        FROM operaciones
        WHERE id_lote = %s
        LIMIT 1
        """,
        (cotizacion["id_lote"],),
        one=True
    ):
        flash("Ya existe una operación creada para este lote.", "warning")
        return redirect(url_for("admin_cotizaciones"))

    ruta_calculada = calcular_ruta_osrm(
        cotizacion.get("latitud_origen"),
        cotizacion.get("longitud_origen"),
        cotizacion.get("latitud_destino"),
        cotizacion.get("longitud_destino"),
    )
    try:
        distancia_manual = Decimal(request.form.get("distancia_km", ""))
        peso_kg = Decimal(request.form.get("peso_kg", ""))
        peajes = Decimal(request.form.get("peajes", "0") or "0")
    except InvalidOperation:
        if not ruta_calculada:
            flash("Indica una distancia, carga y peajes válidos para calcular el flete.", "warning")
            return redirect(url_for("admin_cotizaciones"))
        try:
            peso_kg = Decimal(request.form.get("peso_kg", ""))
            peajes = Decimal(request.form.get("peajes", "0") or "0")
        except InvalidOperation:
            flash("Indica una carga y peajes válidos para calcular el flete.", "warning")
            return redirect(url_for("admin_cotizaciones"))
        distancia_manual = None

    distancia_km = (
        ruta_calculada["distancia_km"] if ruta_calculada else distancia_manual
    )
    proveedor_ruta = "OSRM" if ruta_calculada else "CAPTURA_ADMIN"
    duracion_estimada_min = (
        ruta_calculada["duracion_estimada_min"] if ruta_calculada else None
    )
    if not ruta_calculada and distancia_manual is None:
        flash(
            "No se pudo calcular la ruta; indica manualmente los kilómetros por carretera.",
            "warning"
        )
        return redirect(url_for("admin_cotizaciones"))
    id_tipo_vehiculo = request.form.get("id_tipo_vehiculo", type=int)
    if (
        distancia_km is None or not distancia_km.is_finite()
        or distancia_km <= 0 or distancia_km > 10000
        or not peso_kg.is_finite() or peso_kg <= 0
        or not peajes.is_finite() or peajes < 0
        or not id_tipo_vehiculo
    ):
        flash("Revisa los kilómetros, la carga, los peajes y el tipo de vehículo.", "warning")
        return redirect(url_for("admin_cotizaciones"))

    tipo_vehiculo = query(
        """
        SELECT *
        FROM vw_vehiculos_logistica
        WHERE id_tipo_vehiculo = %s
        LIMIT 1
        """,
        (id_tipo_vehiculo,),
        one=True
    )
    if not tipo_vehiculo:
        flash("El tipo de vehículo seleccionado no está activo.", "warning")
        return redirect(url_for("admin_cotizaciones"))

    ubicacion_origen = query(
        """
        SELECT departamento, provincia, distrito
        FROM ubigeos
        WHERE id_ubigeo = %s
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        (cotizacion.get("id_ubigeo_origen"),),
        one=True
    ) if cotizacion.get("id_ubigeo_origen") else None
    departamento = (ubicacion_origen or {}).get("departamento")
    provincia = (ubicacion_origen or {}).get("provincia")
    distrito = (ubicacion_origen or {}).get("distrito")
    combustible = query(
        """
        SELECT *
        FROM vw_combustible_logistica_actual
        WHERE LOWER(combustible) LIKE '%diesel%'
          AND (departamento = %s OR departamento IS NULL)
          AND (provincia = %s OR provincia IS NULL)
          AND (distrito = %s OR distrito IS NULL)
        ORDER BY (distrito = %s) DESC,
                 (provincia = %s) DESC,
                 (departamento = %s) DESC,
                 fecha_precio DESC
        LIMIT 1
        """,
        (
            departamento, provincia, distrito,
            distrito, provincia, departamento
        ),
        one=True
    )
    if not combustible:
        flash(
            "No hay un precio vigente de diésel para la ubicación del predio. "
            "Registra el precio en Configuración logística antes de aceptar.",
            "warning"
        )
        return redirect(url_for("admin_cotizaciones"))

    def obtener_parametro_logistico(codigo, predeterminado):
        fila = query(
            """
            SELECT valor
            FROM parametros_logisticos
            WHERE codigo = %s
              AND estado = 'ACTIVO'
            LIMIT 1
            """,
            (codigo,),
            one=True
        )
        return Decimal(str(fila["valor"] if fila else predeterminado))

    try:
        costos_flete = calcular_costos_logisticos(
            distancia_km,
            peso_kg,
            tipo_vehiculo,
            combustible["precio_galon"],
            obtener_parametro_logistico("FACTOR_RUTA_RETORNO", 1),
            obtener_parametro_logistico("RESERVA_CAPACIDAD_VEHICULO", 5),
            obtener_parametro_logistico("CARGA_DESCARGA", 0),
            peajes,
            obtener_parametro_logistico("SEGURO_CARGA", 0)
        )
    except ValueError as error:
        flash(str(error), "warning")
        return redirect(url_for("admin_cotizaciones"))

    margen_transporte_pct = Decimal(str(cotizacion["margen_transporte_pct"] or 5))
    margen_transporte_monto = (
        costos_flete["costo_base_flete"] * margen_transporte_pct / Decimal("100")
    ).quantize(Decimal("0.01"))
    precio_flete_cliente = costos_flete["costo_base_flete"] + margen_transporte_monto
    total_operacion = (
        Decimal(str(cotizacion["precio_cosecha_cliente"]))
        + precio_flete_cliente
        + Decimal(str(cotizacion["comision_venta_monto"]))
    ).quantize(Decimal("0.01"))
    ingreso_total_agrovia = (
        Decimal(str(cotizacion["comision_venta_monto"]))
        + Decimal(str(cotizacion["margen_cosecha_monto"]))
        + margen_transporte_monto
    ).quantize(Decimal("0.01"))

    execute(
        """
        UPDATE cotizaciones_operacion
        SET costo_base_flete = %s,
            margen_transporte_monto = %s,
            precio_flete_cliente = %s,
            ingreso_total_agrovia = %s,
            total_operacion = %s
        WHERE id_cotizacion_operacion = %s
          AND estado = 'VIGENTE'
        """,
        (
            costos_flete["costo_base_flete"],
            margen_transporte_monto,
            precio_flete_cliente,
            ingreso_total_agrovia,
            total_operacion,
            id_cotizacion
        )
    )

    cotizacion_logistica = query(
        """
        SELECT id_cotizacion
        FROM cotizaciones_logisticas
        WHERE id_cotizacion_operacion = %s
        ORDER BY id_cotizacion DESC
        LIMIT 1
        """,
        (id_cotizacion,),
        one=True
    )
    valores_logisticos = (
        distancia_km,
        proveedor_ruta,
        duracion_estimada_min,
        cotizacion.get("latitud_origen"),
        cotizacion.get("longitud_origen"),
        cotizacion.get("latitud_destino"),
        cotizacion.get("longitud_destino"),
        peso_kg,
        tipo_vehiculo["nombre"],
        costos_flete["cantidad_vehiculos"],
        costos_flete["precio_combustible"],
        costos_flete["rendimiento"],
        costos_flete["costo_combustible"],
        costos_flete["costo_operativo"],
        costos_flete["otros_costos"],
        costos_flete["carga_descarga"],
        costos_flete["peajes"],
        costos_flete["seguro"],
        costos_flete["costo_base_flete"],
        margen_transporte_pct,
        margen_transporte_monto,
        precio_flete_cliente
    )
    if cotizacion_logistica:
        execute(
            """
            UPDATE cotizaciones_logisticas
            SET distancia_km = %s, proveedor_ruta = %s,
                duracion_estimada_min = %s,
                latitud_origen = %s, longitud_origen = %s,
                latitud_destino = %s, longitud_destino = %s,
                peso_kg = %s,
                tipo_vehiculo = %s, cantidad_vehiculos = %s,
                precio_combustible = %s, rendimiento = %s,
                costo_combustible = %s, costo_operativo = %s,
                otros_costos = %s, carga_descarga = %s, peajes = %s, seguro = %s,
                costo_base_flete = %s, margen_agrovia_pct = %s,
                margen_agrovia_monto = %s, precio_flete_cliente = %s,
                fecha_calculo = NOW()
            WHERE id_cotizacion = %s
            """,
            valores_logisticos + (cotizacion_logistica["id_cotizacion"],)
        )
    else:
        execute(
            """
            INSERT INTO cotizaciones_logisticas
            (
                codigo_cotizacion, id_cotizacion_operacion, modo_calculo,
                distancia_km, proveedor_ruta, duracion_estimada_min,
                latitud_origen, longitud_origen, latitud_destino, longitud_destino,
                peso_kg, tipo_vehiculo,
                cantidad_vehiculos, precio_combustible, rendimiento,
                costo_combustible, costo_operativo, otros_costos,
                carga_descarga, peajes, seguro, costo_base_flete, margen_agrovia_pct,
                margen_agrovia_monto, precio_flete_cliente, fecha_calculo
            )
            VALUES
            (
                %s, %s, 'AUTOMATICO', %s, %s, %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, NOW()
            )
            """,
            (
                generar_codigo("FLT", "cotizaciones_logisticas", "codigo_cotizacion"),
                id_cotizacion
            ) + valores_logisticos
        )

    cotizacion["costo_base_flete"] = costos_flete["costo_base_flete"]
    cotizacion["margen_transporte_monto"] = margen_transporte_monto
    cotizacion["precio_flete_cliente"] = precio_flete_cliente
    cotizacion["ingreso_total_agrovia"] = ingreso_total_agrovia
    cotizacion["total_operacion"] = total_operacion

    direccion_id = direccion_entrega["id_direccion"]
    codigo_operacion = generar_codigo("OP", "operaciones", "codigo_operacion")

    id_operacion = execute(
        """
        INSERT INTO operaciones
        (
            codigo_operacion,
            id_lote,
            id_productor,
            id_comprador,
            id_direccion_entrega,
            id_cotizacion_operacion,
            modalidad_operacion,
            valor_produccion,
            costo_base_cosecha,
            margen_cosecha_pct,
            margen_cosecha_monto,
            precio_cosecha_cliente,
            costo_base_flete,
            margen_transporte_pct,
            margen_transporte_monto,
            precio_flete_cliente,
            comision_venta_pct,
            comision_venta_monto,
            ingreso_total_agrovia,
            total_operacion,
            estado,
            fecha_operacion,
            fecha_creacion,
            fecha_actualizacion
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            'COSECHA_PROXIMA',
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            'PENDIENTE_PAGO',
            NOW(),
            NOW(),
            NOW()
        )
        """,
        (
            codigo_operacion,
            cotizacion["id_lote"],
            cotizacion["id_productor"],
            cotizacion["id_comprador"],
            direccion_id,
            cotizacion["id_cotizacion_operacion"],
            cotizacion["valor_produccion"],
            cotizacion["costo_base_cosecha"],
            cotizacion["margen_cosecha_pct"],
            cotizacion["margen_cosecha_monto"],
            cotizacion["precio_cosecha_cliente"],
            cotizacion["costo_base_flete"],
            cotizacion["margen_transporte_pct"],
            cotizacion["margen_transporte_monto"],
            cotizacion["precio_flete_cliente"],
            cotizacion["comision_venta_pct"],
            cotizacion["comision_venta_monto"],
            cotizacion["ingreso_total_agrovia"],
            cotizacion["total_operacion"]
        )
    )

    registrar_historial_estado(
        "OPERACION",
        id_operacion,
        None,
        "PENDIENTE_PAGO",
        (
            "Operación creada tras aceptar la cotización. Flete: "
            f"{distancia_km} km, {peso_kg} kg, {costos_flete['cantidad_vehiculos']} "
            f"vehículo(s) {tipo_vehiculo['nombre']}; combustible {combustible['combustible']} "
            f"a S/ {costos_flete['precio_combustible']}/galón "
            f"({combustible.get('calidad_dato') or 'calidad no indicada'}). "
            f"Flete base S/ {costos_flete['costo_base_flete']}, "
            f"margen S/ {margen_transporte_monto}, "
            f"flete al comprador S/ {precio_flete_cliente}."
        )
    )
    notificar_participantes_operacion(
        id_operacion,
        "Nueva operación",
        f"La operación {codigo_operacion} fue creada y está pendiente de pago.",
        "INFO"
    )

    execute(
        """
        UPDATE cotizaciones_operacion
        SET estado = 'ACEPTADA'
        WHERE id_cotizacion_operacion = %s
        """,
        (id_cotizacion,)
    )

    execute(
        """
        UPDATE reservas_lote
        SET estado = 'CONFIRMADA'
        WHERE id_cotizacion_operacion = %s
          AND estado = 'ACTIVA'
        """,
        (id_cotizacion,)
    )

    execute(
        """
        UPDATE lotes
        SET estado = 'VENDIDO'
        WHERE id_lote = %s
        """,
        (cotizacion["id_lote"],)
    )

    flash("Cotización aceptada y operación creada correctamente.", "success")

    return redirect(url_for("operaciones"))


@app.route("/lotes/<int:id_lote>/cotizar", methods=["POST"])
@login_required
def generar_cotizacion_lote(id_lote):

    comprador = obtener_comprador_actual()

    if not comprador:

        flash(
            "Debes estar registrado como comprador para generar una cotización.",
            "warning"
        )

        return redirect(url_for("lotes"))

    lote = query(
        """
        SELECT *
        FROM lotes
        WHERE id_lote = %s
        LIMIT 1
        """,
        (id_lote,),
        one=True
    )

    if not lote:
        flash("El lote no existe.", "danger")
        return redirect(url_for("lotes"))

    if lote["estado"] not in ("DISPONIBLE", "RESERVADO"):
        flash("Este lote ya no puede cotizarse.", "warning")
        return redirect(url_for("lotes"))

    existente = query(
        """
        SELECT id_cotizacion_operacion
        FROM cotizaciones_operacion
        WHERE id_lote = %s
          AND id_comprador = %s
          AND estado IN ('VIGENTE', 'ACEPTADA')
        LIMIT 1
        """,
        (id_lote, comprador["id_comprador"]),
        one=True
    )

    if existente:
        flash("Ya existe una cotización vigente para este lote.", "warning")
        return redirect(url_for("lotes"))

    direccion_entrega = query(
        """
        SELECT id_direccion
        FROM direcciones_entrega
        WHERE id_comprador = %s
          AND estado = 'ACTIVO'
        ORDER BY predeterminada DESC, id_direccion DESC
        LIMIT 1
        """,
        (comprador["id_comprador"],),
        one=True
    )
    if not direccion_entrega:
        flash("Registra primero una dirección de entrega en Mi cuenta.", "warning")
        return redirect(url_for("cuenta"))

    cotizacion = calcular_cotizacion_lote(id_lote, comprador["id_comprador"])

    if not cotizacion:
        flash("No se pudo construir la cotización del lote.", "danger")
        return redirect(url_for("lotes"))

    id_cotizacion = execute(
        """
        INSERT INTO cotizaciones_operacion
        (
            codigo_cotizacion,
            id_lote,
            id_comprador,
            id_direccion_entrega,
            valor_produccion,
            costo_base_cosecha,
            margen_cosecha_pct,
            margen_cosecha_monto,
            precio_cosecha_cliente,
            costo_base_flete,
            margen_transporte_pct,
            margen_transporte_monto,
            precio_flete_cliente,
            comision_venta_pct,
            comision_venta_monto,
            ingreso_total_agrovia,
            total_operacion,
            fecha_vencimiento,
            estado
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            DATE_ADD(NOW(), INTERVAL 7 DAY),
            'VIGENTE'
        )
        """,
        (
            generar_codigo("COT", "cotizaciones_operacion", "codigo_cotizacion"),
            id_lote,
            comprador["id_comprador"],
            direccion_entrega["id_direccion"],
            cotizacion["valor_produccion"],
            cotizacion["costo_base_cosecha"],
            cotizacion["margen_cosecha_pct"],
            cotizacion["margen_cosecha_monto"],
            cotizacion["precio_cosecha_cliente"],
            cotizacion["costo_base_flete"],
            cotizacion["margen_transporte_pct"],
            cotizacion["margen_transporte_monto"],
            cotizacion["precio_flete_cliente"],
            cotizacion["comision_venta_pct"],
            cotizacion["comision_venta_monto"],
            cotizacion["ingreso_total_agrovia"],
            cotizacion["total_operacion"]
        )
    )

    execute(
        """
        INSERT INTO cotizaciones_logisticas
        (
            codigo_cotizacion,
            id_cotizacion_operacion,
            costo_base_flete,
            margen_agrovia_pct,
            margen_agrovia_monto,
            precio_flete_cliente,
            fecha_calculo
        )
        VALUES
        (
            %s,
            %s,
            0,
            5,
            0,
            0,
            NOW()
        )
        """,
        (
            generar_codigo("FLT", "cotizaciones_logisticas", "codigo_cotizacion"),
            id_cotizacion
        )
    )

    flash("Cotización generada correctamente y enviada para aprobación del administrador.", "success")

    return redirect(url_for("lotes"))


@app.route("/admin/valorizaciones/<int:id_valorizacion>/calcular", methods=["POST"])
@login_required
def recalcular_valorizacion(id_valorizacion):

    if not tiene_permiso("PRECIOS_MODIFICAR"):

        flash(
            "No tienes permisos para recalcular valorizaciones.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    registro = query(
        """
        SELECT id_solicitud, modo_calculo, estado
        FROM valorizaciones
        WHERE id_valorizacion = %s
        LIMIT 1
        """,
        (id_valorizacion,),
        one=True
    )

    if not registro:
        flash("La valorización no existe.", "danger")
        return redirect(url_for("admin_valorizaciones"))

    if registro["modo_calculo"] != "AUTOMATICO":
        flash("Esta valorización es manual y no se puede recalcular automáticamente.", "warning")
        return redirect(url_for("admin_valorizaciones"))

    if registro["estado"] in ("PRESENTADA", "ACEPTADA", "ANULADA"):
        flash("No se puede recalcular una valorización presentada, aceptada o anulada.", "warning")
        return redirect(url_for("admin_valorizaciones"))

    resultado = calcular_valorizacion_solicitud(registro["id_solicitud"])
    if resultado is None:
        flash("No hay precios de mercado activos para recalcular este producto.", "warning")
        return redirect(url_for("admin_valorizaciones"))

    flash("Valorización recalculada con los precios de mercado disponibles.", "success")

    return redirect(url_for("admin_valorizaciones"))


# ==========================================================
# OPERACIONES
# ==========================================================

@app.route("/mis-operaciones")
@login_required
def mis_operaciones():
    operaciones_usuario = query(
        """
        SELECT o.id_operacion, o.codigo_operacion, o.total_operacion, o.estado,
               o.fecha_operacion, p.nombre AS producto, l.codigo_lote,
               t.id_transporte, t.codigo_transporte,
               t.origen AS transporte_origen, t.destino AS transporte_destino,
               t.estado AS estado_transporte,
               de.direccion AS direccion_entrega,
               de.latitud AS latitud_destino,
               de.longitud AS longitud_destino,
               tr.nombre_empresa AS transportista,
               v.placa,
               cl.distancia_km, cl.duracion_estimada_min,
               ubicacion.latitud AS latitud_actual,
               ubicacion.longitud AS longitud_actual,
               ubicacion.fecha_reporte AS fecha_ubicacion,
               pago.id_pago, pago.codigo_pago, pago.metodo,
               pago.referencia, pago.estado AS estado_pago,
               df.tipo_documento, df.numero_documento, df.nombre_razon_social,
               IF(uc.id_usuario = %s, 1, 0) AS puede_pagar,
               IF(
                   o.estado = 'COMPLETADA'
                   AND NOT EXISTS (
                       SELECT 1
                       FROM calificaciones cal
                       WHERE cal.id_operacion = o.id_operacion
                         AND cal.id_usuario_califica = %s
                   ),
                   1,
                   0
               ) AS puede_calificar
        FROM operaciones o
        INNER JOIN lotes l ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p ON p.id_producto = s.id_producto
        INNER JOIN productores pr ON pr.id_productor = o.id_productor
        INNER JOIN usuarios up ON up.id_usuario = pr.id_usuario
        INNER JOIN compradores co ON co.id_comprador = o.id_comprador
        INNER JOIN usuarios uc ON uc.id_usuario = co.id_usuario
        LEFT JOIN transportes t ON t.id_operacion = o.id_operacion
        LEFT JOIN transportistas tr ON tr.id_transportista = t.id_transportista
        LEFT JOIN vehiculos v ON v.id_vehiculo = t.id_vehiculo
        LEFT JOIN direcciones_entrega de ON de.id_direccion = o.id_direccion_entrega
        LEFT JOIN cotizaciones_logisticas cl ON cl.id_cotizacion = (
            SELECT MAX(cl_actual.id_cotizacion)
            FROM cotizaciones_logisticas cl_actual
            WHERE cl_actual.id_cotizacion_operacion = o.id_cotizacion_operacion
        )
        LEFT JOIN ubicaciones_transporte ubicacion ON ubicacion.id_ubicacion = (
            SELECT MAX(ut.id_ubicacion)
            FROM ubicaciones_transporte ut
            WHERE ut.id_transporte = t.id_transporte
        )
        LEFT JOIN (
            SELECT p1.*
            FROM pagos p1
            INNER JOIN (
                SELECT id_operacion, MAX(id_pago) AS id_pago
                FROM pagos
                GROUP BY id_operacion
            ) latest ON latest.id_pago = p1.id_pago
        ) pago ON pago.id_operacion = o.id_operacion
        LEFT JOIN datos_facturacion df ON df.id_facturacion = pago.id_facturacion
        WHERE up.id_usuario = %s OR uc.id_usuario = %s
        ORDER BY o.fecha_operacion DESC
        """,
        (
            session["user_id"],
            session["user_id"],
            session["user_id"],
            session["user_id"]
        )
    )

    if operaciones_usuario:
        ids_operacion = [item["id_operacion"] for item in operaciones_usuario]
        placeholders = ", ".join(["%s"] * len(ids_operacion))
        eventos = query(
            f"""
            SELECT h.entidad_id, h.estado_anterior, h.estado_nuevo, h.fecha,
                   h.observacion, u.nombres, u.apellido_paterno
            FROM historial_estados h
            LEFT JOIN usuarios u ON u.id_usuario = h.id_usuario
            WHERE h.entidad_tipo = 'OPERACION'
              AND h.entidad_id IN ({placeholders})
            ORDER BY h.fecha ASC, h.id_historial ASC
            """,
            tuple(ids_operacion)
        )
        historial_por_operacion = {}
        for evento in eventos:
            historial_por_operacion.setdefault(evento["entidad_id"], []).append(evento)
        for item in operaciones_usuario:
            item["historial"] = historial_por_operacion.get(item["id_operacion"], [])
        incidencias = query(
            f"""
            SELECT id_operacion, codigo_incidencia, estado, solucion
            FROM incidencias
            WHERE id_operacion IN ({placeholders})
            ORDER BY fecha_apertura DESC, id_incidencia DESC
            """,
            tuple(ids_operacion)
        )
        incidencias_por_operacion = {}
        for incidencia in incidencias:
            incidencias_por_operacion.setdefault(
                incidencia["id_operacion"], []
            ).append(incidencia)
        for item in operaciones_usuario:
            item["incidencias"] = incidencias_por_operacion.get(
                item["id_operacion"], []
            )

    return render_template("mis_operaciones.html", operaciones=operaciones_usuario)


@app.route("/mis-operaciones/<int:id_operacion>/pago", methods=["GET", "POST"])
@login_required
def solicitar_pago_operacion(id_operacion):
    comprador = obtener_comprador_actual()
    if not comprador:
        flash("Solo un comprador puede reportar un pago.", "warning")
        return redirect(url_for("dashboard"))

    operacion = query(
        """
        SELECT id_operacion, codigo_operacion, total_operacion, estado
        FROM operaciones
        WHERE id_operacion = %s AND id_comprador = %s
        LIMIT 1
        """,
        (id_operacion, comprador["id_comprador"]),
        one=True
    )
    if not operacion:
        flash("La operación no existe o no te pertenece.", "warning")
        return redirect(url_for("mis_operaciones"))
    if operacion["estado"] != "PENDIENTE_PAGO":
        flash("Esta operación no está pendiente de pago.", "warning")
        return redirect(url_for("mis_operaciones"))

    facturaciones = query(
        """
        SELECT id_facturacion, tipo_documento, numero_documento, nombre_razon_social,
               direccion_fiscal, predeterminado
        FROM datos_facturacion
        WHERE id_comprador = %s
          AND estado = 'ACTIVO'
          AND tipo_documento IN ('DNI', 'RUC')
        ORDER BY predeterminado DESC, id_facturacion DESC
        """,
        (comprador["id_comprador"],)
    )
    if not facturaciones:
        flash("Registra primero tus datos de facturación con DNI o RUC en Mi cuenta.", "warning")
        return redirect(url_for("cuenta"))

    if request.method == "POST":
        try:
            id_facturacion = int(request.form.get("id_facturacion", ""))
        except ValueError:
            flash("Selecciona un comprobante válido.", "warning")
            return redirect(url_for("solicitar_pago_operacion", id_operacion=id_operacion))

        metodo = request.form.get("metodo", "")
        referencia = request.form.get("referencia", "").strip()
        metodos_validos = {"TRANSFERENCIA_BANCARIA", "BILLETERA_DIGITAL", "EFECTIVO", "OTRO"}
        if metodo not in metodos_validos or len(referencia) > 120:
            flash("Selecciona un método válido y revisa la referencia.", "warning")
            return redirect(url_for("solicitar_pago_operacion", id_operacion=id_operacion))
        if metodo in {"TRANSFERENCIA_BANCARIA", "BILLETERA_DIGITAL"} and not referencia:
            flash("Indica el número de operación o referencia del pago.", "warning")
            return redirect(url_for("solicitar_pago_operacion", id_operacion=id_operacion))

        billing = query(
            """
            SELECT id_facturacion, tipo_documento, numero_documento
            FROM datos_facturacion
            WHERE id_facturacion = %s
              AND id_comprador = %s
              AND estado = 'ACTIVO'
              AND tipo_documento IN ('DNI', 'RUC')
            LIMIT 1
            """,
            (id_facturacion, comprador["id_comprador"]),
            one=True
        )
        if not billing or not documento_facturacion_valido(
            billing["tipo_documento"], billing["numero_documento"]
        ):
            flash("El documento fiscal seleccionado no es válido. Actualízalo en Mi cuenta.", "warning")
            return redirect(url_for("cuenta"))

        codigo_pago = generar_codigo("PG", "pagos", "codigo_pago")
        conn = get_connection()
        cur = conn.cursor(dictionary=True)
        try:
            conn.start_transaction()
            cur.execute(
                """
                SELECT estado
                FROM operaciones
                WHERE id_operacion = %s AND id_comprador = %s
                FOR UPDATE
                """,
                (id_operacion, comprador["id_comprador"])
            )
            estado_operacion = cur.fetchone()
            if not estado_operacion or estado_operacion["estado"] != "PENDIENTE_PAGO":
                conn.rollback()
                flash("La operación dejó de estar pendiente de pago.", "warning")
                return redirect(url_for("mis_operaciones"))

            cur.execute(
                """
                SELECT id_pago
                FROM pagos
                WHERE id_operacion = %s
                  AND estado IN ('PENDIENTE', 'RECIBIDO', 'RETENIDO')
                LIMIT 1
                FOR UPDATE
                """,
                (id_operacion,)
            )
            if cur.fetchone():
                conn.rollback()
                flash("Ya existe un pago pendiente o confirmado para esta operación.", "warning")
                return redirect(url_for("mis_operaciones"))

            cur.execute(
                """
                INSERT INTO pagos
                (
                    codigo_pago, id_operacion, id_facturacion, tipo_pago,
                    monto, metodo, referencia, fecha, estado
                )
                VALUES (%s, %s, %s, 'PAGO_COMPRADOR', %s, %s, %s, NOW(), 'PENDIENTE')
                """,
                (
                    codigo_pago,
                    id_operacion,
                    id_facturacion,
                    operacion["total_operacion"],
                    metodo,
                    referencia or None
                )
            )
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()

        notificar_participantes_operacion(
            id_operacion,
            "Pago reportado",
            f"El comprador reportó el pago de la operación {operacion['codigo_operacion']}. "
            "Administración debe verificarlo.",
            "INFO"
        )
        notificar_administradores(
            "Pago pendiente de verificación",
            f"Se reportó un pago para la operación {operacion['codigo_operacion']}.",
            "WARNING",
            "OPERACION",
            id_operacion,
            "PAGOS_AUTORIZAR"
        )

        flash("Pago reportado. Administración verificará el pago y el comprobante seleccionado.", "success")
        return redirect(url_for("mis_operaciones"))

    return render_template(
        "pago_operacion.html",
        operacion=operacion,
        facturaciones=facturaciones
    )


@app.route("/operaciones")
@login_required
def operaciones():

    if not tiene_permiso("REPORTES_VER"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("index"))

    rows = query(
        """
        SELECT
            o.id_operacion,
            o.codigo_operacion,
            o.total_operacion,
            o.estado,
            o.modalidad_operacion,
            o.fecha_operacion,
            p.nombre AS producto,

            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno,'')
            ) AS productor,

            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno,'')
            ) AS comprador

        FROM operaciones o

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        ORDER BY o.fecha_operacion DESC
        """
    )

    if rows:
        ids_operacion = [row["id_operacion"] for row in rows]
        placeholders = ", ".join(["%s"] * len(ids_operacion))
        historial = query(
            f"""
            SELECT entidad_id, estado_anterior, estado_nuevo, fecha, observacion
            FROM historial_estados
            WHERE entidad_tipo = 'OPERACION'
              AND entidad_id IN ({placeholders})
            ORDER BY fecha DESC
            """,
            tuple(ids_operacion)
        )
        historial_por_operacion = {}
        for evento in historial:
            historial_por_operacion.setdefault(evento["entidad_id"], []).append(evento)
        for row in rows:
            row["historial"] = historial_por_operacion.get(row["id_operacion"], [])

    return render_template(
        "operaciones.html",
        operations=rows
    )


@app.route("/admin/pagos")
@login_required
def admin_pagos():

    if not tiene_permiso("PAGOS_VER"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    pagos = query(
        """
        SELECT
            o.id_operacion,
            o.codigo_operacion,
            o.total_operacion,
            o.estado AS estado_operacion,
            p.id_pago,
            p.codigo_pago,
            p.id_facturacion,
            p.monto AS monto_pago,
            p.metodo,
            p.fecha AS fecha_pago,
            p.estado AS estado_pago,
            df.tipo_documento,
            df.numero_documento,
            df.nombre_razon_social,
            CASE
                WHEN df.tipo_documento = 'DNI' THEN 'BOLETA'
                WHEN df.tipo_documento = 'RUC' THEN 'FACTURA'
                ELSE NULL
            END AS tipo_comprobante,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            prd.nombre AS producto

        FROM operaciones o

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos prd
            ON prd.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        LEFT JOIN (
            SELECT p1.*
            FROM pagos p1
            INNER JOIN (
                SELECT id_operacion, MAX(id_pago) AS id_pago
                FROM pagos
                GROUP BY id_operacion
            ) latest
                ON latest.id_pago = p1.id_pago
        ) p
            ON p.id_operacion = o.id_operacion

        LEFT JOIN datos_facturacion df
            ON df.id_facturacion = p.id_facturacion

        ORDER BY o.fecha_operacion DESC
        """
    )

    return render_template(
        "admin_pagos.html",
        pagos=pagos
    )


@app.route("/admin/pagos/<int:id_operacion>/confirmar", methods=["POST"])
@login_required
def confirmar_pago(id_operacion):

    if not tiene_permiso("PAGOS_AUTORIZAR"):

        flash(
            "No tienes permisos para confirmar pagos.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    operacion = query(
        """
        SELECT o.*, p.nombre AS producto
        FROM operaciones o
        INNER JOIN lotes l
            ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        WHERE o.id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not operacion:
        flash("La operación no existe.", "danger")
        return redirect(url_for("admin_pagos"))

    pago_pendiente = query(
        """
        SELECT id_pago
        FROM pagos
        WHERE id_operacion = %s
          AND estado = 'PENDIENTE'
        ORDER BY id_pago DESC
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )
    if pago_pendiente:
        execute(
            """
            UPDATE pagos
            SET estado = 'RECIBIDO', fecha = NOW()
            WHERE id_pago = %s AND estado = 'PENDIENTE'
            """,
            (pago_pendiente["id_pago"],)
        )
    else:
        pago_activo = query(
            """
            SELECT id_pago
            FROM pagos
            WHERE id_operacion = %s
              AND estado IN ('RECIBIDO', 'RETENIDO')
            LIMIT 1
            """,
            (id_operacion,),
            one=True
        )
        if pago_activo:
            flash("Ya existe un pago confirmado para esta operación.", "warning")
            return redirect(url_for("admin_pagos"))

        codigo_pago = generar_codigo("PG", "pagos", "codigo_pago")
        execute(
            """
            INSERT INTO pagos
            (
                codigo_pago, id_operacion, tipo_pago, monto,
                metodo, referencia, fecha, estado
            )
            VALUES (%s, %s, 'PAGO_COMPRADOR', %s, 'TRANSFERENCIA_BANCARIA', 'CONFIRMADO_MANUAL', NOW(), 'RECIBIDO')
            """,
            (codigo_pago, id_operacion, operacion["total_operacion"])
        )

    execute(
        """
        UPDATE operaciones
        SET estado = 'PAGO_CONFIRMADO',
            fecha_actualizacion = NOW()
        WHERE id_operacion = %s
        """,
        (id_operacion,)
    )

    registrar_historial_estado(
        "OPERACION", id_operacion, operacion["estado"], "PAGO_CONFIRMADO",
        "Pago confirmado por administración."
    )
    notificar_participantes_operacion(
        id_operacion,
        "Pago confirmado",
        f"El pago de la operación {operacion['codigo_operacion']} fue confirmado. "
        "La operación continuará con la programación de cosecha.",
        "SUCCESS"
    )

    flash("Pago confirmado correctamente y operación actualizada.", "success")

    return redirect(url_for("admin_pagos"))


@app.route("/admin/cosechas")
@login_required
def admin_cosechas():

    if not tiene_permiso("COSECHA_PROGRAMAR"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    cosechas = query(
        """
        SELECT
            o.id_operacion,
            o.codigo_operacion,
            o.total_operacion,
            o.estado AS estado_operacion,
            o.id_lote,
            p.nombre AS producto,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            c.id_cosecha,
            c.codigo_cosecha,
            c.fecha_programada,
            c.estado AS estado_cosecha,
            c.trabajadores_requeridos_estimados,
            ac.numero_trabajadores,
            cu.nombre AS cuadrilla

        FROM operaciones o

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        LEFT JOIN cosechas c
            ON c.id_operacion = o.id_operacion

        LEFT JOIN asignaciones_cuadrilla ac
            ON ac.id_cosecha = c.id_cosecha
            AND ac.estado IN ('ASIGNADA', 'CONFIRMADA')

        LEFT JOIN cuadrillas cu
            ON cu.id_cuadrilla = ac.id_cuadrilla

        WHERE o.estado IN ('PAGO_CONFIRMADO', 'COSECHA_PROGRAMADA', 'EN_COSECHA', 'COSECHA_VALIDADA')
        ORDER BY o.fecha_operacion DESC
        """
    )

    cuadrillas = query(
        """
        SELECT *
        FROM cuadrillas
        WHERE estado = 'ACTIVO'
        ORDER BY nombre ASC
        """
    )
    referencias_cosecha = query(
        """
        SELECT
            demo.producto,
            demo.calidad_dato,
            demo.jornales_referenciales_ha,
            demo.trabajadores_para_1_dia,
            demo.jornal_diario,
            demo.costo_mano_obra,
            demo.costo_carga_tonelada,
            demo.costo_acondicionamiento_tonelada,
            demo.otros_costos_base,
            demo.costo_base_sin_volumen,
            demo.fuente,
            parametros.duracion_objetivo_dias,
            parametros.url_fuente,
            parametros.fecha_fuente
        FROM vw_cosecha_demo_1ha AS demo
        LEFT JOIN vw_parametros_cosecha_demo AS parametros
            ON parametros.producto = demo.producto
        ORDER BY demo.producto
        """
    )
    referencia_por_producto = {
        item["producto"]: item for item in referencias_cosecha
    }
    for operacion in cosechas:
        operacion["referencia_cosecha"] = referencia_por_producto.get(
            operacion["producto"]
        )

    return render_template(
        "admin_cosechas.html",
        cosechas=cosechas,
        cuadrillas=cuadrillas,
        referencias_cosecha=referencias_cosecha
    )


@app.route("/admin/cosechas/<int:id_operacion>/programar", methods=["POST"])
@login_required
def programar_cosecha(id_operacion):

    if not tiene_permiso("COSECHA_PROGRAMAR"):

        flash(
            "No tienes permisos para programar cosechas.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    operacion = query(
        """
        SELECT o.*, l.id_lote, l.codigo_lote, p.nombre AS producto
        FROM operaciones o
        INNER JOIN lotes l
            ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        WHERE o.id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not operacion:
        flash("La operación no existe.", "danger")
        return redirect(url_for("admin_cosechas"))

    if query(
        """
        SELECT id_cosecha
        FROM cosechas
        WHERE id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    ):
        flash("Ya existe una cosecha programada para esta operación.", "warning")
        return redirect(url_for("admin_cosechas"))

    fecha_programada = request.form.get("fecha_programada")
    trabajadores = request.form.get("trabajadores_requeridos_estimados")
    id_cuadrilla = request.form.get("id_cuadrilla")
    jornal_acordado = request.form.get("jornal_acordado") or 80
    observaciones = request.form.get("observaciones") or "Cosecha programada por administración."

    if not fecha_programada:
        flash("Debes indicar la fecha programada de cosecha.", "warning")
        return redirect(url_for("admin_cosechas"))

    try:
        trabajadores = int(trabajadores or 0)
        jornal_acordado = float(jornal_acordado)
    except ValueError:
        flash("El número de trabajadores o jornal acordado no es válido.", "warning")
        return redirect(url_for("admin_cosechas"))

    if trabajadores <= 0:
        flash("Debes indicar una cantidad de trabajadores válida.", "warning")
        return redirect(url_for("admin_cosechas"))

    costo_base_estimado = round(float(operacion["total_operacion"]) * 0.12, 2)
    codigo_cosecha = generar_codigo("CS", "cosechas", "codigo_cosecha")

    id_cosecha = execute(
        """
        INSERT INTO cosechas
        (
            codigo_cosecha,
            id_operacion,
            fecha_programada,
            trabajadores_requeridos_estimados,
            costo_base_estimado,
            estado,
            observaciones,
            id_responsable
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            'PROGRAMADA',
            %s,
            %s
        )
        """,
        (
            codigo_cosecha,
            id_operacion,
            fecha_programada,
            trabajadores,
            costo_base_estimado,
            observaciones,
            session["user_id"]
        )
    )

    if id_cuadrilla:
        execute(
            """
            INSERT INTO asignaciones_cuadrilla
            (
                id_cosecha,
                id_cuadrilla,
                numero_trabajadores,
                jornal_acordado,
                fecha_asignacion,
                estado
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                NOW(),
                'ASIGNADA'
            )
            """,
            (
                id_cosecha,
                id_cuadrilla,
                trabajadores,
                jornal_acordado
            )
        )

    execute(
        """
        UPDATE operaciones
        SET estado = 'COSECHA_PROGRAMADA',
            fecha_actualizacion = NOW()
        WHERE id_operacion = %s
        """,
        (id_operacion,)
    )

    registrar_historial_estado(
        "OPERACION", id_operacion, operacion["estado"], "COSECHA_PROGRAMADA",
        "Cosecha programada por administración."
    )
    notificar_participantes_operacion(
        id_operacion,
        "Cosecha programada",
        f"La cosecha de la operación {operacion['codigo_operacion']} fue programada "
        f"para {fecha_programada}.",
        "INFO"
    )

    execute(
        """
        UPDATE lotes
        SET estado = 'COSECHA_PROGRAMADA'
        WHERE id_lote = %s
        """,
        (operacion["id_lote"],)
    )

    flash("Cosecha programada y asignada correctamente.", "success")

    return redirect(url_for("admin_cosechas"))


@app.route("/admin/transportes")
@login_required
def admin_transportes():

    if not tiene_permiso("TRANSPORTE_ASIGNAR"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    transportes = query(
        """
        SELECT
            o.id_operacion,
            o.codigo_operacion,
            o.total_operacion,
            o.estado AS estado_operacion,
            (
                SELECT cl.peso_kg
                FROM cotizaciones_logisticas cl
                WHERE cl.id_cotizacion_operacion = o.id_cotizacion_operacion
                ORDER BY cl.id_cotizacion DESC
                LIMIT 1
            ) AS peso_transporte,
            p.nombre AS producto,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            t.id_transporte,
            t.codigo_transporte,
            t.estado AS estado_transporte,
            tr.nombre_empresa,
            v.placa,
            v.tipo AS tipo_vehiculo,
            v.capacidad_kg

        FROM operaciones o

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        LEFT JOIN transportes t
            ON t.id_operacion = o.id_operacion

        LEFT JOIN transportistas tr
            ON tr.id_transportista = t.id_transportista

        LEFT JOIN vehiculos v
            ON v.id_vehiculo = t.id_vehiculo

        WHERE o.estado IN ('COSECHA_VALIDADA', 'EN_TRANSPORTE', 'ENTREGADA', 'COMPLETADA')
        ORDER BY o.fecha_operacion DESC
        """
    )

    transportistas = query(
        """
        SELECT tr.*,
               CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, ''))
                   AS nombre_usuario
        FROM transportistas tr
        LEFT JOIN usuarios u ON u.id_usuario = tr.id_usuario
        WHERE tr.estado = 'ACTIVO'
        ORDER BY tr.nombre_empresa ASC
        """
    )

    usuarios_transportistas = query(
        """
        SELECT DISTINCT u.id_usuario,
               CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, ''))
                   AS nombre_usuario
        FROM usuarios u
        INNER JOIN usuarios_roles ur ON ur.id_usuario = u.id_usuario
        INNER JOIN roles r ON r.id_rol = ur.id_rol
        WHERE u.estado = 'ACTIVO'
          AND ur.estado = 'ACTIVO'
          AND r.estado = 'ACTIVO'
          AND r.nombre = 'TRANSPORTISTA'
        ORDER BY u.nombres, u.apellido_paterno
        """
    )

    vehiculos = query(
        """
        SELECT v.*, t.nombre_empresa
        FROM vehiculos v
        INNER JOIN transportistas t
            ON t.id_transportista = v.id_transportista
        WHERE v.estado = 'ACTIVO'
        ORDER BY v.tipo ASC
        """
    )

    return render_template(
        "admin_transportes.html",
        transportes=transportes,
        transportistas=transportistas,
        vehiculos=vehiculos,
        usuarios_transportistas=usuarios_transportistas
    )


@app.route("/admin/transportistas/<int:id_transportista>/vincular", methods=["POST"])
@login_required
def vincular_transportista(id_transportista):
    if not tiene_permiso("TRANSPORTE_ASIGNAR"):
        flash("No tienes permisos para vincular transportistas.", "danger")
        return redirect(url_for("dashboard"))

    transportista = query(
        """
        SELECT id_transportista
        FROM transportistas
        WHERE id_transportista = %s
          AND estado = 'ACTIVO'
        LIMIT 1
        """,
        (id_transportista,),
        one=True
    )
    if not transportista:
        flash("El transportista no existe o está inactivo.", "warning")
        return redirect(url_for("admin_transportes"))

    id_usuario = request.form.get("id_usuario", type=int)
    if id_usuario:
        usuario_transportista = query(
            """
            SELECT u.id_usuario
            FROM usuarios u
            INNER JOIN usuarios_roles ur ON ur.id_usuario = u.id_usuario
            INNER JOIN roles r ON r.id_rol = ur.id_rol
            WHERE u.id_usuario = %s
              AND u.estado = 'ACTIVO'
              AND ur.estado = 'ACTIVO'
              AND r.estado = 'ACTIVO'
              AND r.nombre = 'TRANSPORTISTA'
            LIMIT 1
            """,
            (id_usuario,),
            one=True
        )
        if not usuario_transportista:
            flash("Selecciona una cuenta activa con el rol TRANSPORTISTA.", "warning")
            return redirect(url_for("admin_transportes"))
        vinculo_existente = query(
            """
            SELECT id_transportista
            FROM transportistas
            WHERE id_usuario = %s
              AND id_transportista <> %s
            LIMIT 1
            """,
            (id_usuario, id_transportista),
            one=True
        )
        if vinculo_existente:
            flash("Esta cuenta ya está vinculada a otro transportista.", "warning")
            return redirect(url_for("admin_transportes"))

    execute(
        """
        UPDATE transportistas
        SET id_usuario = %s
        WHERE id_transportista = %s
          AND estado = 'ACTIVO'
        """,
        (id_usuario, id_transportista)
    )
    flash(
        "Cuenta del transportista actualizada." if id_usuario
        else "Se retiró el vínculo de la cuenta del transportista.",
        "success"
    )
    return redirect(url_for("admin_transportes"))


@app.route("/mis-transportes")
@login_required
def mis_transportes():
    transportes = query(
        """
        SELECT t.id_transporte, t.codigo_transporte, t.estado,
               t.origen, t.destino, t.fecha_asignacion,
               o.codigo_operacion, tr.nombre_empresa, v.placa,
               (
                   SELECT ut.fecha_reporte
                   FROM ubicaciones_transporte ut
                   WHERE ut.id_transporte = t.id_transporte
                   ORDER BY ut.id_ubicacion DESC
                   LIMIT 1
               ) AS fecha_ubicacion
        FROM transportes t
        INNER JOIN transportistas tr ON tr.id_transportista = t.id_transportista
        INNER JOIN vehiculos v ON v.id_vehiculo = t.id_vehiculo
        INNER JOIN operaciones o ON o.id_operacion = t.id_operacion
        WHERE tr.id_usuario = %s
          AND t.estado NOT IN ('ENTREGADO', 'CANCELADO')
        ORDER BY t.fecha_asignacion DESC
        """,
        (session["user_id"],)
    )
    return render_template("mis_transportes.html", transportes=transportes)


def obtener_transporte_seguimiento(id_transporte):
    return query(
        """
        SELECT t.id_transporte, t.id_operacion, t.estado,
               tr.id_usuario AS id_usuario_transportista
        FROM transportes t
        INNER JOIN transportistas tr ON tr.id_transportista = t.id_transportista
        WHERE t.id_transporte = %s
        LIMIT 1
        """,
        (id_transporte,),
        one=True
    )


def usuario_participa_en_operacion(id_operacion):
    return bool(query(
        """
        SELECT 1
        FROM operaciones o
        LEFT JOIN productores pr ON pr.id_productor = o.id_productor
        LEFT JOIN compradores co ON co.id_comprador = o.id_comprador
        WHERE o.id_operacion = %s
          AND (pr.id_usuario = %s OR co.id_usuario = %s)
        LIMIT 1
        """,
        (id_operacion, session["user_id"], session["user_id"]),
        one=True
    ))


@app.route("/api/transportes/<int:id_transporte>/ubicacion")
@login_required
def ubicacion_transporte(id_transporte):
    transporte = obtener_transporte_seguimiento(id_transporte)
    if not transporte:
        abort(404)
    es_transportista = (
        transporte["id_usuario_transportista"] == session["user_id"]
    )
    if (
        not es_transportista
        and not es_admin()
        and not usuario_participa_en_operacion(transporte["id_operacion"])
    ):
        abort(403)

    ubicacion = query(
        """
        SELECT latitud, longitud, precision_m, fecha_reporte
        FROM ubicaciones_transporte
        WHERE id_transporte = %s
        ORDER BY id_ubicacion DESC
        LIMIT 1
        """,
        (id_transporte,),
        one=True
    )
    if not ubicacion:
        return jsonify({"ubicacion": None})
    return jsonify({
        "ubicacion": {
            "latitud": float(ubicacion["latitud"]),
            "longitud": float(ubicacion["longitud"]),
            "precision_m": float(ubicacion["precision_m"])
            if ubicacion["precision_m"] is not None else None,
            "fecha_reporte": (
                ubicacion["fecha_reporte"].isoformat()
                if ubicacion["fecha_reporte"] else None
            ),
        }
    })


@app.route("/api/transportes/<int:id_transporte>/ubicacion", methods=["POST"])
@login_required
def reportar_ubicacion_transporte(id_transporte):
    transporte = obtener_transporte_seguimiento(id_transporte)
    if not transporte:
        abort(404)
    if transporte["id_usuario_transportista"] != session["user_id"]:
        abort(403)
    if transporte["estado"] != "EN_RUTA":
        return jsonify({
            "error": "La ubicación solo se comparte mientras el transporte está en ruta."
        }), 409
    if not request.is_json:
        return jsonify({"error": "Se requiere enviar coordenadas en formato JSON."}), 400
    datos = request.get_json()
    if not isinstance(datos, dict):
        return jsonify({"error": "Las coordenadas enviadas no son válidas."}), 400

    try:
        latitud = float(datos.get("latitud"))
        longitud = float(datos.get("longitud"))
        precision = float(datos.get("precision_m", 0))
    except (TypeError, ValueError):
        return jsonify({"error": "Indica latitud y longitud válidas."}), 400
    if (
        not all(math.isfinite(valor) for valor in (latitud, longitud, precision))
        or not -90 <= latitud <= 90
        or not -180 <= longitud <= 180
        or not 0 <= precision <= 10000
    ):
        return jsonify({"error": "Las coordenadas o precisión están fuera de rango."}), 400

    execute(
        """
        INSERT INTO ubicaciones_transporte
            (id_transporte, latitud, longitud, precision_m, fecha_reporte)
        VALUES (%s, %s, %s, %s, NOW())
        ON DUPLICATE KEY UPDATE
            latitud = VALUES(latitud),
            longitud = VALUES(longitud),
            precision_m = VALUES(precision_m),
            fecha_reporte = NOW()
        """,
        (id_transporte, latitud, longitud, precision)
    )
    return jsonify({"guardada": True})


@app.route("/mis-transportes/<int:id_transporte>/estado", methods=["POST"])
@login_required
def actualizar_estado_transporte(id_transporte):
    transporte = obtener_transporte_seguimiento(id_transporte)
    if not transporte:
        flash("El transporte no existe.", "danger")
        return redirect(url_for("mis_transportes"))

    es_transportista = (
        transporte["id_usuario_transportista"] == session["user_id"]
    )
    puede_gestionar = es_admin() or tiene_permiso("TRANSPORTE_ASIGNAR")
    if not es_transportista and not puede_gestionar:
        abort(403)

    transiciones = {
        "ASIGNADO": "RECOJO_PROGRAMADO",
        "RECOJO_PROGRAMADO": "CARGA_REALIZADA",
        "CARGA_REALIZADA": "EN_RUTA",
        "EN_RUTA": "LLEGADA_DESTINO",
    }
    estado_nuevo = request.form.get("estado")
    if transiciones.get(transporte["estado"]) != estado_nuevo:
        flash("El cambio de estado no corresponde a la etapa actual.", "warning")
        return redirect(url_for("mis_transportes"))

    conn = get_connection()
    cur = conn.cursor()
    try:
        conn.start_transaction()
        cur.execute(
            """
            UPDATE transportes
            SET estado = %s,
                fecha_recojo = CASE WHEN %s = 'CARGA_REALIZADA'
                                    THEN NOW() ELSE fecha_recojo END,
                fecha_salida = CASE WHEN %s = 'EN_RUTA'
                                    THEN NOW() ELSE fecha_salida END,
                fecha_llegada = CASE WHEN %s = 'LLEGADA_DESTINO'
                                     THEN NOW() ELSE fecha_llegada END
            WHERE id_transporte = %s
              AND estado = %s
            """,
            (
                estado_nuevo, estado_nuevo, estado_nuevo, estado_nuevo,
                id_transporte, transporte["estado"]
            )
        )
        if cur.rowcount != 1:
            conn.rollback()
            flash("El transporte cambió de estado; vuelve a cargar la página.", "warning")
            return redirect(url_for("mis_transportes"))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()

    registrar_historial_estado(
        "TRANSPORTE",
        id_transporte,
        transporte["estado"],
        estado_nuevo,
        f"Estado del transporte actualizado a {estado_nuevo}."
    )
    notificar_participantes_operacion(
        transporte["id_operacion"],
        "Seguimiento de entrega actualizado",
        f"El transporte pasó al estado {estado_nuevo}.",
        "INFO"
    )
    flash("Estado del transporte actualizado.", "success")
    return redirect(url_for("mis_transportes"))


@app.route("/admin/transportes/<int:id_operacion>/asignar", methods=["POST"])
@login_required
def asignar_transporte(id_operacion):

    if not tiene_permiso("TRANSPORTE_ASIGNAR"):

        flash(
            "No tienes permisos para asignar transportes.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    operacion = query(
        """
        SELECT o.*, l.id_lote, l.codigo_lote, p.nombre AS producto,
            (
                SELECT cl.peso_kg
                FROM cotizaciones_logisticas cl
                WHERE cl.id_cotizacion_operacion = o.id_cotizacion_operacion
                ORDER BY cl.id_cotizacion DESC
                LIMIT 1
            ) AS peso_transporte
        FROM operaciones o
        INNER JOIN lotes l
            ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        WHERE o.id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not operacion:
        flash("La operación no existe.", "danger")
        return redirect(url_for("admin_transportes"))

    if query(
        """
        SELECT id_transporte
        FROM transportes
        WHERE id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    ):
        flash("Ya existe un transporte asignado para esta operación.", "warning")
        return redirect(url_for("admin_transportes"))

    id_transportista = request.form.get("id_transportista", type=int)
    id_vehiculo = request.form.get("id_vehiculo", type=int)
    origen = request.form.get("origen") or "Predio Agrovia"
    destino = request.form.get("destino") or "Dirección de entrega del comprador"
    observaciones = request.form.get("observaciones") or "Asignado por administración."

    if not id_transportista or not id_vehiculo:
        flash("Debes seleccionar transportista y vehículo.", "warning")
        return redirect(url_for("admin_transportes"))

    codigo_transporte = generar_codigo("TR", "transportes", "codigo_transporte")
    conn = get_connection()
    cur = conn.cursor(dictionary=True)
    try:
        conn.start_transaction()
        cur.execute(
            "SELECT id_operacion FROM operaciones WHERE id_operacion = %s FOR UPDATE",
            (id_operacion,)
        )
        if not cur.fetchone():
            conn.rollback()
            flash("La operación ya no está disponible para asignar transporte.", "warning")
            return redirect(url_for("admin_transportes"))

        cur.execute(
            "SELECT id_transporte FROM transportes WHERE id_operacion = %s LIMIT 1 FOR UPDATE",
            (id_operacion,)
        )
        if cur.fetchone():
            conn.rollback()
            flash("Ya existe un transporte asignado para esta operación.", "warning")
            return redirect(url_for("admin_transportes"))

        cur.execute(
            """
            SELECT v.id_vehiculo, v.capacidad_kg
            FROM vehiculos v
            INNER JOIN transportistas t
                ON t.id_transportista = v.id_transportista
            WHERE v.id_vehiculo = %s
              AND v.id_transportista = %s
              AND v.estado = 'ACTIVO'
              AND v.disponibilidad = 'DISPONIBLE'
              AND t.estado = 'ACTIVO'
            LIMIT 1
            FOR UPDATE
            """,
            (id_vehiculo, id_transportista)
        )
        vehiculo = cur.fetchone()
        if not vehiculo:
            conn.rollback()
            flash(
                "El vehículo debe estar activo, disponible y pertenecer al transportista seleccionado.",
                "warning"
            )
            return redirect(url_for("admin_transportes"))
        if operacion.get("peso_transporte") is not None and (
            Decimal(str(vehiculo["capacidad_kg"])) < Decimal(str(operacion["peso_transporte"]))
        ):
            conn.rollback()
            flash("La capacidad del vehículo es menor que la carga confirmada para esta operación.", "warning")
            return redirect(url_for("admin_transportes"))

        cur.execute(
            """
            INSERT INTO transportes
            (
                codigo_transporte, id_operacion, id_transportista, id_vehiculo,
                fecha_asignacion, origen, destino, estado, observaciones
            )
            VALUES (%s, %s, %s, %s, NOW(), %s, %s, 'ASIGNADO', %s)
            """,
            (
                codigo_transporte, id_operacion, id_transportista, id_vehiculo,
                origen, destino, observaciones
            )
        )
        cur.execute(
            """
            UPDATE operaciones
            SET estado = 'EN_TRANSPORTE',
                fecha_actualizacion = NOW()
            WHERE id_operacion = %s
            """,
            (id_operacion,)
        )
        cur.execute(
            "UPDATE lotes SET estado = 'EN_TRANSPORTE' WHERE id_lote = %s",
            (operacion["id_lote"],)
        )
        cur.execute(
            """
            UPDATE vehiculos
            SET disponibilidad = 'ASIGNADO'
            WHERE id_vehiculo = %s
              AND id_transportista = %s
              AND estado = 'ACTIVO'
              AND disponibilidad = 'DISPONIBLE'
            """,
            (id_vehiculo, id_transportista)
        )
        if cur.rowcount != 1:
            raise RuntimeError("No se pudo reservar el vehículo seleccionado.")
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()

    registrar_historial_estado(
        "OPERACION", id_operacion, operacion["estado"], "EN_TRANSPORTE",
        "Transporte asignado a la operación."
    )
    notificar_participantes_operacion(
        id_operacion,
        "Transporte asignado",
        f"Se asignó transporte a la operación {operacion['codigo_operacion']}; "
        "el lote está en camino al comprador.",
        "INFO"
    )

    flash("Transporte asignado correctamente.", "success")

    return redirect(url_for("admin_transportes"))


@app.route("/admin/entregas")
@login_required
def admin_entregas():

    if not es_admin():

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    entregas = query(
        """
        SELECT
            t.id_transporte,
            t.codigo_transporte,
            t.estado AS estado_transporte,
            de.direccion AS direccion_entrega,
            de.latitud AS latitud_destino,
            de.longitud AS longitud_destino,
            (
                SELECT ut.latitud
                FROM ubicaciones_transporte ut
                WHERE ut.id_transporte = t.id_transporte
                ORDER BY ut.id_ubicacion DESC
                LIMIT 1
            ) AS latitud_actual,
            (
                SELECT ut.longitud
                FROM ubicaciones_transporte ut
                WHERE ut.id_transporte = t.id_transporte
                ORDER BY ut.id_ubicacion DESC
                LIMIT 1
            ) AS longitud_actual,
            (
                SELECT ut.fecha_reporte
                FROM ubicaciones_transporte ut
                WHERE ut.id_transporte = t.id_transporte
                ORDER BY ut.id_ubicacion DESC
                LIMIT 1
            ) AS fecha_ubicacion,
            (
                SELECT ut.fecha_reporte >= DATE_SUB(NOW(), INTERVAL 15 MINUTE)
                FROM ubicaciones_transporte ut
                WHERE ut.id_transporte = t.id_transporte
                ORDER BY ut.id_ubicacion DESC
                LIMIT 1
            ) AS ubicacion_reciente,
            o.codigo_operacion,
            o.total_operacion,
            p.nombre AS producto,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            tr.nombre_empresa,
            v.placa,
            e.id_entrega,
            e.codigo_entrega,
            e.estado AS estado_entrega,
            e.fecha_hora,
            e.nombre_receptor,
            e.conforme

        FROM transportes t

        INNER JOIN operaciones o
            ON o.id_operacion = t.id_operacion

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        INNER JOIN transportistas tr
            ON tr.id_transportista = t.id_transportista

        INNER JOIN vehiculos v
            ON v.id_vehiculo = t.id_vehiculo

        LEFT JOIN entregas e
            ON e.id_transporte = t.id_transporte

        ORDER BY t.fecha_asignacion DESC
        """
    )

    return render_template(
        "admin_entregas.html",
        entregas=entregas
    )


@app.route(
    "/admin/transportes/<int:id_transporte>/verificar-llegada",
    methods=["POST"]
)
@login_required
def verificar_llegada_transporte(id_transporte):
    if not es_admin():
        flash("Solo un administrador puede verificar la llegada del transporte.", "danger")
        return redirect(url_for("dashboard"))

    conn = get_connection()
    cur = conn.cursor(dictionary=True)
    try:
        conn.start_transaction()
        cur.execute(
            """
            SELECT id_operacion, estado
            FROM transportes
            WHERE id_transporte = %s
            LIMIT 1
            FOR UPDATE
            """,
            (id_transporte,)
        )
        transporte = cur.fetchone()
        if not transporte:
            conn.rollback()
            flash("El transporte no existe.", "warning")
            return redirect(url_for("admin_entregas"))
        if transporte["estado"] != "EN_RUTA":
            conn.rollback()
            flash("Solo se puede verificar la llegada de un transporte en ruta.", "warning")
            return redirect(url_for("admin_entregas"))

        cur.execute(
            """
            SELECT fecha_reporte
            FROM ubicaciones_transporte
            WHERE id_transporte = %s
              AND fecha_reporte >= DATE_SUB(NOW(), INTERVAL 15 MINUTE)
            ORDER BY id_ubicacion DESC
            LIMIT 1
            """,
            (id_transporte,)
        )
        if not cur.fetchone():
            conn.rollback()
            flash(
                "Se requiere una ubicación GPS actualizada en los últimos 15 minutos para verificar la llegada.",
                "warning"
            )
            return redirect(url_for("admin_entregas"))

        cur.execute(
            """
            UPDATE transportes
            SET estado = 'LLEGADA_DESTINO',
                fecha_llegada = NOW()
            WHERE id_transporte = %s
              AND estado = 'EN_RUTA'
            """,
            (id_transporte,)
        )
        if cur.rowcount != 1:
            conn.rollback()
            flash("El estado del transporte cambió; vuelve a revisar su ubicación.", "warning")
            return redirect(url_for("admin_entregas"))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()

    registrar_historial_estado(
        "TRANSPORTE",
        id_transporte,
        "EN_RUTA",
        "LLEGADA_DESTINO",
        "Llegada verificada por administración usando la ubicación GPS del transporte."
    )
    notificar_participantes_operacion(
        transporte["id_operacion"],
        "Llegada del transporte verificada",
        "Administración verificó la llegada del transporte mediante su ubicación GPS.",
        "SUCCESS"
    )
    flash("Llegada del transporte verificada con su ubicación GPS actual.", "success")
    return redirect(url_for("admin_entregas"))


@app.route("/admin/entregas/<int:id_transporte>/confirmar", methods=["POST"])
@login_required
def confirmar_entrega(id_transporte):

    if not es_admin():

        flash(
            "No tienes permisos para confirmar entregas.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    transporte = query(
        """
         SELECT t.*, o.id_operacion, o.codigo_operacion, o.total_operacion,
             o.estado AS estado_operacion, l.id_lote
        FROM transportes t
        INNER JOIN operaciones o
            ON o.id_operacion = t.id_operacion
        INNER JOIN lotes l
            ON l.id_lote = o.id_lote
        WHERE t.id_transporte = %s
        LIMIT 1
        """,
        (id_transporte,),
        one=True
    )

    if not transporte:
        flash("El transporte no existe.", "danger")
        return redirect(url_for("admin_entregas"))

    if query(
        """
        SELECT id_entrega
        FROM entregas
        WHERE id_transporte = %s
        LIMIT 1
        """,
        (id_transporte,),
        one=True
    ):
        flash("Ya existe una entrega registrada para este transporte.", "warning")
        return redirect(url_for("admin_entregas"))

    codigo_entrega = generar_codigo("ENT", "entregas", "codigo_entrega")
    nombre_receptor = request.form.get("nombre_receptor") or "Recepción administrativa"
    documento_receptor = request.form.get("documento_receptor") or "N/A"
    observaciones = request.form.get("observaciones") or "Entrega confirmada por administración."

    execute(
        """
        INSERT INTO entregas
        (
            codigo_entrega,
            id_operacion,
            id_transporte,
            fecha_hora,
            nombre_receptor,
            documento_receptor,
            cantidad_entregada,
            conforme,
            id_usuario_confirma,
            observaciones,
            estado
        )
        VALUES
        (
            %s,
            %s,
            %s,
            NOW(),
            %s,
            %s,
            0,
            1,
            %s,
            %s,
            'CONFORME'
        )
        """,
        (
            codigo_entrega,
            transporte["id_operacion"],
            id_transporte,
            nombre_receptor,
            documento_receptor,
            session["user_id"],
            observaciones
        )
    )

    execute(
        """
        UPDATE transportes
        SET estado = 'ENTREGADO',
            fecha_llegada = COALESCE(fecha_llegada, NOW())
        WHERE id_transporte = %s
        """,
        (id_transporte,)
    )

    execute(
        """
        UPDATE operaciones
        SET estado = 'ENTREGADA',
            fecha_actualizacion = NOW()
        WHERE id_operacion = %s
        """,
        (transporte["id_operacion"],)
    )

    registrar_historial_estado(
        "OPERACION", transporte["id_operacion"], transporte["estado_operacion"],
        "ENTREGADA", "Entrega confirmada por administración."
    )
    notificar_participantes_operacion(
        transporte["id_operacion"],
        "Entrega confirmada",
        f"La entrega de la operación {transporte['codigo_operacion']} fue confirmada.",
        "SUCCESS"
    )

    execute(
        """
        UPDATE lotes
        SET estado = 'ENTREGADO'
        WHERE id_lote = %s
        """,
        (transporte["id_lote"],)
    )

    flash("Entrega confirmada correctamente.", "success")

    return redirect(url_for("admin_entregas"))


@app.route("/admin/liquidaciones")
@login_required
def admin_liquidaciones():

    if not tiene_permiso("PAGOS_VER"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    liquidaciones = query(
        """
        SELECT
            o.id_operacion,
            o.codigo_operacion,
            o.total_operacion,
            o.estado AS estado_operacion,
            p.nombre AS producto,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            lq.id_liquidacion,
            lq.codigo_liquidacion,
            lq.monto,
            lq.estado AS estado_liquidacion,
            lq.fecha_programada,
            op.id_orden_pago,
            op.codigo_orden_pago,
            op.metodo_pago,
            op.estado AS estado_orden_pago,
            c.nombre AS cuadrilla,
            t.nombre_empresa AS transportista

        FROM operaciones o

        INNER JOIN lotes lo
            ON lo.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = lo.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        LEFT JOIN liquidaciones lq
            ON lq.id_operacion = o.id_operacion

        LEFT JOIN ordenes_pago op
            ON op.id_liquidacion = lq.id_liquidacion

        LEFT JOIN cuadrillas c
            ON c.id_cuadrilla = lq.id_cuadrilla_beneficiaria

        LEFT JOIN transportistas t
            ON t.id_transportista = lq.id_transportista_beneficiario

        WHERE o.estado IN ('ENTREGADA', 'COMPLETADA')
        ORDER BY o.fecha_operacion DESC
        """
    )

    return render_template(
        "admin_liquidaciones.html",
        liquidaciones=liquidaciones
    )


@app.route("/admin/liquidaciones/<int:id_operacion>/generar", methods=["POST"])
@login_required
def generar_liquidacion(id_operacion):

    if not tiene_permiso("PAGOS_AUTORIZAR"):

        flash(
            "No tienes permisos para generar liquidaciones.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    operacion = query(
        """
        SELECT o.*, l.id_lote, p.nombre AS producto, pr.id_productor
        FROM operaciones o
        INNER JOIN lotes l
            ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor
        WHERE o.id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not operacion:
        flash("La operación no existe.", "danger")
        return redirect(url_for("admin_liquidaciones"))

    if query(
        """
        SELECT id_liquidacion
        FROM liquidaciones
        WHERE id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    ):
        flash("Ya existe una liquidación para esta operación.", "warning")
        return redirect(url_for("admin_liquidaciones"))

    metodo_pago = request.form.get("metodo_pago", "TRANSFERENCIA_BANCARIA")
    metodos_pago = {
        "TRANSFERENCIA_BANCARIA",
        "BILLETERA_DIGITAL",
        "EFECTIVO_PRESENCIAL"
    }
    if metodo_pago not in metodos_pago:
        flash("Selecciona un método de pago válido.", "warning")
        return redirect(url_for("admin_liquidaciones"))

    codigo_liquidacion = generar_codigo("LIQ", "liquidaciones", "codigo_liquidacion")

    conn = get_connection()
    cur = conn.cursor()
    try:
        conn.start_transaction()
        cur.execute(
            """
            INSERT INTO liquidaciones
            (
                codigo_liquidacion,
                id_operacion,
                id_productor_beneficiario,
                beneficiario_agrovia,
                concepto,
                monto,
                metodo,
                estado,
                fecha_programada
            )
            VALUES (%s, %s, %s, 0, 'Pago por operación completada', %s, %s, 'PENDIENTE', NOW())
            """,
            (
                codigo_liquidacion,
                id_operacion,
                operacion["id_productor"],
                operacion["total_operacion"],
                metodo_pago
            )
        )
        id_liquidacion = cur.lastrowid
        cur.execute(
            """
            INSERT INTO ordenes_pago
            (
                codigo_orden_pago,
                id_liquidacion,
                metodo_pago,
                monto,
                id_usuario_solicita,
                estado
            )
            VALUES (%s, %s, %s, %s, %s, 'GENERADA')
            """,
            (
                generar_codigo("OPG", "ordenes_pago", "codigo_orden_pago"),
                id_liquidacion,
                metodo_pago,
                operacion["total_operacion"],
                session["user_id"]
            )
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()

    execute(
        """
        UPDATE operaciones
        SET estado = 'COMPLETADA',
            fecha_actualizacion = NOW()
        WHERE id_operacion = %s
        """,
        (id_operacion,)
    )

    registrar_historial_estado(
        "OPERACION", id_operacion, operacion["estado"], "COMPLETADA",
        "Liquidación generada y operación completada."
    )
    notificar_participantes_operacion(
        id_operacion,
        "Operación completada",
        f"La operación {operacion['codigo_operacion']} se completó y se generó su liquidación.",
        "SUCCESS"
    )
    acreditar_puntos_operacion_calificada(id_operacion)

    execute(
        """
        UPDATE lotes
        SET estado = 'COMPLETADO'
        WHERE id_lote = %s
        """,
        (operacion["id_lote"],)
    )

    flash("Liquidación generada correctamente.", "success")

    return redirect(url_for("admin_liquidaciones"))


@app.route("/admin/ordenes-pago/<int:id_orden_pago>/autorizar", methods=["POST"])
@login_required
def autorizar_orden_pago(id_orden_pago):
    if not tiene_permiso("PAGOS_AUTORIZAR"):
        flash("No tienes permisos para autorizar órdenes de pago.", "danger")
        return redirect(url_for("dashboard"))

    orden = query(
        """
        SELECT op.*, l.id_operacion
        FROM ordenes_pago op
        INNER JOIN liquidaciones l ON l.id_liquidacion = op.id_liquidacion
        WHERE op.id_orden_pago = %s
        LIMIT 1
        """,
        (id_orden_pago,),
        one=True
    )
    if not orden or orden["estado"] != "GENERADA":
        flash("La orden no existe o ya no está pendiente de autorización.", "warning")
        return redirect(url_for("admin_liquidaciones"))

    conn = get_connection()
    cur = conn.cursor()
    try:
        conn.start_transaction()
        cur.execute(
            """
            UPDATE ordenes_pago
            SET estado = 'AUTORIZADA',
                id_usuario_autoriza = %s,
                fecha_autorizacion = NOW()
            WHERE id_orden_pago = %s
              AND estado = 'GENERADA'
            """,
            (session["user_id"], id_orden_pago)
        )
        if cur.rowcount != 1:
            conn.rollback()
            flash("La orden ya fue procesada por otro usuario.", "warning")
            return redirect(url_for("admin_liquidaciones"))
        cur.execute(
            "UPDATE liquidaciones SET estado = 'EN_PROCESO' WHERE id_liquidacion = %s",
            (orden["id_liquidacion"],)
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()
    notificar_participantes_operacion(
        orden["id_operacion"],
        "Orden de pago autorizada",
        f"La orden de pago de la operación {orden['id_operacion']} fue autorizada.",
        "INFO"
    )
    flash("Orden de pago autorizada.", "success")
    return redirect(url_for("admin_liquidaciones"))


@app.route("/admin/ordenes-pago/<int:id_orden_pago>/ejecutar", methods=["POST"])
@login_required
def ejecutar_orden_pago(id_orden_pago):
    if not tiene_permiso("PAGOS_EJECUTAR"):
        flash("No tienes permisos para ejecutar órdenes de pago.", "danger")
        return redirect(url_for("dashboard"))

    orden = query(
        "SELECT * FROM ordenes_pago WHERE id_orden_pago = %s LIMIT 1",
        (id_orden_pago,),
        one=True
    )
    if not orden or orden["estado"] != "AUTORIZADA":
        flash("La orden no existe o aún no está autorizada.", "warning")
        return redirect(url_for("admin_liquidaciones"))

    if orden["metodo_pago"] == "EFECTIVO_PRESENCIAL":
        fecha_entrega = request.form.get("fecha_entrega", "").strip()
        hora_entrega = request.form.get("hora_entrega", "").strip()
        lugar_entrega = request.form.get("lugar_entrega", "").strip()
        nombre_receptor = request.form.get("nombre_receptor", "").strip()
        documento_receptor = request.form.get("documento_receptor", "").strip()
        if not all((fecha_entrega, hora_entrega, lugar_entrega, nombre_receptor, documento_receptor)):
            flash("Completa los datos de entrega en efectivo.", "warning")
            return redirect(url_for("admin_liquidaciones"))
    conn = get_connection()
    cur = conn.cursor()
    try:
        conn.start_transaction()
        cur.execute(
            """
            UPDATE ordenes_pago
            SET estado = 'PAGADA',
                id_usuario_ejecuta = %s,
                fecha_ejecucion = NOW()
            WHERE id_orden_pago = %s
              AND estado = 'AUTORIZADA'
            """,
            (session["user_id"], id_orden_pago)
        )
        if cur.rowcount != 1:
            conn.rollback()
            flash("La orden ya fue procesada por otro usuario.", "warning")
            return redirect(url_for("admin_liquidaciones"))

        if orden["metodo_pago"] == "EFECTIVO_PRESENCIAL":
            cur.execute(
                """
                INSERT INTO pagos_efectivo
                (
                    id_orden_pago, monto_entregado, fecha_entrega,
                    hora_entrega, lugar_entrega, id_usuario_entrega,
                    nombre_receptor, documento_receptor, conforme
                )
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, 1)
                """,
                (
                    id_orden_pago,
                    orden["monto"],
                    fecha_entrega,
                    hora_entrega,
                    lugar_entrega,
                    session["user_id"],
                    nombre_receptor,
                    documento_receptor
                )
            )

        cur.execute(
            """
            UPDATE liquidaciones
            SET estado = 'PAGADA', fecha_pago = NOW()
            WHERE id_liquidacion = %s
            """,
            (orden["id_liquidacion"],)
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()
    operacion_liquidacion = query(
        """
        SELECT id_operacion
        FROM liquidaciones
        WHERE id_liquidacion = %s
        LIMIT 1
        """,
        (orden["id_liquidacion"],),
        one=True
    )
    if operacion_liquidacion:
        notificar_participantes_operacion(
            operacion_liquidacion["id_operacion"],
            "Pago de liquidación ejecutado",
            f"Se ejecutó el pago correspondiente a la operación "
            f"{operacion_liquidacion['id_operacion']}.",
            "SUCCESS"
        )
    flash("Orden de pago ejecutada y liquidación pagada.", "success")
    return redirect(url_for("admin_liquidaciones"))


@app.route("/admin/incidencias")
@login_required
def admin_incidencias():

    if not es_admin():

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    incidencias = query(
        """
        SELECT
            o.id_operacion,
            o.codigo_operacion,
            p.nombre AS producto,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            i.id_incidencia,
            i.codigo_incidencia,
            i.categoria,
            i.descripcion,
            i.prioridad,
            i.estado,
            i.fecha_apertura

        FROM operaciones o

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        LEFT JOIN incidencias i
            ON i.id_operacion = o.id_operacion

        ORDER BY i.fecha_apertura DESC, o.fecha_operacion DESC
        """
    )

    return render_template(
        "admin_incidencias.html",
        incidencias=incidencias
    )


@app.route("/admin/incidencias/<int:id_operacion>/reportar", methods=["POST"])
@login_required
def reportar_incidencia(id_operacion):

    if not es_admin():

        flash(
            "No tienes permisos para reportar incidencias.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    operacion = query(
        """
        SELECT o.*, p.nombre AS producto
        FROM operaciones o
        INNER JOIN lotes l
            ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        WHERE o.id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not operacion:
        flash("La operación no existe.", "danger")
        return redirect(url_for("admin_incidencias"))

    categoria = request.form.get("categoria") or "OPERACIONAL"
    descripcion = request.form.get("descripcion") or "Incidencia reportada por administración."
    prioridad = request.form.get("prioridad") or "MEDIA"

    codigo_incidencia = generar_codigo("INC", "incidencias", "codigo_incidencia")

    id_incidencia = execute(
        """
        INSERT INTO incidencias
        (
            codigo_incidencia,
            id_usuario_reporta,
            id_operacion,
            categoria,
            descripcion,
            prioridad,
            estado,
            id_responsable,
            fecha_apertura
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            %s,
            'ABIERTO',
            %s,
            NOW()
        )
        """,
        (
            codigo_incidencia,
            session["user_id"],
            id_operacion,
            categoria,
            descripcion,
            prioridad,
            session["user_id"]
        )
    )

    registrar_historial_estado(
        "INCIDENCIA",
        id_incidencia,
        None,
        "ABIERTO",
        descripcion
    )
    notificar_participantes_incidencia(
        id_incidencia,
        "Incidencia registrada",
        f"Se registró la incidencia {codigo_incidencia} para la operación "
        f"{operacion['codigo_operacion']}.",
        "WARNING"
    )

    flash("Incidencia registrada correctamente.", "success")

    return redirect(url_for("admin_incidencias"))


@app.route("/admin/incidencias/<int:id_incidencia>/estado", methods=["POST"])
@login_required
def actualizar_estado_incidencia(id_incidencia):
    if not es_admin():
        flash("No tienes permisos para actualizar incidencias.", "danger")
        return redirect(url_for("dashboard"))

    incidencia = query(
        """
        SELECT i.id_incidencia, i.id_operacion, i.codigo_incidencia, i.estado,
               o.codigo_operacion
        FROM incidencias i
        LEFT JOIN operaciones o ON o.id_operacion = i.id_operacion
        WHERE i.id_incidencia = %s
        LIMIT 1
        """,
        (id_incidencia,),
        one=True
    )
    if not incidencia:
        flash("La incidencia no existe.", "warning")
        return redirect(url_for("admin_incidencias"))
    if incidencia["estado"] == "CERRADO":
        flash("Una incidencia cerrada ya no puede cambiar de estado.", "warning")
        return redirect(url_for("admin_incidencias"))

    nuevo_estado = request.form.get("estado", "")
    estados_permitidos = {"EN_ATENCION", "RESUELTO", "CERRADO"}
    if nuevo_estado not in estados_permitidos:
        flash("Selecciona un estado válido para la incidencia.", "warning")
        return redirect(url_for("admin_incidencias"))
    if nuevo_estado == incidencia["estado"]:
        flash("La incidencia ya tiene ese estado.", "info")
        return redirect(url_for("admin_incidencias"))

    solucion = request.form.get("solucion", "").strip() or None
    conn = get_connection()
    cur = conn.cursor()
    try:
        cur.execute(
            """
            UPDATE incidencias
            SET estado = %s,
                fecha_cierre = CASE
                    WHEN %s = 'CERRADO' THEN NOW()
                    ELSE NULL
                END,
                solucion = COALESCE(%s, solucion)
            WHERE id_incidencia = %s
              AND estado = %s
            """,
            (
                nuevo_estado,
                nuevo_estado,
                solucion,
                id_incidencia,
                incidencia["estado"]
            )
        )
        if cur.rowcount != 1:
            conn.rollback()
            flash("La incidencia cambió mientras la actualizabas. Recarga la página.", "warning")
            return redirect(url_for("admin_incidencias"))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        conn.close()
    registrar_historial_estado(
        "INCIDENCIA",
        id_incidencia,
        incidencia["estado"],
        nuevo_estado,
        solucion or f"Estado actualizado a {nuevo_estado}."
    )
    notificar_participantes_incidencia(
        id_incidencia,
        "Estado de incidencia actualizado",
        f"La incidencia {incidencia['codigo_incidencia']} de la operación "
        f"{incidencia['codigo_operacion'] or incidencia['id_operacion']} cambió a {nuevo_estado}."
        + (f" Solución: {solucion}" if solucion else ""),
        "SUCCESS" if nuevo_estado in {"RESUELTO", "CERRADO"} else "INFO"
    )
    flash("Estado de la incidencia actualizado y notificado a los participantes.", "success")
    return redirect(url_for("admin_incidencias"))


@app.route("/admin/devoluciones")
@login_required
def admin_devoluciones():

    if not es_admin():

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    devoluciones = query(
        """
        SELECT
            d.id_devolucion,
            d.codigo_devolucion,
            d.monto,
            d.motivo,
            d.metodo,
            d.estado,
            d.fecha,
            o.id_operacion,
            o.codigo_operacion,
            p.nombre AS producto,
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador

        FROM devoluciones d

        INNER JOIN operaciones o
            ON o.id_operacion = d.id_operacion

        INNER JOIN lotes l
            ON l.id_lote = o.id_lote

        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud

        INNER JOIN productos p
            ON p.id_producto = s.id_producto

        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor

        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario

        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador

        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario

        ORDER BY d.fecha DESC
        """
    )

    return render_template(
        "admin_devoluciones.html",
        devoluciones=devoluciones
    )


@app.route("/admin/devoluciones/<int:id_operacion>/solicitar", methods=["POST"])
@login_required
def solicitar_devolucion(id_operacion):

    if not es_admin():

        flash(
            "No tienes permisos para gestionar devoluciones.",
            "danger"
        )

        return redirect(url_for("dashboard"))

    operacion = query(
        """
        SELECT o.*, p.nombre AS producto
        FROM operaciones o
        INNER JOIN lotes l
            ON l.id_lote = o.id_lote
        INNER JOIN solicitudes_lote s
            ON s.id_solicitud = l.id_solicitud
        INNER JOIN productos p
            ON p.id_producto = s.id_producto
        WHERE o.id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not operacion:
        flash("La operación no existe.", "danger")
        return redirect(url_for("admin_devoluciones"))

    pago = query(
        """
        SELECT *
        FROM pagos
        WHERE id_operacion = %s
        ORDER BY id_pago DESC
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not pago:
        flash("No existe un pago asociado para registrar una devolución.", "warning")
        return redirect(url_for("admin_devoluciones"))

    if query(
        """
        SELECT id_devolucion
        FROM devoluciones
        WHERE id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    ):
        flash("Ya existe una devolución registrada para esta operación.", "warning")
        return redirect(url_for("admin_devoluciones"))

    motivo = request.form.get("motivo") or "Devolución por ajuste o incidencia registrada."
    monto = request.form.get("monto")

    try:
        monto = float(monto) if monto else round(float(pago["monto"]) * 0.10, 2)
    except ValueError:
        flash("El monto de la devolución no es válido.", "warning")
        return redirect(url_for("admin_devoluciones"))

    codigo_devolucion = generar_codigo("DEV", "devoluciones", "codigo_devolucion")

    execute(
        """
        INSERT INTO devoluciones
        (
            codigo_devolucion,
            id_operacion,
            id_pago,
            motivo,
            monto,
            metodo,
            referencia,
            id_usuario_autoriza,
            id_usuario_ejecuta,
            fecha,
            estado
        )
        VALUES
        (
            %s,
            %s,
            %s,
            %s,
            %s,
            'TRANSFERENCIA_BANCARIA',
            'DEVOLUCION_AUTOMATICA',
            %s,
            %s,
            NOW(),
            'AUTORIZADA'
        )
        """,
        (
            codigo_devolucion,
            id_operacion,
            pago["id_pago"],
            motivo,
            monto,
            session["user_id"],
            session["user_id"]
        )
    )

    flash("Devolución registrada correctamente.", "success")

    return redirect(url_for("admin_devoluciones"))


@app.route("/notificaciones")
@login_required
def notificaciones():

    rows = query(
        """
        SELECT *
        FROM notificaciones
        WHERE id_usuario = %s
        ORDER BY fecha DESC
        """,
        (session["user_id"],)
    )
    execute(
        """
        UPDATE notificaciones
        SET leida = 1
        WHERE id_usuario = %s
          AND leida = 0
        """,
        (session["user_id"],)
    )

    return render_template(
        "notificaciones.html",
        notificaciones=rows
    )


@app.route("/admin/notificaciones", methods=["GET", "POST"])
@login_required
def admin_notificaciones():

    if not es_admin():

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    if request.method == "POST":
        usuario_id = request.form.get("id_usuario")
        titulo = request.form.get("titulo", "").strip()
        mensaje = request.form.get("mensaje", "").strip()
        tipo = request.form.get("tipo", "INFO")

        if not usuario_id or not titulo or not mensaje:
            flash("Debes completar el usuario, título y mensaje.", "warning")
            return redirect(url_for("admin_notificaciones"))

        crear_notificacion(
            int(usuario_id),
            titulo,
            mensaje,
            tipo,
            "NOTIFICACION_ADMIN",
            None
        )

        registrar_auditoria(
            usuario_id=session["user_id"],
            rol="ADMINISTRADOR",
            accion="ENVIAR_NOTIFICACION",
            modulo="NOTIFICACIONES",
            entidad_tipo="notificaciones",
            entidad_id=None,
            valor_nuevo={
                "destino": int(usuario_id),
                "titulo": titulo,
                "tipo": tipo
            }
        )

        flash("Notificación enviada correctamente.", "success")
        return redirect(url_for("admin_notificaciones"))

    usuarios = query(
        """
        SELECT
            id_usuario,
            CONCAT(nombres, ' ', COALESCE(apellido_paterno, '')) AS nombre,
            correo AS email
        FROM usuarios
        ORDER BY nombres ASC
        """
    )

    notifs = query(
        """
        SELECT
            n.*,
            CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, '')) AS usuario
        FROM notificaciones n
        LEFT JOIN usuarios u
            ON u.id_usuario = n.id_usuario
        ORDER BY n.fecha DESC
        LIMIT 200
        """
    )

    return render_template(
        "admin_notificaciones.html",
        usuarios=usuarios,
        notificaciones=notifs
    )


REPORTES_DESCARGABLES = {
    "operaciones": {
        "titulo": "Operaciones",
        "columnas": (
            "Código", "Fecha", "Producto", "Productor", "Comprador",
            "Total (S/)", "Estado",
        ),
        "estados": (
            "PENDIENTE_PAGO", "PAGO_CONFIRMADO", "COSECHA_PROGRAMADA",
            "EN_COSECHA", "COSECHA_VALIDADA", "EN_TRANSPORTE",
            "ENTREGADA", "COMPLETADA", "CANCELADA",
        ),
    },
    "pagos": {
        "titulo": "Pagos",
        "columnas": (
            "Código de pago", "Fecha", "Operación", "Comprador",
            "Monto (S/)", "Método", "Estado",
        ),
        "estados": ("PENDIENTE", "RECIBIDO", "RETENIDO", "DEVUELTO", "ANULADO"),
    },
    "productos": {
        "titulo": "Productos",
        "columnas": (
            "Producto", "Descripción", "Estado", "Lotes publicados",
            "Lotes disponibles", "Producción mínima (kg)",
            "Producción máxima (kg)",
        ),
        "estados": ("ACTIVO", "INACTIVO"),
    },
    "productores": {
        "titulo": "Productores",
        "columnas": (
            "Código", "Productor", "Correo", "Verificación", "Estado",
            "Operaciones", "Total de operaciones (S/)",
        ),
        "estados": ("ACTIVO", "INACTIVO"),
    },
    "entregas": {
        "titulo": "Entregas",
        "columnas": (
            "Código de transporte", "Código de entrega", "Fecha", "Operación",
            "Producto", "Comprador", "Transportista", "Estado", "Receptor",
            "Conforme",
        ),
        "estados": (
            "PENDIENTE_CONFIRMACION", "CONFORME", "OBSERVADA", "RECHAZADA",
            "ASIGNADO", "RECOJO_PROGRAMADO", "CARGA_REALIZADA", "EN_RUTA",
            "LLEGADA_DESTINO", "ENTREGADO", "CANCELADO",
        ),
    },
}

MAX_FILAS_REPORTE = 50000


def obtener_datos_reporte(tipo, fecha_desde, fecha_hasta, estado):
    """Carga filas de un reporte usando solo tablas y filtros permitidos."""
    condicion_fecha = []
    parametros_fecha = []
    if fecha_desde:
        condicion_fecha.append("{fecha_columna} >= %s")
        parametros_fecha.append(fecha_desde)
    if fecha_hasta:
        condicion_fecha.append("{fecha_columna} < DATE_ADD(%s, INTERVAL 1 DAY)")
        parametros_fecha.append(fecha_hasta)

    if tipo == "operaciones":
        where = [
            clause.format(fecha_columna="o.fecha_operacion")
            for clause in condicion_fecha
        ]
        params = list(parametros_fecha)
        if estado:
            where.append("o.estado = %s")
            params.append(estado)
        sql = """
            SELECT o.codigo_operacion AS `Código`,
                   o.fecha_operacion AS `Fecha`,
                   p.nombre AS `Producto`,
                   CONCAT(up.nombres, ' ', COALESCE(up.apellido_paterno, '')) AS `Productor`,
                   CONCAT(uc.nombres, ' ', COALESCE(uc.apellido_paterno, '')) AS `Comprador`,
                   o.total_operacion AS `Total (S/)`,
                   o.estado AS `Estado`
            FROM operaciones o
            INNER JOIN lotes l ON l.id_lote = o.id_lote
            INNER JOIN solicitudes_lote s ON s.id_solicitud = l.id_solicitud
            INNER JOIN productos p ON p.id_producto = s.id_producto
            INNER JOIN productores pr ON pr.id_productor = o.id_productor
            INNER JOIN usuarios up ON up.id_usuario = pr.id_usuario
            INNER JOIN compradores co ON co.id_comprador = o.id_comprador
            INNER JOIN usuarios uc ON uc.id_usuario = co.id_usuario
        """
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY o.fecha_operacion DESC, o.id_operacion DESC"
    elif tipo == "pagos":
        where = [
            clause.format(fecha_columna="p.fecha")
            for clause in condicion_fecha
        ]
        params = list(parametros_fecha)
        if estado:
            where.append("p.estado = %s")
            params.append(estado)
        sql = """
            SELECT p.codigo_pago AS `Código de pago`,
                   p.fecha AS `Fecha`,
                   o.codigo_operacion AS `Operación`,
                   CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, '')) AS `Comprador`,
                   p.monto AS `Monto (S/)`,
                   p.metodo AS `Método`,
                   p.estado AS `Estado`
            FROM pagos p
            INNER JOIN operaciones o ON o.id_operacion = p.id_operacion
            INNER JOIN compradores co ON co.id_comprador = o.id_comprador
            INNER JOIN usuarios u ON u.id_usuario = co.id_usuario
        """
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY p.fecha DESC, p.id_pago DESC"
    elif tipo == "productos":
        where = []
        params = []
        lotes_fecha = [
            clause.format(fecha_columna="l.fecha_publicacion")
            for clause in condicion_fecha
        ]
        lotes_sql = """
            SELECT s.id_producto,
                   COUNT(*) AS `Lotes publicados`,
                   COALESCE(SUM(l.produccion_min), 0) AS `Producción mínima (kg)`,
                   COALESCE(SUM(l.produccion_max), 0) AS `Producción máxima (kg)`,
                   SUM(CASE WHEN l.estado = 'DISPONIBLE' THEN 1 ELSE 0 END) AS `Lotes disponibles`
            FROM lotes l
            INNER JOIN solicitudes_lote s ON s.id_solicitud = l.id_solicitud
            WHERE l.estado != 'CANCELADO'
        """
        if lotes_fecha:
            lotes_sql += " AND " + " AND ".join(lotes_fecha)
            params.extend(parametros_fecha)
        lotes_sql += " GROUP BY s.id_producto"
        if estado:
            where.append("p.estado = %s")
            params.append(estado)
        sql = f"""
            SELECT p.nombre AS `Producto`,
                   p.descripcion AS `Descripción`,
                   p.estado AS `Estado`,
                   COALESCE(lotes.`Lotes publicados`, 0) AS `Lotes publicados`,
                   COALESCE(lotes.`Lotes disponibles`, 0) AS `Lotes disponibles`,
                   COALESCE(lotes.`Producción mínima (kg)`, 0) AS `Producción mínima (kg)`,
                   COALESCE(lotes.`Producción máxima (kg)`, 0) AS `Producción máxima (kg)`
            FROM productos p
            LEFT JOIN ({lotes_sql}) lotes ON lotes.id_producto = p.id_producto
        """
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY p.nombre"
    elif tipo == "productores":
        where = []
        params = []
        operaciones_fecha = [
            clause.format(fecha_columna="o.fecha_operacion")
            for clause in condicion_fecha
        ]
        operaciones_join = " AND ".join(
            ["o.id_productor = pr.id_productor"] + operaciones_fecha
        )
        params.extend(parametros_fecha)
        if estado:
            where.append("pr.estado = %s")
            params.append(estado)
        sql = f"""
            SELECT pr.codigo_productor AS `Código`,
                   CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, '')) AS `Productor`,
                   u.correo AS `Correo`,
                   pr.estado_verificacion AS `Verificación`,
                   pr.estado AS `Estado`,
                   COUNT(o.id_operacion) AS `Operaciones`,
                   COALESCE(SUM(o.total_operacion), 0) AS `Total de operaciones (S/)`
            FROM productores pr
            INNER JOIN usuarios u ON u.id_usuario = pr.id_usuario
            LEFT JOIN operaciones o ON {operaciones_join}
        """
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " GROUP BY pr.id_productor, pr.codigo_productor, u.nombres, "
        sql += "u.apellido_paterno, u.correo, pr.estado_verificacion, pr.estado "
        sql += "ORDER BY u.nombres, u.apellido_paterno"
    else:
        where = [
            clause.format(
                fecha_columna="COALESCE(e.fecha_hora, t.fecha_asignacion)"
            )
            for clause in condicion_fecha
        ]
        params = list(parametros_fecha)
        if estado:
            where.append("COALESCE(e.estado, t.estado) = %s")
            params.append(estado)
        sql = """
            SELECT t.codigo_transporte AS `Código de transporte`,
                   e.codigo_entrega AS `Código de entrega`,
                   COALESCE(e.fecha_hora, t.fecha_asignacion) AS `Fecha`,
                   o.codigo_operacion AS `Operación`,
                   p.nombre AS `Producto`,
                   CONCAT(uc.nombres, ' ', COALESCE(uc.apellido_paterno, '')) AS `Comprador`,
                   tr.nombre_empresa AS `Transportista`,
                   COALESCE(e.estado, t.estado) AS `Estado`,
                   e.nombre_receptor AS `Receptor`,
                   e.conforme AS `Conforme`
            FROM transportes t
            INNER JOIN operaciones o ON o.id_operacion = t.id_operacion
            INNER JOIN lotes l ON l.id_lote = o.id_lote
            INNER JOIN solicitudes_lote s ON s.id_solicitud = l.id_solicitud
            INNER JOIN productos p ON p.id_producto = s.id_producto
            INNER JOIN compradores co ON co.id_comprador = o.id_comprador
            INNER JOIN usuarios uc ON uc.id_usuario = co.id_usuario
            LEFT JOIN transportistas tr ON tr.id_transportista = t.id_transportista
            LEFT JOIN (
                SELECT e1.*
                FROM entregas e1
                INNER JOIN (
                    SELECT id_transporte, MAX(id_entrega) AS id_entrega
                    FROM entregas
                    GROUP BY id_transporte
                ) ultima ON ultima.id_entrega = e1.id_entrega
            ) e ON e.id_transporte = t.id_transporte
        """
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY COALESCE(e.fecha_hora, t.fecha_asignacion) DESC"

    sql += f" LIMIT {MAX_FILAS_REPORTE + 1}"
    filas = query(sql, tuple(params))
    return list(filas[0].keys()) if filas else list(
        REPORTES_DESCARGABLES[tipo]["columnas"]
    ), filas


def generar_archivo_excel(titulo, columnas, filas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte"
    hoja.append([titulo])
    hoja.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(columnas), 1))
    hoja["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    hoja["A1"].fill = PatternFill("solid", fgColor="286B3B")
    hoja["A1"].alignment = Alignment(vertical="center")
    hoja.row_dimensions[1].height = 26
    hoja.append(columnas)
    for celda in hoja[2]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="318447")
        celda.alignment = Alignment(wrap_text=True, vertical="center")

    for fila in filas:
        valores = []
        for columna in columnas:
            valor = fila.get(columna)
            if isinstance(valor, Decimal):
                valor = float(valor)
            if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")):
                valor = "'" + valor
            valores.append(valor)
        hoja.append(valores)

    hoja.freeze_panes = "A3"
    if columnas:
        hoja.auto_filter.ref = f"A2:{get_column_letter(len(columnas))}{max(2, hoja.max_row)}"
        for indice, columna in enumerate(columnas, 1):
            longitudes = [len(str(columna))]
            longitudes.extend(
                len(str(hoja.cell(fila, indice).value or ""))
                for fila in range(3, min(hoja.max_row, 100) + 1)
            )
            longitud = max(longitudes)
            hoja.column_dimensions[get_column_letter(indice)].width = min(longitud + 2, 42)

    archivo = io.BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return archivo


def generar_archivo_pdf(titulo, columnas, filas):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    archivo = io.BytesIO()
    documento = SimpleDocTemplate(
        archivo,
        pagesize=landscape(letter),
        leftMargin=0.35 * inch,
        rightMargin=0.35 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
        title=titulo,
    )
    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle(
        "ReportCell",
        parent=estilos["BodyText"],
        fontName="Helvetica",
        fontSize=7,
        leading=8,
        alignment=TA_LEFT,
        spaceAfter=0,
    )
    estilo_titulo = estilos["Title"]
    datos = [[Paragraph(html.escape(str(columna)), estilo_celda) for columna in columnas]]
    for fila in filas:
        datos.append([
            Paragraph(
                html.escape(str(fila.get(columna)) if fila.get(columna) is not None else "-").replace("\n", "<br/>"),
                estilo_celda,
            )
            for columna in columnas
        ])
    ancho_disponible = landscape(letter)[0] - documento.leftMargin - documento.rightMargin
    tabla = LongTable(
        datos,
        colWidths=[ancho_disponible / max(len(columnas), 1)] * len(columnas),
        repeatRows=1,
        hAlign="LEFT",
    )
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#318447")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d9e2d9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f7f3")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    documento.build([
        Paragraph(html.escape(titulo), estilo_titulo),
        Spacer(1, 0.15 * inch),
        Paragraph(f"Filas exportadas: {len(filas)}", estilos["Normal"]),
        Spacer(1, 0.12 * inch),
        tabla,
    ])
    archivo.seek(0)
    return archivo


@app.route("/admin/reportes")
@login_required
def admin_reportes():
    if not tiene_permiso("REPORTES_VER"):
        flash("No tienes permisos para descargar reportes.", "danger")
        return redirect(url_for("dashboard"))

    tipo = request.args.get("tipo", "operaciones")
    if tipo not in REPORTES_DESCARGABLES:
        tipo = "operaciones"
    estado = request.args.get("estado", "")
    if estado not in REPORTES_DESCARGABLES[tipo]["estados"]:
        estado = ""
    return render_template(
        "admin_reportes.html",
        reportes=REPORTES_DESCARGABLES,
        tipo_seleccionado=tipo,
        fecha_desde=request.args.get("fecha_desde", ""),
        fecha_hasta=request.args.get("fecha_hasta", ""),
        estado_seleccionado=estado,
    )


@app.route("/admin/reportes/<tipo>/<formato>")
@login_required
def descargar_reporte(tipo, formato):
    if not tiene_permiso("REPORTES_VER"):
        flash("No tienes permisos para descargar reportes.", "danger")
        return redirect(url_for("dashboard"))
    if tipo not in REPORTES_DESCARGABLES or formato not in {"excel", "pdf"}:
        abort(404)

    try:
        fecha_desde = (
            date.fromisoformat(request.args["fecha_desde"])
            if request.args.get("fecha_desde") else None
        )
        fecha_hasta = (
            date.fromisoformat(request.args["fecha_hasta"])
            if request.args.get("fecha_hasta") else None
        )
    except ValueError:
        abort(400, description="Las fechas del reporte deben tener formato AAAA-MM-DD.")
    if fecha_desde and fecha_hasta and fecha_desde > fecha_hasta:
        abort(400, description="La fecha inicial no puede ser posterior a la fecha final.")

    estado = request.args.get("estado", "").strip()
    if estado and estado not in REPORTES_DESCARGABLES[tipo]["estados"]:
        abort(400, description="El estado seleccionado no es válido para este reporte.")

    columnas, filas = obtener_datos_reporte(tipo, fecha_desde, fecha_hasta, estado)
    if len(filas) > MAX_FILAS_REPORTE:
        abort(
            413,
            description=(
                f"El reporte supera el máximo de {MAX_FILAS_REPORTE} filas. "
                "Aplica un rango de fechas o un estado más específico."
            ),
        )

    titulo = f"Reporte de {REPORTES_DESCARGABLES[tipo]['titulo']}"
    if formato == "excel":
        archivo = generar_archivo_excel(titulo, columnas, filas)
        extension = "xlsx"
        mimetype = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        archivo = generar_archivo_pdf(titulo, columnas, filas)
        extension = "pdf"
        mimetype = "application/pdf"
    nombre = f"reporte_{tipo}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{extension}"
    return send_file(
        archivo,
        mimetype=mimetype,
        as_attachment=True,
        download_name=nombre,
    )


@app.route("/admin/auditoria")
@login_required
def admin_auditoria():

    if not tiene_permiso("AUDITORIA_VER"):

        flash(
            "Esta sección es solo para administradores.",
            "warning"
        )

        return redirect(url_for("dashboard"))

    registros = query(
        """
        SELECT
            a.*,
            CONCAT(u.nombres, ' ', COALESCE(u.apellido_paterno, '')) AS usuario
        FROM auditoria a
        LEFT JOIN usuarios u
            ON u.id_usuario = a.id_usuario
        ORDER BY a.fecha_hora DESC
        LIMIT 200
        """
    )

    return render_template(
        "admin_auditoria.html",
        registros=registros
    )


@app.route("/operaciones/<int:id_operacion>/calificar", methods=["GET", "POST"])
@login_required
def calificar_operacion(id_operacion):

    operacion = query(
        """
        SELECT
            o.*, 
            CONCAT(
                up.nombres,
                ' ',
                COALESCE(up.apellido_paterno, '')
            ) AS productor,
            CONCAT(
                uc.nombres,
                ' ',
                COALESCE(uc.apellido_paterno, '')
            ) AS comprador,
            pr.id_usuario AS id_usuario_productor,
            co.id_usuario AS id_usuario_comprador
        FROM operaciones o
        INNER JOIN productores pr
            ON pr.id_productor = o.id_productor
        INNER JOIN usuarios up
            ON up.id_usuario = pr.id_usuario
        INNER JOIN compradores co
            ON co.id_comprador = o.id_comprador
        INNER JOIN usuarios uc
            ON uc.id_usuario = co.id_usuario
        WHERE o.id_operacion = %s
        LIMIT 1
        """,
        (id_operacion,),
        one=True
    )

    if not operacion:
        flash("La operación no existe.", "danger")
        return redirect(url_for("dashboard"))

    es_productor = session["user_id"] == operacion["id_usuario_productor"]
    es_comprador = session["user_id"] == operacion["id_usuario_comprador"]

    if not (es_productor or es_comprador):
        flash("Solo el productor o comprador de la operación pueden calificarla.", "warning")
        return redirect(url_for("dashboard"))

    if operacion["estado"] != "COMPLETADA":
        flash("Solo se pueden calificar operaciones completadas.", "warning")
        return redirect(url_for("mis_operaciones"))

    ya_califico = query(
        """
        SELECT id_calificacion
        FROM calificaciones
        WHERE id_operacion = %s
          AND id_usuario_califica = %s
        LIMIT 1
        """,
        (id_operacion, session["user_id"]),
        one=True
    )
    if ya_califico and operacion["estado"] == "COMPLETADA":
        acreditar_puntos_operacion_calificada(id_operacion)

    if request.method == "POST":
        if ya_califico:
            flash("Ya registraste una calificación para esta operación.", "warning")
            return redirect(url_for("notificaciones"))

        puntuacion = request.form.get("puntuacion", "5")
        comentario = (request.form.get("comentario") or "").strip()

        try:
            puntuacion = int(puntuacion)
        except ValueError:
            flash("La puntuación debe ser un número válido.", "warning")
            return redirect(url_for("dashboard"))

        if puntuacion < 1 or puntuacion > 5:
            flash("La puntuación debe estar entre 1 y 5.", "warning")
            return redirect(url_for("dashboard"))

        id_usuario_calificado = operacion["id_usuario_comprador"] if es_productor else operacion["id_usuario_productor"]
        tipo = "PRODUCTOR_A_COMPRADOR" if es_productor else "COMPRADOR_A_PRODUCTOR"

        execute(
            """
            INSERT INTO calificaciones
            (
                id_operacion,
                id_usuario_califica,
                id_usuario_calificado,
                puntuacion,
                comentario,
                tipo,
                fecha,
                estado
            )
            VALUES
            (
                %s,
                %s,
                %s,
                %s,
                %s,
                %s,
                NOW(),
                'ACTIVO'
            )
            """,
            (
                id_operacion,
                session["user_id"],
                id_usuario_calificado,
                puntuacion,
                comentario or "Calificación realizada por el usuario.",
                tipo
            )
        )

        registrar_historial_estado(
            "OPERACION",
            id_operacion,
            operacion["estado"],
            operacion["estado"],
            f"Calificación {puntuacion}/5 registrada por {tipo.replace('_', ' ').lower()}."
        )

        crear_notificacion(
            id_usuario_calificado,
            "Nueva calificación recibida",
            f"Has recibido una calificación de {puntuacion} puntos en la operación {operacion['codigo_operacion']}.",
            "INFO",
            "calificaciones",
            id_operacion
        )

        registrar_auditoria(
            usuario_id=session["user_id"],
            rol=session.get("roles", ["USUARIO"])[0],
            accion="CALIFICAR_OPERACION",
            modulo="CALIFICACIONES",
            entidad_tipo="calificaciones",
            entidad_id=id_operacion,
            valor_nuevo={
                "operacion_id": id_operacion,
                "puntuacion": puntuacion,
                "tipo": tipo
            }
        )

        usuarios_acreditados = acreditar_puntos_operacion_calificada(id_operacion)
        if usuarios_acreditados:
            flash(
                f"Calificación registrada. Se acreditaron {PUNTOS_POR_OPERACION_CALIFICADA} puntos a cada participante.",
                "success"
            )
        else:
            flash(
                "Calificación registrada. Los puntos se acreditarán cuando la operación esté completada y ambos participantes hayan calificado.",
                "success"
            )
        return redirect(url_for("notificaciones"))

    return render_template(
        "calificar_operacion.html",
        operacion=operacion,
        ya_califico=bool(ya_califico),
        es_productor=es_productor,
        es_comprador=es_comprador
    )


# ==========================================================
# USUARIOS
# ==========================================================

@app.route("/usuarios")
@login_required
def usuarios():
    if not tiene_permiso("USUARIOS_VER"):
        flash("No tienes permiso para consultar usuarios.", "danger")
        return redirect(url_for("dashboard"))

    users = query(
        """
        SELECT
            u.id_usuario,
            u.codigo_usuario,
            u.nombres,
            u.apellido_paterno,
            u.numero_documento,
            u.correo,
            u.telefono,
            u.estado,

            GROUP_CONCAT(
                r.nombre
                ORDER BY r.nombre
                SEPARATOR ', '
            ) AS roles

        FROM usuarios u

        LEFT JOIN usuarios_roles ur
            ON ur.id_usuario = u.id_usuario
            AND ur.estado = 'ACTIVO'

        LEFT JOIN roles r
            ON r.id_rol = ur.id_rol

        GROUP BY u.id_usuario

        ORDER BY u.fecha_registro DESC
        """
    )

    return render_template(
        "usuarios.html",
        users=users
    )


# ==========================================================
# MI CUENTA
# ==========================================================

@app.route("/mi-cuenta", methods=["GET", "POST"])
@login_required
def cuenta():

    user = query(
        """
        SELECT *
        FROM usuarios
        WHERE id_usuario = %s
        """,
        (session["user_id"],),
        one=True
    )

    comprador = obtener_comprador_actual()
    if request.method == "POST":
        action = request.form.get("action")
        conn = get_connection()
        cur = conn.cursor()
        try:
            conn.start_transaction()
            if action == "guardar_facturacion" and comprador:
                tipo_documento = request.form.get("tipo_documento", "")
                numero_documento = request.form.get("numero_documento", "").strip()
                nombre_razon_social = request.form.get("nombre_razon_social", "").strip()
                direccion_fiscal = request.form.get("direccion_fiscal", "").strip()
                id_ubigeo = request.form.get("id_ubigeo", type=int)
                tipos_documento = {"DNI", "RUC", "CE", "PASAPORTE", "OTRO"}
                if tipo_documento not in tipos_documento or not numero_documento or not nombre_razon_social:
                    conn.rollback()
                    flash("Completa el tipo y número de documento, y el nombre o razón social.", "warning")
                    return redirect(url_for("cuenta"))
                if not documento_facturacion_valido(tipo_documento, numero_documento):
                    conn.rollback()
                    flash("El DNI debe tener 8 dígitos o el RUC 11 dígitos con dígito verificador válido.", "warning")
                    return redirect(url_for("cuenta"))
                if len(direccion_fiscal) > 255 or bool(direccion_fiscal) != bool(id_ubigeo):
                    conn.rollback()
                    flash("La dirección fiscal requiere seleccionar un ubigeo válido.", "warning")
                    return redirect(url_for("cuenta"))
                if id_ubigeo:
                    cur.execute(
                        """
                        SELECT id_ubigeo
                        FROM ubigeos
                        WHERE id_ubigeo = %s
                          AND estado = 'ACTIVO'
                        LIMIT 1
                        """,
                        (id_ubigeo,)
                    )
                    if not cur.fetchone():
                        conn.rollback()
                        flash("El ubigeo fiscal seleccionado no está activo.", "warning")
                        return redirect(url_for("cuenta"))

                cur.execute(
                    """
                    SELECT id_facturacion
                    FROM datos_facturacion
                    WHERE id_comprador = %s
                      AND estado = 'ACTIVO'
                      AND predeterminado = 1
                    LIMIT 1
                    """,
                    (comprador["id_comprador"],)
                )
                tiene_predeterminado = cur.fetchone()
                predeterminado = request.form.get("predeterminado") == "1" or not tiene_predeterminado
                if predeterminado:
                    cur.execute(
                        "UPDATE datos_facturacion SET predeterminado = 0 WHERE id_comprador = %s",
                        (comprador["id_comprador"],)
                    )
                cur.execute(
                    """
                    INSERT INTO datos_facturacion
                    (
                        id_comprador, tipo_documento, numero_documento,
                        nombre_razon_social, direccion_fiscal, id_ubigeo,
                        correo_facturacion, telefono, predeterminado
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        comprador["id_comprador"],
                        tipo_documento,
                        numero_documento,
                        nombre_razon_social,
                        direccion_fiscal or None,
                        id_ubigeo,
                        request.form.get("correo_facturacion", "").strip() or None,
                        request.form.get("telefono_facturacion", "").strip() or None,
                        int(predeterminado)
                    )
                )
                conn.commit()
                flash("Datos de facturación guardados.", "success")
                return redirect(url_for("cuenta"))

            if action == "guardar_direccion_entrega" and comprador:
                id_ubigeo = request.form.get("id_ubigeo", type=int)
                nombre_referencia = request.form.get("nombre_referencia", "").strip()
                direccion = request.form.get("direccion", "").strip()
                referencia = request.form.get("referencia", "").strip()
                if (
                    not id_ubigeo
                    or not nombre_referencia or len(nombre_referencia) > 120
                    or not direccion or len(direccion) > 255
                    or len(referencia) > 255
                ):
                    conn.rollback()
                    flash("Completa la dirección y selecciona un ubigeo válido.", "warning")
                    return redirect(url_for("cuenta"))

                cur.execute(
                    """
                    SELECT id_ubigeo
                    FROM ubigeos
                    WHERE id_ubigeo = %s
                      AND estado = 'ACTIVO'
                    LIMIT 1
                    """,
                    (id_ubigeo,)
                )
                if not cur.fetchone():
                    conn.rollback()
                    flash("El ubigeo de entrega seleccionado no está activo.", "warning")
                    return redirect(url_for("cuenta"))

                cur.execute(
                    """
                    SELECT id_direccion
                    FROM direcciones_entrega
                    WHERE id_comprador = %s
                      AND estado = 'ACTIVO'
                      AND predeterminada = 1
                    LIMIT 1
                    """,
                    (comprador["id_comprador"],)
                )
                tiene_predeterminada = cur.fetchone()
                predeterminada = (
                    request.form.get("predeterminada") == "1"
                    or not tiene_predeterminada
                )
                if predeterminada:
                    cur.execute(
                        "UPDATE direcciones_entrega SET predeterminada = 0 WHERE id_comprador = %s",
                        (comprador["id_comprador"],)
                    )
                cur.execute(
                    """
                    INSERT INTO direcciones_entrega
                    (
                        id_comprador, nombre_referencia, direccion,
                        referencia, id_ubigeo, predeterminada, estado
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, 'ACTIVO')
                    """,
                    (
                        comprador["id_comprador"],
                        nombre_referencia,
                        direccion,
                        referencia or None,
                        id_ubigeo,
                        int(predeterminada)
                    )
                )
                conn.commit()
                flash("Dirección de entrega guardada.", "success")
                return redirect(url_for("cuenta"))

            if action == "guardar_medio_pago":
                tipo_medio = request.form.get("tipo_medio", "")
                titular = request.form.get("titular", "").strip()
                numero_cuenta = request.form.get("numero_cuenta_enmascarado", "").strip()
                cci = request.form.get("cci_enmascarado", "").strip()
                tipos_medio = {
                    "TRANSFERENCIA_BANCARIA",
                    "BILLETERA_DIGITAL",
                    "EFECTIVO_PRESENCIAL"
                }
                if tipo_medio not in tipos_medio or not titular:
                    conn.rollback()
                    flash("Selecciona un método válido e indica el titular.", "warning")
                    return redirect(url_for("cuenta"))
                patron_enmascarado = r"[*xX\s-]*\d{2,4}"
                if any(
                    valor and not re.fullmatch(patron_enmascarado, valor)
                    for valor in (numero_cuenta, cci)
                ):
                    conn.rollback()
                    flash("Registra solo los últimos cuatro dígitos y enmascara el resto.", "warning")
                    return redirect(url_for("cuenta"))

                cur.execute(
                    """
                    SELECT id_medio_pago
                    FROM medios_pago_usuario
                    WHERE id_usuario = %s
                      AND estado = 'ACTIVO'
                      AND predeterminado = 1
                    LIMIT 1
                    """,
                    (session["user_id"],)
                )
                tiene_predeterminado = cur.fetchone()
                predeterminado = request.form.get("predeterminado") == "1" or not tiene_predeterminado
                if predeterminado:
                    cur.execute(
                        "UPDATE medios_pago_usuario SET predeterminado = 0 WHERE id_usuario = %s",
                        (session["user_id"],)
                    )
                cur.execute(
                    """
                    INSERT INTO medios_pago_usuario
                    (
                        id_usuario, tipo_medio, entidad, titular,
                        telefono_asociado, numero_cuenta_enmascarado,
                        cci_enmascarado, predeterminado
                    )
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        session["user_id"],
                        tipo_medio,
                        request.form.get("entidad", "").strip() or None,
                        titular,
                        request.form.get("telefono_asociado", "").strip() or None,
                        numero_cuenta or None,
                        cci or None,
                        int(predeterminado)
                    )
                )
                conn.commit()
                flash("Medio de pago guardado.", "success")
                return redirect(url_for("cuenta"))

            if action == "predeterminar_medio_pago":
                id_medio_pago = request.form.get("id_medio_pago", type=int)
                cur.execute(
                    """
                    SELECT id_medio_pago
                    FROM medios_pago_usuario
                    WHERE id_medio_pago = %s
                      AND id_usuario = %s
                      AND estado = 'ACTIVO'
                    """,
                    (id_medio_pago, session["user_id"])
                )
                if not cur.fetchone():
                    conn.rollback()
                    flash("El medio de pago no existe o no te pertenece.", "warning")
                    return redirect(url_for("cuenta"))
                cur.execute(
                    "UPDATE medios_pago_usuario SET predeterminado = 0 WHERE id_usuario = %s",
                    (session["user_id"],)
                )
                cur.execute(
                    "UPDATE medios_pago_usuario SET predeterminado = 1 WHERE id_medio_pago = %s",
                    (id_medio_pago,)
                )
                conn.commit()
                flash("Medio de pago predeterminado actualizado.", "success")
                return redirect(url_for("cuenta"))

            conn.rollback()
            flash("La acción solicitada no está disponible.", "warning")
        except Exception:
            conn.rollback()
            raise
        finally:
            cur.close()
            conn.close()

    facturaciones = query(
        """
        SELECT df.*, u.departamento, u.provincia, u.distrito
        FROM datos_facturacion df
        LEFT JOIN ubigeos u ON u.id_ubigeo = df.id_ubigeo
        WHERE df.id_comprador = %s
          AND df.estado = 'ACTIVO'
        ORDER BY df.predeterminado DESC, df.id_facturacion DESC
        """,
        (comprador["id_comprador"],)
    ) if comprador else []
    direcciones_entrega = query(
        """
        SELECT d.*, u.departamento, u.provincia, u.distrito
        FROM direcciones_entrega d
        LEFT JOIN ubigeos u ON u.id_ubigeo = d.id_ubigeo
        WHERE d.id_comprador = %s
          AND d.estado = 'ACTIVO'
        ORDER BY d.predeterminada DESC, d.id_direccion DESC
        """,
        (comprador["id_comprador"],)
    ) if comprador else []
    departamentos = query(
        """
        SELECT DISTINCT departamento
        FROM ubigeos
        WHERE estado = 'ACTIVO'
        ORDER BY departamento
        """
    ) if comprador else []
    medios_pago = query(
        """
        SELECT *
        FROM medios_pago_usuario
        WHERE id_usuario = %s
          AND estado = 'ACTIVO'
        ORDER BY predeterminado DESC, id_medio_pago DESC
        """,
        (session["user_id"],)
    )

    return render_template(
        "cuenta.html",
        user=user,
        facturaciones=facturaciones,
        direcciones_entrega=direcciones_entrega,
        departamentos=departamentos,
        medios_pago=medios_pago,
        es_comprador=bool(comprador)
    )

@app.route('/vida-agrovia')
def vida_agrovia():
    capacitaciones = query(
        """
        SELECT id_contenido, titulo, descripcion, categoria, tipo, imagen, video
        FROM contenidos_educativos
        WHERE estado = 'ACTIVO'
        ORDER BY fecha_creacion DESC
        """
    )

    eventos = query(
        """
        SELECT
            c.id_capacitacion,
            c.titulo,
            c.descripcion,
            c.fecha,
            DATE_FORMAT(c.fecha, '%%d') AS dia,
            CASE MONTH(c.fecha)
                WHEN 1 THEN 'ENE' WHEN 2 THEN 'FEB' WHEN 3 THEN 'MAR'
                WHEN 4 THEN 'ABR' WHEN 5 THEN 'MAY' WHEN 6 THEN 'JUN'
                WHEN 7 THEN 'JUL' WHEN 8 THEN 'AGO' WHEN 9 THEN 'SEP'
                WHEN 10 THEN 'OCT' WHEN 11 THEN 'NOV' WHEN 12 THEN 'DIC'
            END AS mes,
            COALESCE(TIME_FORMAT(c.hora, '%%H:%%i'), 'Hora por confirmar') AS hora,
            c.lugar AS modalidad,
            c.capacidad,
            (
                SELECT COUNT(*)
                FROM capacitacion_asistentes ca
                WHERE ca.id_capacitacion = c.id_capacitacion
                  AND ca.estado_asistencia <> 'CANCELADO'
            ) AS inscritos
        FROM capacitaciones c
        WHERE TIMESTAMP(c.fecha, c.hora) >= NOW()
          AND c.estado IN ('PROGRAMADA', 'ABIERTA')
        ORDER BY c.fecha, c.hora
        """
    )
    publicaciones = query(
        """
        SELECT id_publicacion, tipo, titulo, contenido, fecha_publicacion
        FROM publicaciones
        WHERE estado = 'PUBLICADA'
          AND (fecha_publicacion IS NULL OR fecha_publicacion <= NOW())
        ORDER BY fecha_publicacion DESC, id_publicacion DESC
        LIMIT 20
        """
    )

    return render_template(
        'vida_agrovia.html',
        capacitaciones=capacitaciones,
        eventos=eventos,
        publicaciones=publicaciones,
        today=datetime.now().strftime('%Y-%m-%d')
    )


@app.route('/vida-agrovia/capacitaciones', methods=['POST'])
@login_required
def crear_capacitacion():
    if not es_admin():
        flash('Solo el administrador puede publicar capacitaciones.', 'danger')
        return redirect(url_for('vida_agrovia'))

    titulo = request.form.get('titulo', '').strip()
    descripcion = request.form.get('descripcion', '').strip()
    fecha_texto = request.form.get('fecha', '').strip()
    hora_texto = request.form.get('hora', '').strip()
    lugar = request.form.get('lugar', '').strip()
    capacidad_texto = request.form.get('capacidad', '').strip()

    try:
        fecha_hora = datetime.strptime(
            f'{fecha_texto} {hora_texto}',
            '%Y-%m-%d %H:%M'
        )
        capacidad = int(capacidad_texto) if capacidad_texto else None
    except ValueError:
        flash('Ingresa una fecha, hora y capacidad válidas.', 'warning')
        return redirect(url_for('vida_agrovia'))

    if not titulo or not lugar or fecha_hora <= datetime.now():
        flash('Completa título y lugar, y elige una fecha y hora futuras.', 'warning')
        return redirect(url_for('vida_agrovia'))

    if capacidad is not None and capacidad < 1:
        flash('La capacidad debe ser mayor que cero.', 'warning')
        return redirect(url_for('vida_agrovia'))

    execute(
        """
        INSERT INTO capacitaciones
            (codigo_capacitacion, titulo, descripcion, fecha, hora, lugar, capacidad, estado)
        VALUES (%s, %s, %s, %s, %s, %s, %s, 'ABIERTA')
        """,
        (
            f'CAP-{uuid.uuid4().hex[:12].upper()}',
            titulo,
            descripcion or None,
            fecha_hora.date(),
            fecha_hora.time(),
            lugar,
            capacidad
        )
    )

    flash('La capacitación fue publicada y ya aparece en próximas capacitaciones.', 'success')
    return redirect(url_for('vida_agrovia'))


@app.route('/vida-agrovia/contenidos', methods=['POST'])
@login_required
def crear_contenido_educativo():
    if not es_admin():
        flash('Solo el administrador puede cargar contenidos.', 'danger')
        return redirect(url_for('vida_agrovia'))

    titulo = request.form.get('titulo', '').strip()
    descripcion = request.form.get('descripcion', '').strip()
    categoria = request.form.get('categoria', '').strip()
    tipo = request.form.get('tipo', '').strip()
    imagen = request.form.get('imagen', '').strip()
    video = request.files.get('video')
    categorias_permitidas = {'CAPACITACIONES', 'BUENAS PRÁCTICAS', 'TECNOLOGÍA'}
    extensiones_video = {'.mp4', '.webm', '.ogg'}
    imagenes_permitidas = {
        'cosecha choclo.png',
        'cosecha papa.png',
        'cosecha zanahoria.png'
    }

    if (
        not titulo or not descripcion or not tipo
        or categoria not in categorias_permitidas
        or imagen not in imagenes_permitidas
    ):
        flash('Completa todos los campos del contenido con opciones válidas.', 'warning')
        return redirect(url_for('vida_agrovia'))

    nombre_video = None
    if video and video.filename:
        nombre_original = secure_filename(video.filename)
        extension = os.path.splitext(nombre_original)[1].lower()
        if extension not in extensiones_video:
            flash('El video debe estar en formato MP4, WebM u OGG.', 'warning')
            return redirect(url_for('vida_agrovia'))

        nombre_video = f'{uuid.uuid4().hex}{extension}'
        directorio_videos = os.path.join(app.static_folder, 'uploads', 'contenidos')
        os.makedirs(directorio_videos, exist_ok=True)
        video.save(os.path.join(directorio_videos, nombre_video))

    try:
        execute(
            """
            INSERT INTO contenidos_educativos
                (titulo, descripcion, categoria, tipo, imagen, video, id_usuario_creador, estado)
            VALUES (%s, %s, %s, %s, %s, %s, %s, 'ACTIVO')
            """,
            (
                titulo,
                descripcion,
                categoria,
                tipo,
                imagen,
                nombre_video,
                session['user_id']
            )
        )
    except Exception:
        if nombre_video:
            ruta_video = os.path.join(directorio_videos, nombre_video)
            if os.path.exists(ruta_video):
                os.remove(ruta_video)
        raise

    flash('El contenido educativo ya está publicado.', 'success')
    return redirect(url_for('vida_agrovia'))


@app.route('/vida-agrovia/contenidos/<int:id_contenido>/video/eliminar', methods=['POST'])
@login_required
def eliminar_video_contenido(id_contenido):
    if not es_admin():
        flash('Solo el administrador puede eliminar videos.', 'danger')
        return redirect(url_for('vida_agrovia'))

    contenido = query(
        """
        SELECT video
        FROM contenidos_educativos
        WHERE id_contenido = %s AND estado = 'ACTIVO'
        """,
        (id_contenido,),
        one=True
    )
    if not contenido or not contenido['video']:
        flash('Este contenido no tiene un video para eliminar.', 'warning')
        return redirect(url_for('vida_agrovia'))

    execute(
        "UPDATE contenidos_educativos SET video = NULL WHERE id_contenido = %s",
        (id_contenido,)
    )

    nombre_video = os.path.basename(contenido['video'])
    ruta_video = os.path.join(app.static_folder, 'uploads', 'contenidos', nombre_video)
    if os.path.isfile(ruta_video):
        os.remove(ruta_video)

    flash('Se eliminó el video; el contenido permanece publicado.', 'success')
    return redirect(url_for('vida_agrovia'))


@app.route('/vida-agrovia/contenidos/<int:id_contenido>/editar', methods=['POST'])
@login_required
def actualizar_contenido_educativo(id_contenido):
    if not tiene_permiso('PUBLICACIONES_GESTIONAR'):
        flash('No tienes permiso para editar contenidos educativos.', 'danger')
        return redirect(url_for('vida_agrovia'))

    titulo = request.form.get('titulo', '').strip()
    descripcion = request.form.get('descripcion', '').strip()
    categoria = request.form.get('categoria', '').strip()
    tipo = request.form.get('tipo', '').strip()
    imagen = request.form.get('imagen', '').strip()
    if (
        not titulo or len(titulo) > 180 or not descripcion or not tipo or len(tipo) > 80
        or categoria not in {'CAPACITACIONES', 'BUENAS PRÁCTICAS', 'TECNOLOGÍA'}
        or imagen not in {'cosecha choclo.png', 'cosecha papa.png', 'cosecha zanahoria.png'}
    ):
        flash('Revisa los campos del contenido educativo.', 'warning')
        return redirect(url_for('vida_agrovia'))

    execute(
        """
        UPDATE contenidos_educativos
        SET titulo = %s, descripcion = %s, categoria = %s, tipo = %s, imagen = %s
        WHERE id_contenido = %s AND estado = 'ACTIVO'
        """,
        (titulo, descripcion, categoria, tipo, imagen, id_contenido)
    )
    flash('Contenido educativo actualizado.', 'success')
    return redirect(url_for('vida_agrovia'))


@app.route('/vida-agrovia/contenidos/<int:id_contenido>/eliminar', methods=['POST'])
@login_required
def eliminar_contenido_educativo(id_contenido):
    if not tiene_permiso('PUBLICACIONES_GESTIONAR'):
        flash('No tienes permiso para retirar contenidos educativos.', 'danger')
        return redirect(url_for('vida_agrovia'))

    execute(
        "UPDATE contenidos_educativos SET estado = 'INACTIVO' WHERE id_contenido = %s AND estado = 'ACTIVO'",
        (id_contenido,)
    )
    flash('Contenido retirado de Vida Agrovia.', 'success')
    return redirect(url_for('vida_agrovia'))


@app.route('/admin/capacitaciones')
@login_required
def admin_capacitaciones():
    if not tiene_permiso('CAPACITACIONES_GESTIONAR'):
        flash('No tienes permiso para gestionar capacitaciones.', 'danger')
        return redirect(url_for('dashboard'))

    capacitaciones = query(
        """
         SELECT c.id_capacitacion, c.codigo_capacitacion, c.titulo, c.fecha,
             c.hora, c.lugar, c.capacidad, c.estado,
             (
                 SELECT COUNT(*)
                 FROM capacitacion_asistentes ca
                 WHERE ca.id_capacitacion = c.id_capacitacion
                AND ca.estado_asistencia <> 'CANCELADO'
             ) AS inscritos
        FROM capacitaciones c
        ORDER BY c.fecha DESC, c.hora DESC
        """
    )
    return render_template("admin_capacitaciones.html", capacitaciones=capacitaciones)


@app.route('/admin/capacitaciones/<int:id_capacitacion>/asistentes', methods=['GET', 'POST'])
@login_required
def admin_capacitacion_asistentes(id_capacitacion):
    if not tiene_permiso('CAPACITACIONES_GESTIONAR'):
        flash('No tienes permiso para gestionar asistentes.', 'danger')
        return redirect(url_for('dashboard'))

    capacitacion = query(
        "SELECT * FROM capacitaciones WHERE id_capacitacion = %s",
        (id_capacitacion,),
        one=True
    )
    if not capacitacion:
        flash('La capacitación no existe.', 'warning')
        return redirect(url_for('admin_capacitaciones'))

    if request.method == 'POST':
        id_asistencia = request.form.get('id_asistencia', type=int)
        estado = request.form.get('estado_asistencia', '')
        estados_validos = {'REGISTRADO', 'CONFIRMADO', 'ASISTIO', 'NO_ASISTIO', 'CANCELADO'}
        if not id_asistencia or estado not in estados_validos:
            flash('El estado de asistencia no es válido.', 'warning')
            return redirect(url_for('admin_capacitacion_asistentes', id_capacitacion=id_capacitacion))
        execute(
            """
            UPDATE capacitacion_asistentes
            SET estado_asistencia = %s,
                confirmado = IF(%s IN ('CONFIRMADO', 'ASISTIO', 'NO_ASISTIO'), 1, confirmado)
            WHERE id_asistencia = %s AND id_capacitacion = %s
            """,
            (estado, estado, id_asistencia, id_capacitacion)
        )
        flash('Estado del asistente actualizado.', 'success')
        return redirect(url_for('admin_capacitacion_asistentes', id_capacitacion=id_capacitacion))

    asistentes = query(
        """
        SELECT ca.id_asistencia, ca.estado_asistencia, ca.confirmado, ca.fecha_registro,
               u.nombres, u.apellido_paterno, u.correo, u.telefono
        FROM capacitacion_asistentes ca
        INNER JOIN usuarios u ON u.id_usuario = ca.id_usuario
        WHERE ca.id_capacitacion = %s
        ORDER BY ca.fecha_registro, ca.id_asistencia
        """,
        (id_capacitacion,)
    )
    return render_template(
        "admin_capacitacion_asistentes.html",
        capacitacion=capacitacion,
        asistentes=asistentes,
        estados_asistencia=('REGISTRADO', 'CONFIRMADO', 'ASISTIO', 'NO_ASISTIO', 'CANCELADO')
    )


@app.route('/vida-agrovia/capacitaciones/<int:id_capacitacion>/inscribirse', methods=['POST'])
@login_required
def inscribir_capacitacion(id_capacitacion):
    capacitacion = query(
        """
        SELECT c.id_capacitacion, c.capacidad
        FROM capacitaciones c
        WHERE c.id_capacitacion = %s
          AND TIMESTAMP(c.fecha, c.hora) >= NOW()
          AND c.estado IN ('PROGRAMADA', 'ABIERTA')
        """,
        (id_capacitacion,),
        one=True
    )

    if not capacitacion:
        flash('Esta capacitación ya no está disponible para inscripción.', 'warning')
        return redirect(url_for('vida_agrovia'))

    existente = query(
        """
        SELECT id_asistencia
        FROM capacitacion_asistentes
        WHERE id_capacitacion = %s AND id_usuario = %s
          AND estado_asistencia <> 'CANCELADO'
        """,
        (id_capacitacion, session['user_id']),
        one=True
    )
    if existente:
        flash('Ya estás inscrito en esta capacitación.', 'info')
        return redirect(url_for('vida_agrovia'))

    if capacitacion['capacidad'] is not None:
        inscritos = query(
            """
            SELECT COUNT(*) AS total
            FROM capacitacion_asistentes
            WHERE id_capacitacion = %s AND estado_asistencia <> 'CANCELADO'
            """,
            (id_capacitacion,),
            one=True
        )
        if inscritos['total'] >= capacitacion['capacidad']:
            flash('La capacitación ya completó sus cupos.', 'warning')
            return redirect(url_for('vida_agrovia'))

    insertado = execute(
        """
        INSERT IGNORE INTO capacitacion_asistentes
            (id_capacitacion, id_usuario, estado_asistencia)
        VALUES (%s, %s, 'REGISTRADO')
        """,
        (id_capacitacion, session['user_id'])
    )

    if insertado:
        flash('Tu inscripción quedó registrada.', 'success')
    else:
        flash('Ya estás inscrito en esta capacitación.', 'info')

    return redirect(url_for('vida_agrovia'))

# ==========================================================
# EJECUCIÓN
# ==========================================================

if __name__ == "__main__":

    app.run(
        debug=False,
        host="127.0.0.1",
        port=5000
    )